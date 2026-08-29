# -*- coding: utf-8 -*-
"""PDF release versions for all deliverables (reportlab). Blueprint palette."""
import os, re
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable)
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.abspath(__file__))
DLV = os.path.join(ROOT, "deliverables")
NAVY = colors.HexColor("#14283C")
BLUE = colors.HexColor("#1E5A8A")
ORANGE = colors.HexColor("#C97B3D")
GREY = colors.HexColor("#5A6A7A")

def esc(t):
    return (t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

def inline(t):
    t = esc(t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = t.replace("✓", "PASS").replace("✗", "FAIL")
    return t

ST = {
 "title": ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=17, textColor=NAVY, leading=21, spaceAfter=2),
 "sub":   ParagraphStyle("s", fontName="Helvetica", fontSize=9.5, textColor=GREY, leading=13, spaceAfter=6),
 "h2":    ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=13, textColor=BLUE, leading=17, spaceBefore=10, spaceAfter=4),
 "h3":    ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=11, textColor=NAVY, leading=14, spaceBefore=8, spaceAfter=3),
 "body":  ParagraphStyle("b", fontName="Helvetica", fontSize=9, leading=12.5, spaceAfter=3),
 "cell":  ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10.5),
 "cellh": ParagraphStyle("ch", fontName="Helvetica-Bold", fontSize=8, leading=10.5, textColor=colors.white),
}

def md_to_flow(md_text):
    flow, table_buf = [], []
    lines = md_text.splitlines()

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = [[Paragraph(inline(c.strip()), ST["cellh"] if i == 0 else ST["cell"])
                 for c in r] for i, r in enumerate(table_buf)]
        t = Table(rows, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), BLUE),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C6D4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        flow.append(t)
        flow.append(Spacer(1, 6))
        table_buf = []

    for ln in lines:
        s = ln.strip()
        if s.startswith("|"):
            cells = [c for c in s.strip("|").split("|")]
            table_buf.append(cells)
        elif re.match(r"^\s*\|[\s\-:|]+\|\s*$", s) or s == "---" or not s:
            continue
        else:
            flush_table()
            if s.startswith("# "):
                flow.append(Paragraph(inline(s[2:]), ST["title"]))
            elif s.startswith("## "):
                flow.append(HRFlowable(width="100%", thickness=1.2, color=ORANGE, spaceBefore=2, spaceAfter=4))
                flow.append(Paragraph(inline(s[3:]), ST["h2"]))
            elif s.startswith("### "):
                flow.append(Paragraph(inline(s[4:]), ST["h3"]))
            else:
                flow.append(Paragraph(inline(s), ST["body"]))
    flush_table()
    return flow

def md_pdf(src, dst):
    doc = SimpleDocTemplate(dst, pagesize=A4,
                            leftMargin=16*mm, rightMargin=16*mm,
                            topMargin=14*mm, bottomMargin=14*mm,
                            title=os.path.basename(dst))
    doc.build(md_to_flow(open(src, encoding="utf-8").read()))

def xlsx_pdf(src, dst):
    wb = load_workbook(src, read_only=True)
    flow = []
    for ws in wb.worksheets:
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        rows = [[("" if v is None else str(v)) for v in r] for r in rows if any(v is not None for v in r)]
        if not rows:
            continue
        flow.append(Paragraph(esc(ws.title), ST["h2"]))
        tbl = [[Paragraph(inline(c), ST["cellh"] if i == 0 else ST["cell"]) for c in r]
               for i, r in enumerate(rows)]
        t = Table(tbl, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), BLUE),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C6D4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        flow.append(t)
        flow.append(Spacer(1, 8))
    doc = SimpleDocTemplate(dst, pagesize=A4,
                            leftMargin=16*mm, rightMargin=16*mm,
                            topMargin=14*mm, bottomMargin=14*mm,
                            title=os.path.basename(dst))
    doc.build(flow)

pairs = [
    ("design_spec.md", "design_spec.pdf"),
    ("datasheet.md", "datasheet.pdf"),
    ("dvpr.md", "dvpr.pdf"),
    ("dfmea.md", "dfmea.pdf"),
    ("delivery_index.md", "delivery_index.pdf"),
    ("bom.xlsx", "bom.pdf"),
    ("calc.xlsx", "calc.pdf"),
]
for src, dst in pairs:
    sp = os.path.join(DLV, src); dp = os.path.join(DLV, dst)
    if src.endswith(".xlsx"):
        xlsx_pdf(sp, dp)
    else:
        md_pdf(sp, dp)
    print(dst, os.path.getsize(dp), "bytes")
print("PDF release generation complete")
