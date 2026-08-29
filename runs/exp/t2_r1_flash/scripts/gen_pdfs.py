"""Generate PDF release versions of all deliverables (reportlab, engineering blueprint style #14283C/#1E5A8A/#C97B3D)."""
import re
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

DELIV = Path(__file__).resolve().parent.parent / "deliverables"

DEEP = colors.HexColor("#14283C")
MED = colors.HexColor("#1E5A8A")
COPPER = colors.HexColor("#C97B3D")

styles = {
    "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=15, textColor=DEEP, spaceAfter=8),
    "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=12, textColor=MED, spaceBefore=10, spaceAfter=4),
    "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.5, textColor=MED, spaceBefore=8, spaceAfter=3),
    "body": ParagraphStyle("body", fontName="Helvetica", fontSize=8.5, leading=11, spaceAfter=3),
    "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=7.5, leading=9),
    "cellb": ParagraphStyle("cellb", fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white),
}


def escape(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def md_to_flowables(md_text):
    fl = []
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            fl.append(Paragraph(escape(line[2:]), styles["title"]))
        elif line.startswith("## "):
            fl.append(Paragraph(escape(line[3:]), styles["h1"]))
        elif line.startswith("### "):
            fl.append(Paragraph(escape(line[4:]), styles["h2"]))
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not any(cells):
                continue
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
                continue  # separator row
            rows_data = [[Paragraph(escape(c), styles["cell"]) for c in cells]]
            if len(fl) > 0 and isinstance(fl[-1], list):
                fl[-1].append(rows_data[0])
            else:
                fl.append(rows_data)
        elif line.startswith("- "):
            fl.append(Paragraph("• " + escape(line[2:]), styles["body"]))
        elif line.startswith("**"):
            fl.append(Paragraph(escape(line), styles["body"]))
        else:
            fl.append(Paragraph(escape(line), styles["body"]))
    out = []
    for item in fl:
        if isinstance(item, list):
            ncols = max(len(r) for r in item)
            tbl = Table([r + [Paragraph("", styles["cell"])] * (ncols - len(r)) for r in item],
                        repeatRows=1, hAlign="LEFT")
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), MED),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9FB4C7")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F7")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]))
            out.append(tbl)
            out.append(Spacer(1, 5))
        else:
            out.append(item)
    return out


def xlsx_to_flowables(path):
    wb = load_workbook(path, data_only=True)
    fl = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        fl.append(Paragraph(escape(ws.title), styles["h1"]))
        data = [[Paragraph(escape("" if c is None else str(c)), styles["cell"]) for c in r] for r in rows]
        ncols = max(len(r) for r in data)
        data = [r + [Paragraph("", styles["cell"])] * (ncols - len(r)) for r in data]
        tbl = Table(data, repeatRows=1, hAlign="LEFT")
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), DEEP),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9FB4C7")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F7")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        fl.append(tbl)
        fl.append(Spacer(1, 8))
    return fl


def make_pdf(out_name, flowables):
    doc = SimpleDocTemplate(str(DELIV / out_name), pagesize=A4,
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=out_name)
    doc.build(flowables)
    print("pdf:", out_name)


md_files = {
    "design_spec.md": "design_spec.pdf",
    "datasheet.md": "datasheet.pdf",
    "dvpr.md": "dvpr.pdf",
    "dfmea.md": "dfmea.pdf",
    "delivery_index.md": "delivery_index.pdf",
}
for src, dst in md_files.items():
    make_pdf(dst, md_to_flowables((DELIV / src).read_text(encoding="utf-8")))

make_pdf("bom.pdf", xlsx_to_flowables(DELIV / "bom.xlsx"))
make_pdf("calc.pdf", xlsx_to_flowables(DELIV / "calc.xlsx"))
print("ALL PDFs generated")
