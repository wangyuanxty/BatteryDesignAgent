# One-off generator for bom.xlsx and calc.xlsx (openpyxl). All values from tool outputs
# (cell/r8_d6_*.json, parameter-set dump, calc-energy contract formulas).
import openpyxl
from openpyxl.styles import Font

OUT = r"D:\research\degradation_prognostics\Battery_Design_Agent\runs\exp\t6_r1\deliverables"
AREA = 0.1027          # cell/r8_d6_energy_dfn.json:area_m2
E_WH = 25.40591018     # cell/r8_d6_energy_dfn.json:energy_wh
E_KWH = E_WH / 1000.0
MASS_KG = 0.05149989784  # energy json:mass_kg (electrolyte excluded)

LAYERS = {  # energy json:layer_kg_m2
    "Positive electrode coating (LNMO)": 0.2212056,
    "Negative electrode coating (graphite)": 0.12700905,
    "Positive current collector (Al)": 0.0432,
    "Negative current collector (Cu)": 0.10752,
    "Separator": 0.00252492,
}
ELEC_MASS_G = 6.27  # design estimate: 5.804 cm3 pore x 1.2 g/cm3 x 0.9 fill (literature density)

# ---------------- bom.xlsx ----------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BOM"
ws.append(["Component", "Mass g/cell", "Mass kg/kWh", "Basis / source"])
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True); ws["C1"].font = Font(bold=True); ws["D1"].font = Font(bold=True)
for name, kgpm2 in LAYERS.items():
    g = kgpm2 * AREA * 1000.0
    ws.append([name, round(g, 2), round(g / 1000.0 / E_KWH, 3),
               f"cell/r8_d6_energy_dfn.json:layer_kg_m2[{name}] x area_m2"])
ws.append(["Electrolyte (EC/EMC + LiPF6)", ELEC_MASS_G, round(ELEC_MASS_G / 1000.0 / E_KWH, 3),
           "design estimate: pore volume 5.804 cm3 x 1.2 g/cm3 (literature) x 0.9 fill"])
ws.append(["Conductive additive (positive)", "estimate 0.34 (1.5 wt% of 22.72 g)", "—", "literature default 1.5 wt% (estimate, not parameterized)"])
ws.append(["Binder (positive)", "estimate 0.45 (2.0 wt% of 22.72 g)", "—", "literature default 2.0 wt% (estimate, not parameterized)"])
ws.append(["Conductive additive (negative)", "estimate 0.13 (1.0 wt% of 13.04 g)", "—", "literature default 1.0 wt% (estimate, not parameterized)"])
ws.append(["Binder (negative)", "estimate 0.39 (3.0 wt% of 13.04 g)", "—", "literature default 3.0 wt% (estimate, not parameterized)"])
ws.append(["Enclosure and tabs", "Not modeled", "—", "beyond pure simulation boundary"])
ws.append(["TOTAL (electrolyte excluded, contract caliber)", round(MASS_KG * 1000, 2), round(MASS_KG / E_KWH, 3),
           "cell/r8_d6_energy_dfn.json:mass_kg"])
ws.append(["TOTAL incl. electrolyte estimate", round(MASS_KG * 1000 + ELEC_MASS_G, 2), round((MASS_KG * 1000 + ELEC_MASS_G) / 1000.0 / E_KWH, 3), "computed"])
ws.append(["Total energy", round(E_WH, 2), "Wh", "cell/r8_d6_energy_dfn.json:energy_wh"])
wb.save(OUT + r"\bom.xlsx")

# ---------------- calc.xlsx ----------------
wb2 = openpyxl.Workbook()
# Sheet 1: inputs
ws = wb2.active
ws.title = "inputs"
ws.append(["Parameter", "Value", "Source"])
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True); ws["C1"].font = Font(bold=True)
inputs = [
    ("Positive electrode thickness [m]", 7.56e-5, "parameter set dump"),
    ("Negative electrode thickness [m]", 1.022e-4, "bridge/params_final.json"),
    ("Separator thickness [m]", 1.2e-5, "parameter set dump"),
    ("Positive CC thickness [m]", 1.6e-5, "parameter set dump"),
    ("Negative CC thickness [m]", 1.2e-5, "parameter set dump"),
    ("Positive electrode porosity", 0.335, "parameter set dump"),
    ("Negative electrode porosity", 0.25, "parameter set dump"),
    ("Separator porosity", 0.47, "parameter set dump"),
    ("Positive active material volume fraction", 0.665, "parameter set dump"),
    ("Negative active material volume fraction", 0.75, "parameter set dump"),
    ("Positive particle radius [m]", 5.22e-6, "parameter set dump"),
    ("Negative particle radius [m]", 3.5e-6, "bridge/params_final.json"),
    ("Positive electrode density [kg.m-3]", 4400.0, "parameter set dump"),
    ("Negative electrode density [kg.m-3]", 1657.0, "parameter set dump"),
    ("Maximum concentration in positive electrode [mol.m-3]", 63104.0, "parameter set dump (Chen2020; LNMO.json key inert - recorded)"),
    ("Maximum concentration in negative electrode [mol.m-3]", 33133.0, "parameter set dump"),
    ("Electrode height [m]", 0.065, "parameter set dump"),
    ("Electrode width [m]", 1.58, "parameter set dump"),
    ("Electrode area [m2]", AREA, "cell/r8_d6_energy_dfn.json:area_m2"),
    ("Lower / Upper voltage cut-off [V]", "2.5 / 4.7", "parameter set dump"),
    ("Cation transference number", 0.2594, "parameter set dump"),
    ("Electrolyte diffusivity [m2.s-1]", 6.0e-10, "bridge/params_final.json"),
    ("EC initial concentration in electrolyte [mol.m-3]", 2270.5, "bridge/params_final.json"),
    ("EC diffusivity [m2.s-1]", 1.0e-18, "bridge/params_final.json"),
    ("Total heat transfer coefficient [W.m-2.K-1]", 400.0, "bridge/params_final.json"),
    ("Faraday constant [C.mol-1]", 96485.0, "physical constant"),
]
for row in inputs:
    ws.append(list(row))
# Sheet 2: capacity and energy
ws = wb2.create_sheet("capacity_energy")
ws.append(["Quantity", "Value", "Formula", "Source"])
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True); ws["C1"].font = Font(bold=True); ws["D1"].font = Font(bold=True)
ws.append(["1C discharge capacity [Ah]", 6.1976, "time integral of I at 1C to 2.5 V", "cell/r8_d6_1c_dfn.json:capacity_ah"])
ws.append(["Discharge energy [Wh]", 25.4059, "integral(V x I_1C) dt", "cell/r8_d6_energy_dfn.json:energy_wh"])
ws.append(["Midpoint voltage [V]", 4.1145, "V at t=len(t)/2 (DFN uniform grid, ~70% depth)", "cell/r8_d6_energy_dfn.json:midpoint_voltage_v"])
ws.append(["Cell volume [m3]", 2.2368e-5, "sum(layer thickness) x area", "cell/r8_d6_energy_dfn.json:volume_m3"])
ws.append(["DCR [Ohm]", 6.38e-3, "midpoint DCR (calc-energy)", "cell/r8_d6_energy_dfn.json:dcr_ohm"])
# Sheet 3: energy density
ws = wb2.create_sheet("energy_density")
ws.append(["Quantity", "Value", "Formula", "Source"])
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True); ws["C1"].font = Font(bold=True); ws["D1"].font = Font(bold=True)
ws.append(["Volumetric ED [Wh/L]", 1135.8, "energy_wh / volume_m3 / 1000", "cell/r8_d6_energy_dfn.json:energy_density_wh_l"])
ws.append(["Gravimetric ED [Wh/kg]", 493.3, "energy_wh / mass_kg", "cell/r8_d6_energy_dfn.json:energy_density_wh_kg"])
ws.append(["Mass [kg]", 0.0515, "sum(layer_kg_m2 x area_m2)", "cell/r8_d6_energy_dfn.json:mass_kg"])
ws.append(["Criterion check", "1135.8 >= 950 -> PASS (+19.6%)", "", "entry 0 criteria.stage2"])
# Sheet 4: N/P and mass
ws = wb2.create_sheet("np_mass")
ws.append(["Quantity", "Value", "Formula", "Source"])
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True); ws["C1"].font = Font(bold=True); ws["D1"].font = Font(bold=True)
ws.append(["Positive areal capacity [Ah/m2]", 85.03, "F x c_max_pos x eps_am_pos x L_pos / 3600 = 96485 x 63104 x 0.665 x 75.6e-6 / 3600", "parameter set dump"])
ws.append(["Negative areal capacity [Ah/m2]", 68.07, "F x c_max_neg x eps_am_neg x L_neg / 3600 = 96485 x 33133 x 0.75 x 102.2e-6 / 3600", "parameter set dump"])
ws.append(["N/P ratio", 0.80, "68.07 / 85.03 (capacity ratio; supersedes round shorthand 'N/P 1.2' = thickness-ratio mislabel - recorded)", "computed from parameter set"])
for name, kgpm2 in LAYERS.items():
    ws.append([f"Layer mass {name} [kg]", round(kgpm2 * AREA, 6), "layer_kg_m2 x area_m2", "cell/r8_d6_energy_dfn.json"])
ws.append(["Total mass [kg]", 0.0515, "sum of layers", "cell/r8_d6_energy_dfn.json:mass_kg"])
# Sheet 5: process parameters
ws = wb2.create_sheet("process")
ws.append(["Parameter", "Value", "Formula", "Source"])
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True); ws["C1"].font = Font(bold=True); ws["D1"].font = Font(bold=True)
ws.append(["Positive areal density [g/m2]", 221.2, "75.6e-6 x 0.665 x 4400 x 1000", "parameter set dump"])
ws.append(["Negative areal density [g/m2]", 127.0, "102.2e-6 x 0.75 x 1657 x 1000", "parameter set dump"])
ws.append(["Positive compaction density [g/cm3]", 2.93, "4400 x 0.665 / 1000", "parameter set dump"])
ws.append(["Negative compaction density [g/cm3]", 1.24, "1657 x 0.75 / 1000", "parameter set dump"])
ws.append(["Pore volume [cm3]", 5.804, "sum(thickness x porosity x area) x 1e6", "parameter set dump"])
ws.append(["Electrolyte fill [g]", 6.27, "5.804 x 1.2 x 0.9 (density = literature 1.2 g/cm3)", "design estimate"])
ws.append(["Formation", "0.1C CC to 4.7 V, 25 C, 2 cycles", "", "design-recommended value; production tuning required"])
wb2.save(OUT + r"\calc.xlsx")
print("bom.xlsx + calc.xlsx written")
