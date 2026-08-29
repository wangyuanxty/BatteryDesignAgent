# -*- coding: utf-8 -*-
"""Generate bom.xlsx, calc.xlsx, datasheet.docx for t6_r1_noceiling deliverables."""
import sys
sys.path.insert(0, r"D:\research\degradation_prognostics\Battery_Design_Agent\runs\exp\t6_r1_noceiling")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt

from deliverable_data import (DELIV, BOM_ROWS, CALC_SHEETS, DATASHEET_FIELDS,
                              MASS_G, MASS_ELEC_G, ENERGY_WH, g_to_kg_kwh,
                              MASS_POS_BINDER_G, MASS_POS_CARBON_G, MASS_NEG_BINDER_G, MASS_NEG_CARBON_G)

HDR_FILL = PatternFill("solid", fgColor="1E5A8A")
HDR_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="9DB2C8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_sheet(ws, headers, widths):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill, c.font, c.alignment = HDR_FILL, HDR_FONT, Alignment(vertical="center", wrap_text=True)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BORDER
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    ws.freeze_panes = "A2"


def make_bom():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    headers = ["Component", "g/cell", "kg/kWh", "Annotation / source"]
    rows = []
    for comp, g, kgkwh, note in BOM_ROWS:
        if g is None:
            rows.append([comp, "Not modeled", "Not modeled", note])
        else:
            rows.append([comp, round(g, 4), round(kgkwh, 4), note])
    grand = MASS_G + MASS_ELEC_G + MASS_POS_BINDER_G + MASS_POS_CARBON_G + MASS_NEG_BINDER_G + MASS_NEG_CARBON_G
    rows.append(["TOTAL (contract caliber, electrolyte excluded)", round(MASS_G, 4), round(g_to_kg_kwh(MASS_G), 4),
                 "electrodes + current collectors + separator (calc-energy mass_kg)"])
    rows.append(["TOTAL (all modeled components)", round(grand, 4), round(g_to_kg_kwh(grand), 4),
                 "contract total + electrolyte + binder/additive literature defaults"])
    rows.append(["Cell energy (kWh)", "-", round(ENERGY_WH / 1000.0, 6),
                 "calc-energy output energy_wh / 1000 (18.6012 Wh)"])
    for r in rows:
        ws.append(r)
    style_sheet(ws, headers, [46, 12, 12, 80])
    return wb


def make_calc():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in CALC_SHEETS:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append([r[0], r[1], r[2], r[3], r[4]])
        style_sheet(ws, ["Quantity", "Expression", "Value", "Unit", "Source"], [34, 46, 16, 12, 52])
    return wb


def make_datasheet():
    doc = Document()
    doc.add_heading("Technical Datasheet - VBF-T6R1NOCEILING-DSH-01", level=0)
    doc.add_paragraph("ED-Compact B-p smartphone cell (virtual design). All values mechanically taken from "
                      "parameter set / simulation outputs and annotated line by line; missing items written honestly.")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Field", "Value (with source)"
    for field, value in DATASHEET_FIELDS:
        row = table.add_row().cells
        row[0].text = field
        row[1].text = value
    for r in table.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return doc


def main():
    DELIV.mkdir(exist_ok=True)
    make_bom().save(str(DELIV / "bom.xlsx"))
    print("bom.xlsx written")
    make_calc().save(str(DELIV / "calc.xlsx"))
    print("calc.xlsx written")
    make_datasheet().save(str(DELIV / "datasheet.docx"))
    print("datasheet.docx written")


if __name__ == "__main__":
    main()
