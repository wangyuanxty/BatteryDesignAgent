# -*- coding: utf-8 -*-
"""t4_r1 deliverables generator: bom.xlsx, calc.xlsx, delivery_index.md, and PDF releases
(md -> PDF via a small markdown-subset renderer; xlsx -> PDF from the same data structures).
Values are mechanically taken from tool outputs / parameter set; no numbers from memory."""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

WS = Path(__file__).resolve().parent
DEL = WS / "deliverables"
CELL = WS / "cell"
DEL.mkdir(exist_ok=True)

# ---------------------------------------------------------------- source data
def load(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)

en = load(CELL / "r6_t2_energy.json")
c1 = load(CELL / "r6_t2_1c_dfn.json")
cl = load(CELL / "r6_t2_lowT_dfn.json")
c4 = load(CELL / "r6_t2_4c45C_dfn.json")
ret = load(CELL / "r6_t2_retention.json")

AREA = 0.065 * 1.58
L = {"pos": 75.6e-6, "sep": 10e-6, "neg": 85.2e-6}
EPS = {"pos": 0.335, "sep": 0.47, "neg": 0.25}
RHO = {"pos": 3262.0, "neg": 1657.0}
KWH = en["energy_wh"] / 1000.0
pore_ml = (L["pos"] * EPS["pos"] + L["sep"] * EPS["sep"] + L["neg"] * EPS["neg"]) * AREA * 1e6
m_elyte = pore_ml * 1.2  # g (1.2 g/cm3 literature default)
lkg = en["layer_kg_m2"]
m_pos, m_neg, m_al, m_cu, m_sep = (lkg["positive_electrode"] * AREA * 1000,
                                   lkg["negative_electrode"] * AREA * 1000,
                                   lkg["positive_cc"] * AREA * 1000,
                                   lkg["negative_cc"] * AREA * 1000,
                                   lkg["separator"] * AREA * 1000)
total_nofill = en["mass_kg"] * 1000
total_fill = total_nofill + m_elyte

# ================================================================ BOM xlsx
wb = Workbook()
ws = wb.active
ws.title = "BOM"
hdr = ["Component", "Basis / formula", "g/cell", "kg/kWh", "Source"]
ws.append(hdr)
bom_rows = [
    ("Positive active material (NMC811)", "96 wt% of pos coating (literature default)", m_pos * 0.96, "cell/r6_t2_energy.json:layer_kg_m2 x area x 0.96"),
    ("Positive conductive additive (carbon)", "2 wt% (literature default)", m_pos * 0.02, "same x 0.02"),
    ("Positive binder (PVDF)", "2 wt% (literature default)", m_pos * 0.02, "same x 0.02"),
    ("Negative active material (graphite)", "95.5 wt% of neg coating (literature default)", m_neg * 0.955, "cell/r6_t2_energy.json:layer_kg_m2 x area x 0.955"),
    ("Negative conductive additive (carbon)", "1 wt% (literature default)", m_neg * 0.01, "same x 0.01"),
    ("Negative binder (SBR/CMC)", "3.5 wt% (literature default)", m_neg * 0.035, "same x 0.035"),
    ("Separator (polyolefin)", "layer mass formula", m_sep, "cell/r6_t2_energy.json:layer_kg_m2"),
    ("Electrolyte", "pore volume x 1.2 g/cm3 (literature default)", m_elyte, "pore volume = (L_pos*eps_pos + L_sep*eps_sep + L_neg*eps_neg) x area = %.3f mL" % pore_ml),
    ("Positive current collector (Al)", "10 um x 2700 kg/m3", m_al, "cell/r6_t2_energy.json:layer_kg_m2"),
    ("Negative current collector (Cu)", "8 um x 8960 kg/m3", m_cu, "cell/r6_t2_energy.json:layer_kg_m2"),
    ("Enclosure / tabs", "Not modeled", None, "beyond pure simulation boundary"),
    ("TOTAL (electrolyte excluded)", "contract cell mass", total_nofill, "cell/r6_t2_energy.json:mass_kg"),
    ("TOTAL (incl. electrolyte fill)", "sum of above", total_fill, "sum"),
    ("Cell energy", "simulated 1C discharge, 25 C", None, "cell/r6_t2_energy.json:energy_wh = %.4f Wh" % en["energy_wh"]),
]
for name, basis, g, src in bom_rows:
    kgkwh = (g / 1000 / KWH) if g is not None else None
    ws.append([name, basis, round(g, 3) if g is not None else "Not provided",
               round(kgkwh, 3) if kgkwh is not None else "Not provided", src])
for c_idx, wdt in zip(range(1, 6), [42, 55, 12, 12, 60]):
    ws.column_dimensions[get_column_letter(c_idx)].width = wdt
for cell in ws[1]:
    cell.font = Font(bold=True)
wb.save(DEL / "bom.xlsx")

# ================================================================ CALC xlsx
wb2 = Workbook()
def sheet(title, header, rows):
    w = wb2.create_sheet(title)
    w.append(header)
    for r in rows:
        w.append(r)
    for c_idx, wdt in zip(range(1, len(header) + 1), [38, 60, 16, 10, 60][: len(header)]):
        w.column_dimensions[get_column_letter(c_idx)].width = wdt
    for cell in w[1]:
        cell.font = Font(bold=True)

inp = [
    ("Current collectors / separator / particle / transport / cooling overrides", "design overrides (bridge/params_r6_t2.json)", "see bridge file", "-", "9 keys: CC 10/8 um, sep 10 um, R_neg 2.5 um, t+ 0.6, kappa 1.1, D_e 3e-10, h 80"),
    ("Positive electrode thickness", "parameter set", "75.6", "um", "Chen2020"),
    ("Negative electrode thickness", "parameter set", "85.2", "um", "Chen2020"),
    ("Separator thickness", "override", "10.0", "um", "bridge/params_r6_t2.json"),
    ("Positive electrode porosity", "parameter set", "0.335", "-", "Chen2020"),
    ("Negative electrode porosity", "parameter set", "0.25", "-", "Chen2020"),
    ("Separator porosity", "parameter set", "0.47", "-", "Chen2020"),
    ("Positive electrode density", "parameter set", "3262", "kg/m3", "Chen2020"),
    ("Negative electrode density", "parameter set", "1657", "kg/m3", "Chen2020"),
    ("Electrode height x width", "parameter set", "0.065 x 1.58", "m", "Chen2020 -> area 0.1027 m2"),
    ("Voltage window", "parameter set", "2.5 - 4.2", "V", "Chen2020 cut-offs"),
    ("Nominal cell capacity", "parameter set", "5.0", "Ah", "Chen2020"),
]
cap_rows = [
    ("1C discharge capacity (25 C)", "run-pyamm 1C_discharge DFN", "5.03677", "Ah", "cell/r6_t2_1c_dfn.json:capacity_ah"),
    ("-20 C discharge capacity (cold-soak)", "run-pyamm lowT_discharge DFN, T_init 253.15 K, h=80", "4.99987", "Ah", "cell/r6_t2_lowT_dfn.json:capacity_ah"),
    ("-20 C retention", "capacity(lowT)/capacity(1C)", "0.99267", "-", "cell/r6_t2_retention.json:lowT_retention"),
    ("Discharge energy", "time integration of V*I over 1C discharge", "17.9505", "Wh", "cell/r6_t2_energy.json:energy_wh"),
    ("Midpoint voltage", "calc-energy", "3.742", "V", "cell/r6_t2_energy.json:midpoint_voltage_v"),
    ("DC resistance", "calc-energy DCR caliber", "0.004065", "ohm", "cell/r6_t2_energy.json:dcr_ohm"),
]
ed_rows = [
    ("Cell mass (contract)", "sum layer thickness x (1-porosity) x density x area, electrolyte excluded", "0.0380668", "kg", "cell/r6_t2_energy.json:mass_kg"),
    ("Cell volume", "sum layer thickness x area = 188.8e-6 x 0.1027", "1.938976e-5", "m3", "cell/r6_t2_energy.json:volume_m3"),
    ("Energy density", "energy / mass", "471.55", "Wh/kg", "cell/r6_t2_energy.json:energy_density_wh_kg"),
    ("Volumetric energy density", "energy / volume", "925.77", "Wh/L", "cell/r6_t2_energy.json:energy_density_wh_l"),
    ("Power density", "calc-energy", "26903.6", "W/kg", "cell/r6_t2_energy.json:power_density_w_kg"),
]
np_rows = [
    ("N/P (total capacity density basis)", "(eps_am_neg x c_max_neg x L_neg) / (eps_am_pos x c_max_pos x L_pos)", "0.667", "-",
     "(0.75x33133x85.2)/(0.665x63104x75.6); parameter-set values"),
    ("N/P operational balance", "same 5.037 Ah window on both electrodes", "~1.0", "-", "neg 90.1->3.7%, pos 27.0->84.7% (initial stoichiometries, measured window)"),
    ("Positive electrode mass", "layer_kg_m2 x area", "16.842", "g", "cell/r6_t2_energy.json"),
    ("Negative electrode mass", "layer_kg_m2 x area", "10.874", "g", "cell/r6_t2_energy.json"),
    ("Al collector mass", "layer_kg_m2 x area", "2.773", "g", "cell/r6_t2_energy.json"),
    ("Cu collector mass", "layer_kg_m2 x area", "7.362", "g", "cell/r6_t2_energy.json"),
    ("Separator mass", "layer_kg_m2 x area", "0.216", "g", "cell/r6_t2_energy.json"),
    ("Total mass", "sum", "38.067", "g", "cell/r6_t2_energy.json:mass_kg"),
]
proc_rows = [
    ("Areal density positive", "L x (1-eps) x rho", "163.99", "g/m2", "cell/r6_t2_energy.json:layer_kg_m2"),
    ("Areal density negative", "L x (1-eps) x rho", "105.88", "g/m2", "cell/r6_t2_energy.json:layer_kg_m2"),
    ("Compaction density positive", "rho x (1-eps) / 1000", "2.169", "g/cm3", "parameter set (divide by 1000)"),
    ("Compaction density negative", "rho x (1-eps) / 1000", "1.243", "g/cm3", "parameter set (divide by 1000)"),
    ("Electrolyte fill amount", "pore volume x 1.2 g/cm3 (literature default)", "6.325", "g", "pore volume %.3f mL" % pore_ml),
    ("Formation recommendation", "0.1C CC charge to 4.2 V, 25 C, 2 cycles", "-", "-", "design recommended value; production-line value requires tuning"),
]
sheet("1. Input parameters", ["Quantity", "Basis", "Value", "Unit", "Source"], inp)
sheet("2. Capacity and energy", ["Quantity", "Basis", "Value", "Unit", "Source"], cap_rows)
sheet("3. Energy density", ["Quantity", "Basis", "Value", "Unit", "Source"], ed_rows)
sheet("4. NP ratio and mass", ["Quantity", "Basis", "Value", "Unit", "Source"], np_rows)
sheet("5. Process parameters", ["Quantity", "Basis", "Value", "Unit", "Source"], proc_rows)
del wb2["Sheet"]
wb2.save(DEL / "calc.xlsx")

# ================================================================ delivery_index.md
files = [
    ("design_spec.md", "VBF-T4R1-DS-01", "md", "design_spec generated per deliverable-design-spec spec; values from parameter set / cell/r6_t2_*.json"),
    ("design_spec.pdf", "VBF-T4R1-DS-01", "pdf", "PDF release of design_spec.md (reportlab)"),
    ("bom.xlsx", "VBF-T4R1-BOM-01", "xlsx", "bill of materials, dual caliber g/cell + kg/kWh (openpyxl); literature-default mass fractions annotated"),
    ("bom.pdf", "VBF-T4R1-BOM-01", "pdf", "PDF release of bom.xlsx (reportlab)"),
    ("datasheet.md", "VBF-T4R1-DSH-01", "md", "technical datasheet per deliverable-datasheet spec; simulation values only"),
    ("datasheet.pdf", "VBF-T4R1-DSH-01", "pdf", "PDF release of datasheet.md (reportlab)"),
    ("calc.xlsx", "VBF-T4R1-CALC-01", "xlsx", "design calculation sheet, 5 sheets: inputs / capacity-energy / energy density / N-P-mass / process (openpyxl)"),
    ("calc.pdf", "VBF-T4R1-CALC-01", "pdf", "PDF release of calc.xlsx (reportlab)"),
    ("dvpr.md", "VBF-T4R1-DVPR-01", "md", "design verification plan and report, virtual-test version per deliverable-dvpr spec"),
    ("dvpr.pdf", "VBF-T4R1-DVPR-01", "pdf", "PDF release of dvpr.md (reportlab)"),
    ("dfmea.md", "VBF-T4R1-DFMEA-01", "md", "qualitative design FMEA per deliverable-dfmea spec; simulation signals only"),
    ("dfmea.pdf", "VBF-T4R1-DFMEA-01", "pdf", "PDF release of dfmea.md (reportlab)"),
    ("delivery_index.md", "VBF-T4R1-IDX-01", "md", "delivery package index (this file)"),
    ("delivery_index.pdf", "VBF-T4R1-IDX-01", "pdf", "PDF release of delivery_index.md (reportlab, blueprint cover)"),
]
idx_md = []
idx_md.append("# Delivery Package Index — Case t4_r1")
idx_md.append("")
idx_md.append("## Cover information")
idx_md.append("")
idx_md.append("| Field | Value |")
idx_md.append("|---|---|")
idx_md.append("| Case name | t4_r1 (battery for extreme-cold environment equipment) |")
idx_md.append("| Numbering scheme | VBF-T4R1-DOC-NN (case ID uppercased, non-alphanumeric removed) |")
idx_md.append("| Generation date | 2026-08-25 |")
idx_md.append("| Prepared |  |")
idx_md.append("| Reviewed |  |")
idx_md.append("| Approved |  |")
idx_md.append("")
idx_md.append("## Document code reference (fixed by protocol)")
idx_md.append("")
idx_md.append("| Document code | Meaning | Corresponding file |")
idx_md.append("|---|---|---|")
idx_md.append("| DS | Specification | design_spec.md |")
idx_md.append("| BOM | Bill of Materials | bom.xlsx |")
idx_md.append("| DSH | Datasheet | datasheet.md |")
idx_md.append("| CALC | Calculation sheet | calc.xlsx |")
idx_md.append("| DVPR | Design verification report | dvpr.md |")
idx_md.append("| DFMEA | Failure analysis | dfmea.md |")
idx_md.append("| IDX | Delivery index | delivery_index.md |")
idx_md.append("")
idx_md.append("## File list")
idx_md.append("")
idx_md.append("| File name | Number | Format | Source description |")
idx_md.append("|---|---|---|---|")
for fn, num, fmt, src in files:
    idx_md.append("| %s | %s | %s | %s |" % (fn, num, fmt, src))
idx_md.append("")
idx_md.append("Note: CAD structure model not generated (not requested for this headless case; per deliverable-package spec the index registers only actually generated files). Signature fields intentionally left blank for manual signing.")
(DEL / "delivery_index.md").write_text("\n".join(idx_md) + "\n", encoding="utf-8")

# ================================================================ PDF rendering
BLUE_D = colors.HexColor("#14283C")
BLUE_M = colors.HexColor("#1E5A8A")
COPPER = colors.HexColor("#C97B3D")

styles = getSampleStyleSheet()
H1 = ParagraphStyle("H1x", parent=styles["Heading1"], fontSize=16, textColor=BLUE_D, spaceAfter=10)
H2 = ParagraphStyle("H2x", parent=styles["Heading2"], fontSize=12, textColor=BLUE_M, spaceBefore=8, spaceAfter=4)
BODY = ParagraphStyle("Bodyx", parent=styles["BodyText"], fontSize=9, leading=13)
CELLS = ParagraphStyle("Cellx", parent=styles["BodyText"], fontSize=7.5, leading=9.5)

def _clean(s):
    return s.replace("**", "").strip()

def md_flowables(text):
    out = []
    tbl = []
    for line in text.splitlines():
        ls = line.rstrip()
        if not ls.strip():
            if tbl:
                out.append(_table(tbl)); tbl = []
            continue
        if ls.startswith("|"):
            cells = [c.strip() for c in ls.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            tbl.append(cells)
        else:
            if tbl:
                out.append(_table(tbl)); tbl = []
            if ls.startswith("# "):
                out.append(Paragraph(_clean(ls[2:]), H1))
            elif ls.startswith("## "):
                out.append(Paragraph(_clean(ls[3:]), H2))
            else:
                out.append(Paragraph(_clean(ls), BODY))
    if tbl:
        out.append(_table(tbl))
    return out

def _table(rows):
    maxc = max(len(r) for r in rows)
    rows = [r + [""] * (maxc - len(r)) for r in rows]
    widths = []
    for c in range(maxc):
        w = max((len(_clean(r[c])) for r in rows), default=1)
        widths.append(min(max(w * 4.6 + 8, 28), 220))
    tot = sum(widths)
    avail = 523
    widths = [w * avail / tot for w in widths]
    data = [[Paragraph(_clean(cell), CELLS) for cell in r] for r in rows]
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BLUE_M),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA7B4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F7")]),
    ]))
    return t

def md_to_pdf(md_name, pdf_name):
    text = (DEL / md_name).read_text(encoding="utf-8")
    doc = SimpleDocTemplate(str(DEL / pdf_name), pagesize=A4,
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=pdf_name)
    doc.build(md_flowables(text))

for md_name, pdf_name in [("design_spec.md", "design_spec.pdf"),
                          ("datasheet.md", "datasheet.pdf"),
                          ("dvpr.md", "dvpr.pdf"),
                          ("dfmea.md", "dfmea.pdf")]:
    md_to_pdf(md_name, pdf_name)

# xlsx twin PDFs from the same data structures
def twin_pdf(pdf_name, title, header, rows):
    doc = SimpleDocTemplate(str(DEL / pdf_name), pagesize=A4,
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=pdf_name)
    els = [Paragraph(title, H1)]
    data = [[Paragraph(str(h), CELLS) for h in header]]
    for r in rows:
        data.append([Paragraph(_clean(str(x)) if x is not None else "Not provided", CELLS) for x in r])
    ncol = len(header)
    widths = [523 / ncol] * ncol
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BLUE_M),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA7B4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F7")]),
    ]))
    els.append(t)
    doc.build(els)

bom_pdf_rows = []
for name, basis, g, src in bom_rows:
    kgkwh = (g / 1000 / KWH) if g is not None else None
    bom_pdf_rows.append([name, basis,
                         round(g, 3) if g is not None else "Not provided",
                         round(kgkwh, 3) if kgkwh is not None else "Not provided", src])
twin_pdf("bom.pdf", "Bill of Materials — VBF-T4R1-BOM-01",
         ["Component", "Basis / formula", "g/cell", "kg/kWh", "Source"], bom_pdf_rows)

calc_pdf_rows = []
for s_title, s_rows in [("1. Input parameters", inp), ("2. Capacity and energy", cap_rows),
                        ("3. Energy density", ed_rows), ("4. NP ratio and mass", np_rows),
                        ("5. Process parameters", proc_rows)]:
    calc_pdf_rows.append([s_title, "", "", "", ""])
    calc_pdf_rows.extend(s_rows)
twin_pdf("calc.pdf", "Design Calculation Sheet — VBF-T4R1-CALC-01",
         ["Quantity", "Basis", "Value", "Unit", "Source"], calc_pdf_rows)

# delivery_index.pdf with blueprint cover
doc = SimpleDocTemplate(str(DEL / "delivery_index.pdf"), pagesize=A4,
                        leftMargin=14 * mm, rightMargin=14 * mm,
                        topMargin=14 * mm, bottomMargin=14 * mm,
                        title="delivery_index.pdf")
els = []
cover = Table([[Paragraph("VIRTUAL BATTERY FACTORY", ParagraphStyle("cv1", parent=styles["Normal"], fontSize=10, textColor=COPPER))],
               [Paragraph("Design Delivery Package", ParagraphStyle("cv2", parent=styles["Heading1"], fontSize=22, textColor=colors.white))],
               [Paragraph("Case t4_r1 — battery for extreme-cold environment equipment", ParagraphStyle("cv3", parent=styles["Normal"], fontSize=12, textColor=colors.white))],
               [Paragraph("Numbering scheme: VBF-T4R1-&lt;DOC&gt;-&lt;NN&gt;  |  Generated 2026-08-25", ParagraphStyle("cv4", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#B8C6D4")))]],
              colWidths=[523])
cover.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, -1), BLUE_D),
    ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 14),
    ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
]))
els.append(cover)
els.append(Spacer(1, 10))
els.append(Paragraph("Signature block (left blank for manual signing)", H2))
sig = Table([[Paragraph("Prepared", CELLS), Paragraph("Reviewed", CELLS), Paragraph("Approved", CELLS)],
             ["", "", ""]], colWidths=[174, 174, 175])
sig.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA7B4")),
                         ("BACKGROUND", (0, 0), (-1, 0), BLUE_M), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                         ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 30)]))
els.append(sig)
els.append(Spacer(1, 10))
els.append(Paragraph("File list", H2))
data = [[Paragraph("File name", CELLS), Paragraph("Number", CELLS), Paragraph("Format", CELLS), Paragraph("Source description", CELLS)]]
for fn, num, fmt, src in files:
    data.append([Paragraph(fn, CELLS), Paragraph(num, CELLS), Paragraph(fmt, CELLS), Paragraph(src, CELLS)])
t = Table(data, colWidths=[95, 95, 40, 293], repeatRows=1)
t.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), BLUE_M), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA7B4")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F7")]),
]))
els.append(t)
doc.build(els)

print("deliverables generated in", DEL)
