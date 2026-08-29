# Generate bom.xlsx, calc.xlsx and PDF releases for deliverables/.
# Every value is read mechanically from the tool output JSONs (cell/*.json),
# the candidate params JSON, the Chen2020 parameter set, or the diag_r4a_np.py output.
import json
import re
import sys

sys.path.insert(0, r".claude/skills/virtual-battery-factory/scripts")

CASE = "runs/exp/t5_r1_noceiling"
DEL = f"{CASE}/deliverables"
CALC = json.load(open(f"{CASE}/cell/r4_A_calc_dfn.json", encoding="utf-8"))
PARAMS = json.load(open(f"{CASE}/candidates/r4_A_params.json", encoding="utf-8"))

import pybamm

pv = pybamm.ParameterValues("Chen2020")
EPS_P = float(pv["Positive electrode porosity"])
EPS_N = float(pv["Negative electrode porosity"])
EPS_S = float(pv["Separator porosity"])
RHO_P = float(pv["Positive electrode density [kg.m-3]"])
RHO_N = float(pv["Negative electrode density [kg.m-3]"])
RHO_S = float(pv["Separator density [kg.m-3]"])
RHO_AL = float(pv["Positive current collector density [kg.m-3]"])
RHO_CU = float(pv["Negative current collector density [kg.m-3]"])
C_MAX_P = float(pv["Maximum concentration in positive electrode [mol.m-3]"])
C_MAX_N = float(pv["Maximum concentration in negative electrode [mol.m-3]"])

AREA = float(CALC["area_m2"])
L_P = PARAMS["Positive electrode thickness [m]"]
L_N = PARAMS["Negative electrode thickness [m]"]
L_S = PARAMS["Separator thickness [m]"]
LCC_P = PARAMS["Positive current collector thickness [m]"]
LCC_N = PARAMS["Negative current collector thickness [m]"]
ENERGY_WH = float(CALC["energy_wh"])
ENERGY_KWH = ENERGY_WH / 1000.0
MASS_KG = float(CALC["mass_kg"])

# per-layer masses (g) = layer_kg_m2 x area — mechanical from calc-energy output
LK = CALC["layer_kg_m2"]
m_pos = LK["positive_electrode"] * AREA * 1000.0
m_neg = LK["negative_electrode"] * AREA * 1000.0
m_al = LK["positive_cc"] * AREA * 1000.0
m_cu = LK["negative_cc"] * AREA * 1000.0
m_sep = LK["separator"] * AREA * 1000.0

# additive/binder literature defaults: 2 wt% CB + 2 wt% binder of coating solids (estimate)
m_cb_p = m_pos * 2.0 / 96.0
m_bd_p = m_pos * 2.0 / 96.0
m_cb_n = m_neg * 2.0 / 96.0
m_bd_n = m_neg * 2.0 / 96.0

# electrolyte: pore volume x literature density 1.2 g/cm3 (estimate, fill factor 1.0)
pore_vol_m3 = (EPS_P * L_P + EPS_N * L_N + EPS_S * L_S) * AREA
m_elyte = pore_vol_m3 * 1200.0 * 1000.0  # g

industry_mass = m_pos + m_neg + m_al + m_cu + m_sep + m_cb_p + m_bd_p + m_cb_n + m_bd_n + m_elyte

def kg_kwh(g):
    return g / 1000.0 / ENERGY_KWH

BOM_ROWS = [
    ("Positive electrode active material (NMC811)", m_pos, "active volume fraction 0.665, density 3262 kg/m3 (parameter set); mass = layer_kg_m2 x area (cell/r4_A_calc_dfn.json)"),
    ("Negative electrode active material (graphite)", m_neg, "active volume fraction 0.75, density 1657 kg/m3 (parameter set); mass = layer_kg_m2 x area"),
    ("Positive conductive additive (carbon black)", m_cb_p, "ESTIMATE: literature default 2 wt% of positive coating solids (not in parameter set; excluded from simulation mass)"),
    ("Positive binder (PVDF)", m_bd_p, "ESTIMATE: literature default 2 wt% of positive coating solids (not in parameter set; excluded from simulation mass)"),
    ("Negative conductive additive (carbon black)", m_cb_n, "ESTIMATE: literature default 2 wt% of negative coating solids (not in parameter set; excluded from simulation mass)"),
    ("Negative binder (CMC/SBR)", m_bd_n, "ESTIMATE: literature default 2 wt% of negative coating solids (not in parameter set; excluded from simulation mass)"),
    ("Separator", m_sep, "porosity 0.47, density 397 kg/m3 (parameter set); mass = layer_kg_m2 x area"),
    ("Electrolyte (1 M LiPF6 in EC:EMC 3:7 w/w)", m_elyte, "ESTIMATE: pore volume 9.64 mL x literature density 1.2 g/cm3, fill factor 1.0 (not in parameter set; excluded from simulation mass)"),
    ("Positive current collector (Al, 6 um)", m_al, "density 2700 kg/m3 (parameter set); mass = layer_kg_m2 x area"),
    ("Negative current collector (Cu, 4 um)", m_cu, "density 8960 kg/m3 (parameter set); mass = layer_kg_m2 x area"),
    ("Enclosure (can/pouch)", None, "Not modeled (beyond pure simulation boundary)"),
    ("Tabs", None, "Not modeled (beyond pure simulation boundary)"),
]

# ---------------- xlsx generation ----------------
def write_xlsx():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    HDR_FILL = PatternFill("solid", fgColor="1E5A8A")
    HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
    BODY_FONT = Font(size=10)
    SUB_FILL = PatternFill("solid", fgColor="DCE6F1")

    def style_header(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    # --- bom.xlsx ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["Component", "Mass (g/cell)", "kg/kWh", "Volume fraction", "Density (kg/m3)", "Source / note"])
    fracs = {
        "Positive electrode active material (NMC811)": "0.665 (active fraction)", "Negative electrode active material (graphite)": "0.75 (active fraction)",
        "Separator": "0.53 (solid fraction)", "Positive current collector (Al, 6 um)": "1.0", "Negative current collector (Cu, 4 um)": "1.0",
    }
    dens = {
        "Positive electrode active material (NMC811)": 3262.0, "Negative electrode active material (graphite)": 1657.0,
        "Separator": 397.0, "Positive current collector (Al, 6 um)": 2700.0, "Negative current collector (Cu, 4 um)": 8960.0,
    }
    for name, mass, note in BOM_ROWS:
        ws.append([
            name,
            None if mass is None else round(mass, 3),
            None if mass is None else round(kg_kwh(mass), 4),
            fracs.get(name, "-"),
            dens.get(name, "-"),
            note,
        ])
    ws.append(["SUMMARY: total mass, simulation contract caliber (electrolyte excluded)", round(MASS_KG * 1000.0, 3), round(MASS_KG / ENERGY_KWH, 4), "-", "-", "cell/r4_A_calc_dfn.json:mass_kg; formula = sum layer thickness x (1-porosity) x density x area"])
    ws.append(["SUMMARY: total mass, industry caliber incl. electrolyte + additive/binder estimates", round(industry_mass, 3), round(industry_mass / 1000.0 / ENERGY_KWH, 4), "-", "-", "estimate (electrolyte + additive/binder literature defaults added to contract mass)"])
    ws.append(["SUMMARY: cell energy", f"{ENERGY_WH:.3f} Wh = {ENERGY_KWH:.6f} kWh", "-", "-", "-", "cell/r4_A_calc_dfn.json:energy_wh (time integration of V x I at 5 A)"])
    style_header(ws, 6)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 95
    wb.save(f"{DEL}/bom.xlsx")
    print("wrote bom.xlsx")

    # --- calc.xlsx ---
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Inputs"
    ws.append(["Parameter", "Value", "Unit", "Source"])
    inputs = [
        ("Positive electrode thickness", L_P * 1e6, "um", "candidates/r4_A_params.json"),
        ("Negative electrode thickness", L_N * 1e6, "um", "candidates/r4_A_params.json"),
        ("Separator thickness", L_S * 1e6, "um", "candidates/r4_A_params.json"),
        ("Positive current collector thickness (Al)", LCC_P * 1e6, "um", "candidates/r4_A_params.json"),
        ("Negative current collector thickness (Cu)", LCC_N * 1e6, "um", "candidates/r4_A_params.json"),
        ("Electrode width", PARAMS["Electrode width [m]"], "m", "candidates/r4_A_params.json"),
        ("Electrode height", 0.065, "m", "parameter set (Chen2020)"),
        ("Electrode area", AREA, "m2", "cell/r4_A_calc_dfn.json:area_m2"),
        ("Positive particle radius", PARAMS["Positive particle radius [m]"] * 1e6, "um", "candidates/r4_A_params.json"),
        ("Negative particle radius", PARAMS["Negative particle radius [m]"] * 1e6, "um", "candidates/r4_A_params.json"),
        ("Electrolyte conductivity", PARAMS["Electrolyte conductivity [S.m-1]"], "S/m", "candidates/r4_A_params.json (literature-advanced estimate)"),
        ("Cation transference number", PARAMS["Cation transference number"], "-", "candidates/r4_A_params.json"),
        ("Electrolyte diffusivity", PARAMS["Electrolyte diffusivity [m2.s-1]"], "m2/s", "candidates/r4_A_params.json (literature-advanced estimate)"),
        ("Total heat transfer coefficient", PARAMS["Total heat transfer coefficient [W.m-2.K-1]"], "W/m2/K", "candidates/r4_A_params.json"),
        ("Positive electrode porosity", EPS_P, "-", "parameter set (Chen2020)"),
        ("Negative electrode porosity", EPS_N, "-", "parameter set (Chen2020)"),
        ("Separator porosity", EPS_S, "-", "parameter set (Chen2020)"),
        ("Positive electrode density", RHO_P, "kg/m3", "parameter set (Chen2020)"),
        ("Negative electrode density", RHO_N, "kg/m3", "parameter set (Chen2020)"),
        ("Separator density", RHO_S, "kg/m3", "parameter set (Chen2020)"),
        ("Positive current collector density (Al)", RHO_AL, "kg/m3", "parameter set (Chen2020)"),
        ("Negative current collector density (Cu)", RHO_CU, "kg/m3", "parameter set (Chen2020)"),
        ("Maximum concentration, positive", C_MAX_P, "mol/m3", "parameter set (Chen2020)"),
        ("Maximum concentration, negative", C_MAX_N, "mol/m3", "parameter set (Chen2020)"),
        ("Lower / upper voltage cut-off", "2.5 / 4.2", "V", "parameter set (Chen2020)"),
        ("1C current (nominal)", 5.0, "A", "protocol definition (nominal 5 Ah)"),
        ("4C charge current", 20.0, "A", "protocol 4C_charge_45C (C_rate 4.0)"),
    ]
    for row in inputs:
        ws.append(list(row))
    style_header(ws, 4)
    for col, w in zip("ABCD", (52, 16, 12, 60)):
        ws.column_dimensions[col].width = w
    # number format for value column
    for r in range(2, len(inputs) + 2):
        cell = ws.cell(row=r, column=2)
        if isinstance(cell.value, float):
            cell.number_format = "0.000E+00" if abs(cell.value) < 1e-3 else "0.0000"

    ws = wb.create_sheet("CapacityEnergy")
    ws.append(["Quantity", "Value", "Formula", "Source"])
    cap_rows = [
        ("Capacity (1C to 2.5 V)", CALC["capacity_ah"], "time integral of 5 A until 2.5 V (calc-energy reads 1C discharge sim)", "cell/r4_A_calc_dfn.json:capacity_ah"),
        ("Discharge energy", CALC["energy_wh"], "E = 5 A x integral(V dt) to 2.5 V", "cell/r4_A_calc_dfn.json:energy_wh"),
        ("Midpoint voltage", CALC["midpoint_voltage_v"], "V at half discharged energy", "cell/r4_A_calc_dfn.json:midpoint_voltage_v"),
        ("DC resistance", CALC["dcr_ohm"], "quasi-static DCR from calc-energy", "cell/r4_A_calc_dfn.json:dcr_ohm"),
        ("Power density", CALC["power_density_w_kg"], "from DCR and cell mass", "cell/r4_A_calc_dfn.json:power_density_w_kg"),
        ("4C charge accepted (DFN)", 8.12, "span from argmin-V (charge start) to t_end x 20 A / 3600 = 1461.0 s", "cell/r4_A_4c_dfn.json (span 1461.0 s, V_end 4.2000)"),
    ]
    for row in cap_rows:
        ws.append(list(row))
    style_header(ws, 4)
    for col, w in zip("ABCD", (34, 22, 60, 60)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("EnergyDensity")
    ws.append(["Quantity", "Value", "Formula", "Source"])
    ed_rows = [
        ("Cell mass (contract caliber)", MASS_KG, "sum over layers of thickness x (1-porosity) x density x area; electrolyte excluded", "cell/r4_A_calc_dfn.json:mass_kg"),
        ("Gravimetric energy density", CALC["energy_density_wh_kg"], "ED = energy Wh / mass kg", "cell/r4_A_calc_dfn.json:energy_density_wh_kg"),
        ("Volumetric energy density", CALC["energy_density_wh_l"], "EDvol = energy Wh / stack volume (0.2208 mm x 0.15405 m2)", "cell/r4_A_calc_dfn.json:energy_density_wh_l"),
        ("Contract threshold", 500.94, ">= 500.94 Wh/kg (task text, entry 0)", "log.jsonl entry 0"),
        ("Positive electrode layer mass", m_pos, "layer_kg_m2 0.206727619 kg/m2 x area 0.15405 m2", "cell/r4_A_calc_dfn.json:layer_kg_m2"),
        ("Negative electrode layer mass", m_neg, "layer_kg_m2 0.133595625 kg/m2 x area", "cell/r4_A_calc_dfn.json:layer_kg_m2"),
        ("Al current collector mass", m_al, "layer_kg_m2 0.0162 kg/m2 x area", "cell/r4_A_calc_dfn.json:layer_kg_m2"),
        ("Cu current collector mass", m_cu, "layer_kg_m2 0.03584 kg/m2 x area", "cell/r4_A_calc_dfn.json:layer_kg_m2"),
        ("Separator mass", m_sep, "layer_kg_m2 0.00168328 kg/m2 x area", "cell/r4_A_calc_dfn.json:layer_kg_m2"),
    ]
    for row in ed_rows:
        ws.append(list(row))
    style_header(ws, 4)
    for col, w in zip("ABCD", (36, 22, 66, 46)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("NPandMass")
    ws.append(["Quantity", "Value", "Formula", "Source"])
    np_rows = [
        ("x_n at charge start (2.5 V)", 0.0283, "X-averaged negative particle concentration / c_max_n (DFN 4C experiment)", "diag_r4a_np.py"),
        ("x_n at charge end (4.2 V)", 0.7622, "same", "diag_r4a_np.py"),
        ("x_p at charge start (2.5 V)", 0.8485, "X-averaged positive particle concentration / c_max_p", "diag_r4a_np.py"),
        ("x_p at charge end (4.2 V)", 0.3805, "same", "diag_r4a_np.py"),
        ("Negative areal capacity", 52.546, "Q_n = c_max_n x (x_n_end - x_n_start) x (1 - eps_n) x L_n x F / 3600 [Ah/m2]", "diag_r4a_np.py"),
        ("Positive areal capacity", 50.161, "Q_p = c_max_p x (x_p_start - x_p_end) x (1 - eps_p) x L_p x F / 3600 [Ah/m2]", "diag_r4a_np.py"),
        ("N/P ratio", 1.0475, "N/P = Q_n / Q_p (2.5-4.2 V charge window)", "diag_r4a_np.py"),
        ("Stack thickness", CALC["thickness_m"] * 1e6, "L_p + L_n + L_s + L_cc_p + L_cc_n = 95.3 + 107.5 + 8 + 6 + 4 um", "cell/r4_A_calc_dfn.json:thickness_m"),
    ]
    for row in np_rows:
        ws.append(list(row))
    style_header(ws, 4)
    for col, w in zip("ABCD", (36, 22, 72, 40)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("ProcessParams")
    ws.append(["Parameter", "Value", "Formula", "Source"])
    proc_rows = [
        ("Positive areal density", 206.73, "L_p x (1 - eps_p) x rho_p = 95.3 um x 0.665 x 3262 kg/m3", "cell/r4_A_calc_dfn.json:layer_kg_m2.positive_electrode"),
        ("Negative areal density", 133.60, "L_n x (1 - eps_n) x rho_n = 107.5 um x 0.75 x 1657 kg/m3", "cell/r4_A_calc_dfn.json:layer_kg_m2.negative_electrode"),
        ("Positive compaction density", 2.169, "rho_p x (1 - eps_p) / 1000 = 3262 x 0.665 / 1000 [g/cm3]", "parameter set"),
        ("Negative compaction density", 1.243, "rho_n x (1 - eps_n) / 1000 = 1657 x 0.75 / 1000 [g/cm3]", "parameter set"),
        ("Electrolyte fill amount (volume)", round(pore_vol_m3 * 1e6, 2), "pore volume = (eps_p L_p + eps_n L_n + eps_s L_s) x area [mL]", "mechanical derivation"),
        ("Electrolyte fill amount (mass)", round(m_elyte, 3), "pore volume x 1.2 g/cm3 (literature density) x fill factor 1.0", "literature density estimate"),
        ("Formation recommendation", "0.1C CC to 4.2 V at 25 degC, 2 cycles", "design recommended value; production-line value requires tuning", "design recommendation (estimate)"),
    ]
    for row in proc_rows:
        ws.append(list(row))
    style_header(ws, 4)
    for col, w in zip("ABCD", (38, 20, 68, 52)):
        ws.column_dimensions[col].width = w

    wb.save(f"{DEL}/calc.xlsx")
    print("wrote calc.xlsx")


# ---------------- PDF generation (reportlab) ----------------
def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def md_blocks(text):
    lines = text.splitlines()
    blocks, table = [], []
    for ln in lines:
        s = ln.strip()
        if s.startswith("|"):
            table.append(s)
        else:
            if table:
                blocks.append(("table", table))
                table = []
            if s.startswith("# "):
                blocks.append(("h1", s[2:]))
            elif s.startswith("## "):
                blocks.append(("h2", s[3:]))
            elif s.startswith("### "):
                blocks.append(("h3", s[4:]))
            elif s.startswith("- "):
                blocks.append(("bullet", s[2:]))
            elif s.startswith("> "):
                blocks.append(("quote", s[2:]))
            elif s:
                blocks.append(("para", s))
    if table:
        blocks.append(("table", table))
    return blocks


def parse_table(rows):
    data, is_sep = [], True
    for r in rows:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
            continue  # separator row
        data.append(cells)
    return data


def render_md_pdf(md_path, out_path, title_color="#14283C"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    text = open(md_path, encoding="utf-8").read()
    blocks = md_blocks(text)
    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=text.splitlines()[0].lstrip("# ").strip())
    st_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.HexColor(title_color), spaceAfter=8)
    st_h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#1E5A8A"), spaceBefore=8, spaceAfter=4)
    st_h3 = ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor("#14283C"), spaceBefore=6, spaceAfter=3)
    st_p = ParagraphStyle("p", fontName="Helvetica", fontSize=9.5, leading=12.5, spaceAfter=4)
    st_b = ParagraphStyle("b", fontName="Helvetica", fontSize=9.5, leading=12.5, leftIndent=10, bulletIndent=2, spaceAfter=2)
    st_q = ParagraphStyle("q", fontName="Helvetica-Oblique", fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"), spaceAfter=4)
    st_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=9.6)
    st_cellb = ParagraphStyle("cb", fontName="Helvetica-Bold", fontSize=8, leading=9.6, textColor=colors.white)

    story = []
    for kind, payload in blocks:
        if kind == "h1":
            story.append(Paragraph(_esc(payload), st_title))
            story.append(Spacer(1, 2))
        elif kind == "h2":
            story.append(Paragraph(_esc(payload), st_h2))
        elif kind == "h3":
            story.append(Paragraph(_esc(payload), st_h3))
        elif kind == "para":
            story.append(Paragraph(_esc(payload), st_p))
        elif kind == "bullet":
            story.append(Paragraph(_esc(payload), st_b, bulletText="•"))
        elif kind == "quote":
            story.append(Paragraph(_esc(payload), st_q))
        elif kind == "table":
            data = parse_table(payload)
            if not data:
                continue
            avail = 182 * mm
            ncol = max(len(r) for r in data)
            widths = [avail / ncol] * ncol
            tdata = []
            for i, row in enumerate(data):
                st = st_cellb if i == 0 else st_cell
                tdata.append([Paragraph(_esc(c), st) for c in row])
            tbl = Table(tdata, colWidths=widths, repeatRows=1)
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E5A8A")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C6D6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F8")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]
            tbl.setStyle(TableStyle(style))
            story.append(tbl)
            story.append(Spacer(1, 5))
    doc.build(story)
    print("wrote", out_path)


def xlsx_to_pdf(sheet_rows, out_path, headers, title):
    """sheet_rows: list of (sheet_title, rows: list of tuples)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=13 * mm, bottomMargin=13 * mm, title=title)
    st_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=14, leading=17, textColor=colors.HexColor("#14283C"), spaceAfter=2)
    st_sub = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=11, textColor=colors.HexColor("#1E5A8A"), spaceAfter=6)
    st_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=7.5, leading=9)
    st_cellb = ParagraphStyle("cb", fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white)
    st_sheet = ParagraphStyle("sh", fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor("#C97B3D"), spaceBefore=8, spaceAfter=3)

    story = [Paragraph(_esc(title), st_title), Paragraph("PDF release of the xlsx deliverable; values mechanically taken from tool outputs (cell/*.json, parameter set).", st_sub)]
    avail = 186 * mm
    for si, (sheet_title, rows) in enumerate(sheet_rows):
        if si > 0:
            story.append(PageBreak())
        story.append(Paragraph(_esc(f"Sheet: {sheet_title}"), st_sheet))
        data = [list(headers)] + [list(r) for r in rows]
        ncol = len(headers)
        widths = [avail / ncol] * ncol
        tdata = []
        for i, row in enumerate(data):
            st = st_cellb if i == 0 else st_cell
            tdata.append([Paragraph(_esc("" if c is None else str(c)), st) for c in row])
        tbl = Table(tdata, colWidths=widths, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E5A8A")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C6D6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F8")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        story.append(tbl)
    doc.build(story)
    print("wrote", out_path)


def render_index_pdf(md_path, out_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    text = open(md_path, encoding="utf-8").read()
    blocks = md_blocks(text)
    # cover page: engineering blueprint style
    c = canvas.Canvas(out_path, pagesize=A4)
    W, H = A4
    c.setFillColor(colors.HexColor("#14283C"))
    c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColor(colors.HexColor("#C97B3D"))
    c.rect(20 * mm, H - 46 * mm, 40 * mm, 1.6 * mm, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(20 * mm, H - 38 * mm, "Virtual Battery Factory - Design Package")
    c.setFont("Helvetica", 11)
    c.setFillColor(colors.HexColor("#9FB6CC"))
    c.drawString(20 * mm, H - 47.5 * mm, "DELIVERY INDEX")
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(20 * mm, H - 60 * mm, "Case: t5_r1_noceiling")
    c.setFont("Helvetica", 10.5)
    c.drawString(20 * mm, H - 67 * mm, "Next-generation flagship vehicle battery")
    c.drawString(20 * mm, H - 74 * mm, "Energy density >= 500.94 Wh/kg | 4C fast charge, no lithium plating | T_max <= 60 degC")
    c.drawString(20 * mm, H - 81 * mm, "Numbering scheme: VBF-<CASE-ID>-<DOC-CODE>-<SEQ>  (case ID: T5R1NOCEILING)")
    c.setFillColor(colors.HexColor("#9FB6CC"))
    c.setFont("Helvetica", 9)
    c.drawString(20 * mm, H - 96 * mm, "Generation date: 2026-08-25")
    c.setFillColor(colors.white)
    c.setFont("Helvetica", 9)
    c.drawString(20 * mm, H - 116 * mm, "Prepared: ____________________")
    c.drawString(20 * mm, H - 124 * mm, "Reviewed: ____________________")
    c.drawString(20 * mm, H - 132 * mm, "Approved: ____________________")
    c.setFillColor(colors.HexColor("#1E5A8A"))
    c.rect(0, 12 * mm, W, 8 * mm, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Helvetica", 8)
    c.drawString(20 * mm, 15 * mm, "Virtual battery design package - simulation-caliber deliverables (virtual test). Source files and PDF releases.")
    c.showPage()

    # content pages: headings/paragraphs + the file list table
    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title="Delivery Index")
    st_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.HexColor("#14283C"), spaceAfter=8)
    st_h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#1E5A8A"), spaceBefore=8, spaceAfter=4)
    st_p = ParagraphStyle("p", fontName="Helvetica", fontSize=9.5, leading=12.5, spaceAfter=4)
    st_b = ParagraphStyle("b", fontName="Helvetica", fontSize=9.5, leading=12.5, leftIndent=10, bulletIndent=2, spaceAfter=2)
    st_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=9.6)
    st_cellb = ParagraphStyle("cb", fontName="Helvetica-Bold", fontSize=8, leading=9.6, textColor=colors.white)

    story = []
    for kind, payload in blocks:
        if kind == "h1":
            story.append(Paragraph(_esc(payload), st_title))
        elif kind == "h2":
            story.append(Paragraph(_esc(payload), st_h2))
        elif kind == "para":
            story.append(Paragraph(_esc(payload), st_p))
        elif kind == "bullet":
            story.append(Paragraph(_esc(payload), st_b, bulletText="•"))
        elif kind == "table":
            data = parse_table(payload)
            if not data:
                continue
            avail = 182 * mm
            ncol = max(len(r) for r in data)
            widths = [avail / ncol] * ncol
            tdata = []
            for i, row in enumerate(data):
                st = st_cellb if i == 0 else st_cell
                tdata.append([Paragraph(_esc(c), st) for c in row])
            tbl = Table(tdata, colWidths=widths, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E5A8A")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C6D6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF3F8")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 5))
    doc.build(story, onFirstPage=lambda d, doc: None)
    print("wrote", out_path)


def xlsx_data_for_pdf():
    bom_rows = []
    for name, mass, note in BOM_ROWS:
        bom_rows.append([
            name,
            "" if mass is None else f"{mass:.3f}",
            "" if mass is None else f"{kg_kwh(mass):.4f}",
            note,
        ])
    bom_rows.append([
        "SUMMARY: total mass, simulation contract caliber (electrolyte excluded)",
        f"{MASS_KG * 1000.0:.3f}", f"{MASS_KG / ENERGY_KWH:.4f}",
        "cell/r4_A_calc_dfn.json:mass_kg; formula = sum layer thickness x (1-porosity) x density x area",
    ])
    bom_rows.append([
        "SUMMARY: total mass, industry caliber incl. electrolyte + additive/binder estimates",
        f"{industry_mass:.3f}", f"{industry_mass / 1000.0 / ENERGY_KWH:.4f}",
        "estimate (electrolyte + additive/binder literature defaults added to contract mass)",
    ])
    bom_rows.append([
        "SUMMARY: cell energy", f"{ENERGY_WH:.3f} Wh = {ENERGY_KWH:.6f} kWh", "",
        "cell/r4_A_calc_dfn.json:energy_wh (time integration of V x I at 5 A)",
    ])
    calc_rows = {
        "Inputs": [
            ["Positive electrode thickness", f"{L_P * 1e6:.1f} um", "candidates/r4_A_params.json"],
            ["Negative electrode thickness", f"{L_N * 1e6:.1f} um", "candidates/r4_A_params.json"],
            ["Separator thickness", f"{L_S * 1e6:.1f} um", "candidates/r4_A_params.json"],
            ["Positive current collector thickness (Al)", f"{LCC_P * 1e6:.1f} um", "candidates/r4_A_params.json"],
            ["Negative current collector thickness (Cu)", f"{LCC_N * 1e6:.1f} um", "candidates/r4_A_params.json"],
            ["Electrode width", f"{PARAMS['Electrode width [m]']:.2f} m", "candidates/r4_A_params.json"],
            ["Electrode height", "0.065 m", "parameter set (Chen2020)"],
            ["Electrode area", f"{AREA:.5f} m2", "cell/r4_A_calc_dfn.json:area_m2"],
            ["Positive / negative particle radius", f"{PARAMS['Positive particle radius [m]'] * 1e6:.1f} / {PARAMS['Negative particle radius [m]'] * 1e6:.1f} um", "candidates/r4_A_params.json"],
            ["Electrolyte conductivity", f"{PARAMS['Electrolyte conductivity [S.m-1]']:.1f} S/m", "candidates/r4_A_params.json (literature-advanced estimate)"],
            ["Cation transference number", f"{PARAMS['Cation transference number']:.2f}", "candidates/r4_A_params.json"],
            ["Electrolyte diffusivity", f"{PARAMS['Electrolyte diffusivity [m2.s-1]']:.2e} m2/s", "candidates/r4_A_params.json (literature-advanced estimate)"],
            ["Total heat transfer coefficient", f"{PARAMS['Total heat transfer coefficient [W.m-2.K-1]']:.0f} W/m2/K", "candidates/r4_A_params.json"],
            ["Positive / negative porosity", f"{EPS_P} / {EPS_N}", "parameter set (Chen2020)"],
            ["Positive / negative density", f"{RHO_P} / {RHO_N} kg/m3", "parameter set (Chen2020)"],
            ["Current collector densities (Al/Cu)", f"{RHO_AL} / {RHO_CU} kg/m3", "parameter set (Chen2020)"],
            ["Voltage cut-offs", "2.5 / 4.2 V", "parameter set (Chen2020)"],
            ["1C / 4C charge current", "5 / 20 A", "protocol definitions"],
        ],
        "CapacityEnergy": [
            ["Capacity (1C to 2.5 V)", f"{CALC['capacity_ah']:.4f} Ah", "cell/r4_A_calc_dfn.json:capacity_ah"],
            ["Discharge energy", f"{CALC['energy_wh']:.3f} Wh", "cell/r4_A_calc_dfn.json:energy_wh"],
            ["Midpoint voltage", f"{CALC['midpoint_voltage_v']:.4f} V", "cell/r4_A_calc_dfn.json:midpoint_voltage_v"],
            ["DC resistance", f"{CALC['dcr_ohm']:.6f} Ohm", "cell/r4_A_calc_dfn.json:dcr_ohm"],
            ["Power density", f"{CALC['power_density_w_kg']:.1f} W/kg", "cell/r4_A_calc_dfn.json:power_density_w_kg"],
            ["4C charge accepted (DFN)", "8.12 Ah in 1461.0 s to 4.2 V", "cell/r4_A_4c_dfn.json span x 20 A / 3600"],
        ],
        "EnergyDensity": [
            ["Cell mass (contract caliber)", f"{MASS_KG:.5f} kg", "cell/r4_A_calc_dfn.json:mass_kg"],
            ["Gravimetric energy density", f"{CALC['energy_density_wh_kg']:.2f} Wh/kg", "cell/r4_A_calc_dfn.json:energy_density_wh_kg"],
            ["Volumetric energy density", f"{CALC['energy_density_wh_l']:.2f} Wh/L", "cell/r4_A_calc_dfn.json:energy_density_wh_l"],
            ["Contract threshold", ">= 500.94 Wh/kg", "log.jsonl entry 0 (task text verbatim)"],
            ["Positive / negative layer mass", f"{m_pos:.3f} / {m_neg:.3f} g", "calc layer_kg_m2 x area"],
            ["Al / Cu collector mass", f"{m_al:.3f} / {m_cu:.3f} g", "calc layer_kg_m2 x area"],
            ["Separator mass", f"{m_sep:.3f} g", "calc layer_kg_m2 x area"],
        ],
        "NPandMass": [
            ["x_n charge window (2.5 -> 4.2 V)", "0.0283 -> 0.7622", "diag_r4a_np.py (DFN 4C experiment)"],
            ["x_p charge window (2.5 -> 4.2 V)", "0.8485 -> 0.3805", "diag_r4a_np.py"],
            ["Negative areal capacity", "52.546 Ah/m2", "diag_r4a_np.py (Q_n formula)"],
            ["Positive areal capacity", "50.161 Ah/m2", "diag_r4a_np.py (Q_p formula)"],
            ["N/P ratio", "1.0475", "diag_r4a_np.py (Q_n / Q_p)"],
            ["Stack thickness", f"{CALC['thickness_m'] * 1e6:.1f} um", "cell/r4_A_calc_dfn.json:thickness_m"],
        ],
        "ProcessParams": [
            ["Positive / negative areal density", "206.73 / 133.60 g/m2", "calc layer_kg_m2"],
            ["Positive / negative compaction density", "2.169 / 1.243 g/cm3", "parameter set: rho x (1-eps) / 1000"],
            ["Electrolyte fill (volume / mass)", f"{pore_vol_m3 * 1e6:.2f} mL / {m_elyte:.3f} g", "pore volume x 1.2 g/cm3 (literature density), fill factor 1.0"],
            ["Formation recommendation", "0.1C CC to 4.2 V at 25 degC, 2 cycles", "design recommended value (estimate)"],
        ],
    }
    return (
        [("BOM", bom_rows)],
        [("Inputs", calc_rows["Inputs"]), ("CapacityEnergy", calc_rows["CapacityEnergy"]),
         ("EnergyDensity", calc_rows["EnergyDensity"]), ("NPandMass", calc_rows["NPandMass"]),
         ("ProcessParams", calc_rows["ProcessParams"])],
    )


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "all"
    if mode in ("xlsx", "all"):
        write_xlsx()
    if mode in ("pdf", "all"):
        bom_sheets, calc_sheets = xlsx_data_for_pdf()
        xlsx_to_pdf(bom_sheets, f"{DEL}/bom.pdf", ["Component", "Mass (g/cell)", "kg/kWh", "Source / note"], "Bill of Materials - VBF-T5R1NOCEILING-BOM-01")
        xlsx_to_pdf(calc_sheets, f"{DEL}/calc.pdf", ["Quantity", "Value", "Source"], "Design Calculation Sheet - VBF-T5R1NOCEILING-CALC-01")
        render_md_pdf(f"{DEL}/design_spec.md", f"{DEL}/design_spec.pdf")
        render_md_pdf(f"{DEL}/datasheet.md", f"{DEL}/datasheet.pdf")
        render_md_pdf(f"{DEL}/dvpr.md", f"{DEL}/dvpr.pdf")
        render_md_pdf(f"{DEL}/dfmea.md", f"{DEL}/dfmea.pdf")
        render_index_pdf(f"{DEL}/delivery_index.md", f"{DEL}/delivery_index.pdf")
