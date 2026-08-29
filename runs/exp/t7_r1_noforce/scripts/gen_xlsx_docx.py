"""Generate bom.xlsx, calc.xlsx, datasheet.docx from tool outputs + Chen2020 dump values."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from docx import Document
from docx.shared import Pt

ws = Path(r"D:/research/degradation_prognostics/Battery_Design_Agent/runs/exp/t7_r1_noforce")
deliv = ws / "deliverables"
deliv.mkdir(exist_ok=True)

energy = json.loads((ws / "cell/r5_slim_energy.json").read_text(encoding="utf-8-sig"))
AREA = energy["area_m2"]  # 0.1027
E_WH = energy["energy_wh"]  # 17.8042
E_KWH = E_WH / 1000.0
TOT_G = energy["mass_kg"] * 1000.0  # 36.2264 g (electrolyte excluded)
layer_g = {k: v * AREA * 1000.0 for k, v in energy["layer_kg_m2"].items()}

# Literature weight-fraction splits (annotated defaults; Chen2020 has no additive/binder parameters)
pos_coat = layer_g["positive_electrode"]
neg_coat = layer_g["negative_electrode"]
pos_am, pos_cb, pos_binder = pos_coat * 0.96, pos_coat * 0.02, pos_coat * 0.02
neg_am, neg_cb, neg_binder = neg_coat * 0.958, neg_coat * 0.012, neg_coat * 0.03
# Electrolyte: pore volume x 1.2 g/cm3 (literature density, annotated)
pore_m3 = AREA * (75.6e-6 * 0.335 + 85.2e-6 * 0.25 + 10e-6 * 0.47)
elec_g = pore_m3 * 1.2e6
al_g = layer_g["positive_cc"]
cu_g = layer_g["negative_cc"]
sep_g = layer_g["separator"]
total_g = TOT_G + elec_g

bom_rows = [
    ("Positive electrode active material (NMC811)", pos_am, "coating mass x 0.96 wt (literature default 96/2/2 AM/CB/PVDF; parameter set has no additive/binder fractions)", "calc-energy layer_kg_m2.positive_electrode"),
    ("Positive electrode conductive additive (carbon black)", pos_cb, "coating mass x 0.02 wt (literature default)", "同上"),
    ("Positive electrode binder (PVDF)", pos_binder, "coating mass x 0.02 wt (literature default)", "同上"),
    ("Negative electrode active material (graphite)", neg_am, "coating mass x 0.958 wt (literature default 95.8/1.2/3.0 AM/CB/CMC-SBR)", "calc-energy layer_kg_m2.negative_electrode"),
    ("Negative electrode conductive additive (carbon black)", neg_cb, "coating mass x 0.012 wt (literature default)", "同上"),
    ("Negative electrode binder (CMC/SBR)", neg_binder, "coating mass x 0.03 wt (literature default)", "同上"),
    ("Separator (polyolefin)", sep_g, "thickness x area x (1-porosity) x density (tool contract caliber; Chen2020 rho=397 kg/m3, eps=0.47)", "calc-energy layer_kg_m2.separator"),
    ("Electrolyte (LiPF6/carbonate)", elec_g, "pore volume x 1.2 g/cm3 (literature density, annotated; not in contract mass caliber)", "pore volume from parameter-set geometry"),
    ("Positive current collector (Al, 10 um)", al_g, "thickness x area x density (2700 kg/m3)", "calc-energy layer_kg_m2.positive_cc"),
    ("Negative current collector (Cu, 6 um)", cu_g, "thickness x area x density (8960 kg/m3)", "calc-energy layer_kg_m2.negative_cc"),
    ("Enclosure / tabs", None, "Not modeled", "—"),
]

wb = Workbook()
sh = wb.active
sh.title = "BOM"
sh.append(["Component", "Mass (g/cell)", "kg/kWh", "Formula", "Source"])
bold = Font(bold=True)
for c in sh[1]:
    c.font = bold
for name, g, formula, src in bom_rows:
    if g is None:
        sh.append([name, "Not modeled", "—", formula, src])
    else:
        sh.append([name, round(g, 4), round(g / 1000.0 / E_KWH, 4), formula, src])
sh.append(["TOTAL (incl. electrolyte estimate)", round(total_g, 4), round(total_g / 1000.0 / E_KWH, 4),
           "36.226 g (contract, electrolyte excluded) + electrolyte estimate", "calc-energy + design_spec sec.3"])
sh.append(["Cell energy", E_WH, "kWh:", round(E_KWH, 6), "V*I time integration", "calc-energy energy_wh"])
wb.save(deliv / "bom.xlsx")

# ---- calc.xlsx ----
def calc_sheet(wb, title, header, rows):
    s = wb.create_sheet(title)
    s.append(header)
    for c in s[1]:
        c.font = bold
    for r in rows:
        s.append(r)
    widths = [max(len(str(x)) for x in col) for col in zip(header, *rows)]
    for i, w in enumerate(widths):
        s.column_dimensions[chr(65 + i)].width = min(max(w + 2, 12), 80)
    return s

wb2 = Workbook()
wb2.remove(wb2.active)
calc_sheet(wb2, "1_Input_parameters", ["Parameter", "Value", "Unit", "Source"], [
    ["Positive electrode thickness", 75.6e-6, "m", "Chen2020"],
    ["Negative electrode thickness", 85.2e-6, "m", "params_r5_slim.json"],
    ["Separator thickness", 10e-6, "m", "params_r5_slim.json"],
    ["Positive CC thickness (Al)", 10e-6, "m", "params_r5_slim.json"],
    ["Negative CC thickness (Cu)", 6e-6, "m", "params_r5_slim.json"],
    ["Electrode area (height x width)", AREA, "m2", "Chen2020 0.065 x 1.58"],
    ["Positive porosity", 0.335, "—", "Chen2020"],
    ["Negative porosity", 0.25, "—", "Chen2020"],
    ["Separator porosity", 0.47, "—", "Chen2020"],
    ["Positive electrode density", 3262.0, "kg/m3", "Chen2020"],
    ["Negative electrode density", 1657.0, "kg/m3", "Chen2020"],
    ["Separator density", 397.0, "kg/m3", "Chen2020"],
    ["c_max positive", 63104.0, "mol/m3", "Chen2020"],
    ["c_max negative", 33133.0, "mol/m3", "Chen2020"],
    ["Heat transfer coefficient h", 100.0, "W/m2/K", "params_r5_slim.json"],
    ["Cooling surface area", 0.00531, "m2", "Chen2020"],
    ["Ambient (4C/aging)", 318.15, "K", "pybamm_runner.PROTOCOLS"],
])
calc_sheet(wb2, "2_Capacity_energy", ["Quantity", "Value", "Unit", "Formula", "Source"], [
    ["1C discharge capacity", 5.0065, "Ah", "simulation result", "cell/r5_slim_1c_dfn.json"],
    ["Rated energy", 17.8042, "Wh", "simulation V*I integral", "cell/r5_slim_energy.json"],
    ["Discharge midpoint voltage", 3.8263, "V", "simulation", "cell/r5_slim_energy.json"],
    ["4C charge CC duration", 379.7, "s", "capacity_ah*3600/C_rate = 0.42187*900", "cell/r5_slim_4c45.json (protocol detail, dvpr item 3)"],
])
calc_sheet(wb2, "3_Energy_density", ["Quantity", "Value", "Unit", "Formula", "Source"], [
    ["Gravimetric energy density", 491.471, "Wh/kg", "17.8042 Wh / 0.0362264 kg", "cell/r5_slim_energy.json"],
    ["Volumetric energy density", 928.06, "Wh/L", "17.8042 Wh / 0.019184 L", "cell/r5_slim_energy.json"],
    ["Mass (contract caliber)", 36.2264, "g", "sum of layer masses, electrolyte excluded", "cell/r5_slim_energy.json"],
    ["Criterion ED >= 327.18", "PASS (+164.3)", "Wh/kg", "mechanical log-evaluate R5", "eval_batch_r5.json + log.jsonl"],
])
calc_sheet(wb2, "4_NP_mass", ["Quantity", "Value", "Unit", "Formula", "Source"], [
    ["Negative capacity density", 666017, "Ah/m3", "33133*0.75*96485/3600", "Chen2020 c_max_neg x (1-eps)"],
    ["Positive capacity density", 1124723, "Ah/m3", "63104*0.665*96485/3600", "Chen2020 c_max_pos x (1-eps)"],
    ["Negative areal capacity", 56.74, "Ah/m2", "666017 x 85.2e-6", "formula"],
    ["Positive areal capacity", 85.03, "Ah/m2", "1124723 x 75.6e-6", "formula"],
    ["N/P", 0.667, "—", "56.74 / 85.03 (spec formula; usable-window N/P = 1.00 by charge balance)", "design_spec sec.2"],
    ["Positive coating mass", round(pos_coat, 4), "g/cell", "layer_kg_m2 x area", "calc-energy"],
    ["Negative coating mass", round(neg_coat, 4), "g/cell", "layer_kg_m2 x area", "calc-energy"],
    ["Al CC mass", round(al_g, 4), "g/cell", "layer_kg_m2 x area", "calc-energy"],
    ["Cu CC mass", round(cu_g, 4), "g/cell", "layer_kg_m2 x area", "calc-energy"],
    ["Separator mass", round(sep_g, 4), "g/cell", "layer_kg_m2 x area", "calc-energy"],
    ["Total mass (no electrolyte)", round(TOT_G, 4), "g/cell", "sum", "calc-energy mass_kg"],
])
calc_sheet(wb2, "5_Process_parameters", ["Parameter", "Value", "Unit", "Formula", "Source"], [
    ["Positive areal density", 163.99, "g/m2", "thickness x (1-eps) x density", "calc-energy layer_kg_m2"],
    ["Negative areal density", 105.88, "g/m2", "thickness x (1-eps) x density", "calc-energy layer_kg_m2"],
    ["Positive compaction density", 2.169, "g/cm3", "3262 x 0.665 / 1000  (kg/m3 -> g/cm3, divide by 1000)", "formula"],
    ["Negative compaction density", 1.243, "g/cm3", "1657 x 0.75 / 1000", "formula"],
    ["Electrolyte fill amount", 6.33, "g/cell", "pore volume (5.27 cm3) x 1.2 g/cm3 (literature density) x fill factor 1.0 (assumed)", "design_spec sec.3"],
    ["Formation recommendation", "0.1C CC to 4.2V, 25C, 2 cycles", "—", "design recommended value; production value requires tuning", "design_spec sec.3"],
])
wb2.save(deliv / "calc.xlsx")

# ---- datasheet.docx ----
doc = Document()
doc.add_heading("Technical Datasheet — VBF-T7R1NOFORCE-DSH-01", 0)
p = doc.add_paragraph()
p.add_run("Case t7_r1_noforce (HEV cell, NMC811/graphite) — 2026-08-25. Values from parameter set / simulation outputs; see datasheet.md for the editable source.")
rows = [
    ("Rated capacity", "5.0 Ah nominal / 5.0065 Ah (1C sim)"),
    ("Voltage window", "2.5 – 4.2 V (midpoint 3.826 V)"),
    ("Rated energy", "17.804 Wh"),
    ("Energy density", "491.47 Wh/kg; 928.06 Wh/L"),
    ("Fast charge", "4C @45C: no plating (anode min +0.0153 V), T_max 325.54 K"),
    ("Max continuous discharge", "1C verified (5.0065 Ah, T_max 299.43 K); higher not simulated"),
    ("Operating temperature", "25 – 45 °C verified by simulation"),
    ("Cycle life", "Not determined (100 cyc 45C SEI simulated: 510.5 nm)"),
    ("Safety", "4C no plating ✓; nail 10 W triggered=false (T_final 317.0 K) ✓"),
    ("Dimensions", "65 mm x 1580 mm (unfolded); 186.8 µm stack"),
    ("Mass", "36.23 g (no electrolyte) / ~42.6 g incl. electrolyte est."),
    ("DCR / power density", "2.6415 mOhm / 43.24 kW/kg"),
    ("Cooling requirement", "h >= 100 W/m2/K (load-bearing design element)"),
]
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Table Grid"
tbl.rows[0].cells[0].text = "Field"
tbl.rows[0].cells[1].text = "Value"
for k, v in rows:
    r = tbl.add_row().cells
    r[0].text = k
    r[1].text = v
doc.save(deliv / "datasheet.docx")
print("bom.xlsx, calc.xlsx, datasheet.docx written")
print("total_g incl electrolyte:", round(total_g, 4), "kg/kWh:", round(total_g / 1000.0 / E_KWH, 4))
