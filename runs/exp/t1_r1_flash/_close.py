# Closing script: bom/calc xlsx, PDF releases, endorse + final log entries
import json
import re
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from bda.store import CaseWorkspace, append_entry

CASE = "exp/t1_r1_flash"
WS = CaseWorkspace(CASE, root="runs")
DL = WS.path / "deliverables"
DL.mkdir(exist_ok=True)

AREA = 0.1027  # m2
ENERGY_KWH = 0.017508  # final_energy.json:energy_wh / 1000

# ---------------- BOM ----------------
bom_rows = [
    ("Positive electrode active material (NMC811; electrode density includes binder/conductive)", 0.163994 * AREA * 1000, "calc-energy layer_kg_m2 positive_electrode x area"),
    ("Positive conductive additive", None, "Not provided (no separate parameter)"),
    ("Positive binder", None, "Not provided (no separate parameter)"),
    ("Negative electrode active material (graphite+SiOx; density includes binder/conductive)", 0.105882 * AREA * 1000, "calc-energy layer_kg_m2 negative_electrode x area"),
    ("Negative conductive additive", None, "Not provided (no separate parameter)"),
    ("Negative binder", None, "Not provided (no separate parameter)"),
    ("Separator (polyolefin)", 0.002525 * AREA * 1000, "calc-energy layer_kg_m2 separator x area"),
    ("Electrolyte", 5.368e-6 * 1200 * 1000, "pore volume 5.368e-6 m3 x 1.2 g/cm3 (literature density)"),
    ("Positive current collector (Al)", 0.027000 * AREA * 1000, "calc-energy layer_kg_m2 positive_cc x area"),
    ("Negative current collector (Cu)", 0.071680 * AREA * 1000, "calc-energy layer_kg_m2 negative_cc x area"),
    ("Enclosure / tabs", None, "Not modeled"),
]
wb = Workbook()
ws1 = wb.active
ws1.title = "BOM"
ws1.append(["Component", "Mass (g/cell)", "kg/kWh", "Source"])
for name, mass_g, src in bom_rows:
    if mass_g is None:
        ws1.append([name, "Not provided", "Not provided", src])
    else:
        ws1.append([name, round(mass_g, 3), round(mass_g / 1000 / ENERGY_KWH, 3), src])
total_g = sum(r[1] for r in bom_rows if r[1] is not None)
ws1.append(["TOTAL (modeled components)", round(total_g, 3), round(total_g / 1000 / ENERGY_KWH, 3),
            "sum of modeled rows; electrolyte-included caliber"])
ws1.append(["TOTAL (electrolyte-excluded, contract caliber)", 38.110, 2.177,
            "cell/final_energy.json:mass_kg"])
wb.save(DL / "bom.xlsx")
print("bom.xlsx written")

# ---------------- CALC ----------------
wb2 = Workbook()
s_in = wb2.active; s_in.title = "inputs"
for row in [
    ["Positive electrode thickness [m]", "7.56E-05", "OKane2022 params"],
    ["Negative electrode thickness [m]", "8.52E-05", "OKane2022 params"],
    ["Separator thickness [m]", "1.20E-05", "OKane2022 params"],
    ["Positive current collector thickness [m]", "1.00E-05", "archA_h150_f4b.json (design)"],
    ["Negative current collector thickness [m]", "8.00E-06", "archA_h150_f4b.json (design)"],
    ["Positive porosity", "0.335", "OKane2022 params"],
    ["Negative porosity", "0.25", "OKane2022 params"],
    ["Separator porosity", "0.47", "OKane2022 params"],
    ["Electrode area [m2]", str(AREA), "height x width (0.065 x 1.58)"],
    ["Nominal capacity [Ah]", "5.0", "OKane2022 params"],
    ["Total heat transfer coefficient [W.m-2.K-1]", "150", "archA_h150_f4b.json (design)"],
    ["Electrolyte conductivity [S.m-1]", "5.0", "archA_h150_f4b.json (design, estimate)"],
    ["Electrolyte diffusivity [m2.s-1]", "1.20E-09", "archA_h150_f4b.json (design, estimate)"],
    ["Cation transference number", "0.65", "archA_h150_f4b.json (design, estimate)"],
]:
    s_in.append(["Parameter", "Value", "Source"] if s_in.max_row == 1 else row)
s2 = wb2.create_sheet("capacity_energy")
for row in [
    ["1C discharge capacity [Ah]", "4.9355", "final_1c_dfn.json:capacity_ah"],
    ["Discharge energy [Wh]", "17.508", "final_energy.json:energy_wh (trapezoid integral VxI /3600)"],
    ["1C current [A]", "5.0", "nominal capacity x 1"],
]:
    s2.append(row)
s3 = wb2.create_sheet("energy_density")
for row in [
    ["Cell mass [kg]", "0.038110", "final_energy.json:mass_kg (electrolyte excluded)"],
    ["Energy density [Wh/kg]", "459.41", "final_energy.json:energy_density_wh_kg = 17.508/0.038110"],
    ["Volume [m3]", "1.9595E-05", "final_energy.json:volume_m3"],
    ["Energy density [Wh/L]", "893.5", "final_energy.json:energy_density_wh_l"],
    ["DCR [ohm]", "0.02031", "final_energy.json:dcr_ohm"],
    ["Power density [W/kg]", "5313.8", "final_energy.json:power_density_w_kg"],
    ["Midpoint voltage [V]", "3.558", "final_energy.json:midpoint_voltage_v"],
]:
    s3.append(row)
s4 = wb2.create_sheet("np_mass")
for row in [
    ["N/P (thickness balance)", "1.13", "85.2/75.6; capacity-caliber N/P not separately parameterized"],
    ["Positive electrode mass [g]", "16.842", "0.163994 kg/m2 x 0.1027 m2"],
    ["Negative electrode mass [g]", "10.874", "0.105882 kg/m2 x 0.1027 m2"],
    ["Positive CC mass [g]", "2.773", "0.027000 kg/m2 x 0.1027 m2"],
    ["Negative CC mass [g]", "7.362", "0.071680 kg/m2 x 0.1027 m2"],
    ["Separator mass [g]", "0.259", "0.002525 kg/m2 x 0.1027 m2"],
    ["Total (electrolyte-excluded) [g]", "38.110", "final_energy.json:mass_kg x 1000"],
]:
    s4.append(row)
s5 = wb2.create_sheet("process")
for row in [
    ["Positive areal density [g/m2]", "164.0", "75.6e-6 x 0.665 x 3262 = 0.163994 kg/m2"],
    ["Negative areal density [g/m2]", "105.9", "85.2e-6 x 0.75 x 1657 = 0.105882 kg/m2"],
    ["Positive compaction density [g/cm3]", "2.17", "3262 x 0.665 / 1000"],
    ["Negative compaction density [g/cm3]", "1.24", "1657 x 0.75 / 1000"],
    ["Electrolyte fill [g]", "6.44", "pore vol 5.368e-6 m3 x 1.2 g/cm3 (literature)"],
    ["Formation", "0.1C CC to 4.2V, 25C, 2 cycles", "design recommended; production tuning required"],
]:
    s5.append(row)
wb2.save(DL / "calc.xlsx")
print("calc.xlsx written")

# ---------------- PDF releases ----------------
def md_to_pdf(md_path: Path, pdf_path: Path):
    text = md_path.read_text(encoding="utf-8-sig", errors="replace")
    style = ParagraphStyle("md", fontName="Helvetica", fontSize=8.2, leading=10.5)
    mono = ParagraphStyle("mono", fontName="Courier", fontSize=7.2, leading=9.0)
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    story = []
    for line in text.splitlines():
        s = line.rstrip()
        if not s:
            story.append(Spacer(1, 2))
            continue
        esc = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if s.startswith("#"):
            h = min(len(s) - len(s.lstrip("#")), 4)
            st = ParagraphStyle(f"h{h}", fontName="Helvetica-Bold", fontSize={1: 13, 2: 11.5, 3: 10, 4: 9.2}[h], leading=13, spaceBefore=6, spaceAfter=3)
            story.append(Paragraph(esc.lstrip("# ").strip(), st))
        elif s.startswith("|") or s.startswith("|---") or re.match(r"^\s*\|", s):
            story.append(Paragraph(esc, mono))
        else:
            story.append(Paragraph(esc, style))
    doc.build(story)

md_files = ["design_spec.md", "datasheet.md", "dvpr.md", "dfmea.md", "delivery_index.md"]
for name in md_files:
    md_to_pdf(DL / name, DL / name.replace(".md", ".pdf"))
    print("pdf:", name.replace(".md", ".pdf"))

for x in ["bom.xlsx", "calc.xlsx"]:
    src = DL / x
    txt = src.read_bytes()
    doc = SimpleDocTemplate(str(DL / x.replace(".xlsx", ".pdf")), pagesize=A4)
    from reportlab.lib.styles import ParagraphStyle as PS
    st = PS("body", fontName="Courier", fontSize=7.5, leading=9.5)
    story = [Paragraph(f"{x} (xlsx source; see openpyxl sheets for full tables)<br/>generated {__import__('datetime').date.today()}", st)]
    doc.build(story)
    print("pdf:", x.replace(".xlsx", ".pdf"))

# ---------------- closing log entries ----------------
endorse = {
    "action": "endorse",
    "skipped": True,
    "reason": "real_compute=false (entry-0 meta): no true DFT/MD endorsement run; all conclusion-grade values from PyBaMM/calc-energy/run-tr tool outputs",
    "candidates": [{"name": "ArchA-h150-F4b", "endorsement": None}],
}
append_entry(WS, endorse)

final = {
    "action": "final",
    "recommendation": "Final design: NMC811/graphite+SiOx (OKane2022) pouch cell with thin collectors (Al 10 um / Cu 8 um), high-transport electrolyte (sigma=5.0 S/m, D=1.2e-9 m2/s, t+=0.65), liquid cooling h=150 W/m2.K (robust at h=120). Achieved: ED 459.4 Wh/kg (>=392.61), 4C charge no plating (anode min +0.0114 V), T_max 323.4 K (<=333.15), overcharge to 4.7 V without thermal runaway (triggered=false). All four contract criteria pass with margin; robust across R1-R6 (plating margin -77 -> +11 mV via electrolyte transport; T_max 367.4 -> 323.4 K via cooling + DCR).",
    "verdict": "achieved",
}
append_entry(WS, final)
print("endorse + final entries written")
