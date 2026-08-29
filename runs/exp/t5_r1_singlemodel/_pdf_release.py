# -*- coding: utf-8 -*-
"""PDF release versions for all 7 deliverable categories (reportlab; DejaVuSans if available)."""
from pathlib import Path

import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

WS = Path(r"D:\research\degradation_prognostics\Battery_Design_Agent\runs\exp\t5_r1_singlemodel")
DLV = WS / "deliverables"

# font: try DejaVuSans (ships with matplotlib) for full unicode; else Helvetica + transliteration
FONT = "Helvetica"
FONT_B = "Helvetica-Bold"
try:
    from matplotlib import font_manager
    path = font_manager.findfont("DejaVu Sans")
    pdfmetrics.registerFont(TTFont("DejaVu", str(path)))
    FONT, FONT_B = "DejaVu", "DejaVu"
except Exception:
    pass

def _trans(s):
    if FONT == "DejaVu":
        return s
    for a, b in [("✓", "[PASS]"), ("✗", "[FAIL]"), ("≤", "<="), ("≥", ">="), ("→", "->"),
                 ("⁻", "-"), ("⁺", "+"), ("⁹", "9"), ("²", "2"), ("³", "3"),
                 ("·", "-"), ("×", "x"), ("µm", "um"), ("°", " deg "), ("—", "-")]:
        s = s.replace(a, b)
    return s

def md_pdf(md_name, pdf_name, title):
    doc = SimpleDocTemplate(str(DLV / pdf_name), pagesize=A4,
                            leftMargin=15 * mm, rightMargin=15 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    h1 = ParagraphStyle("h1", fontName=FONT_B, fontSize=13, leading=16, spaceAfter=8)
    h2 = ParagraphStyle("h2", fontName=FONT_B, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("body", fontName=FONT, fontSize=8, leading=10.5, spaceAfter=3)
    tbl = ParagraphStyle("tbl", fontName=FONT, fontSize=6.2, leading=8.2)
    story = [Paragraph(_trans(f"# {title}"), h1)]
    for line in (DLV / md_name).read_text(encoding="utf-8").splitlines():
        s = line.rstrip()
        if not s.strip():
            story.append(Spacer(1, 3))
        elif s.startswith("## "):
            story.append(Paragraph(_trans(s[3:]), h2))
        elif s.startswith("# "):
            continue
        elif s.startswith("|"):
            story.append(Preformatted(_trans(s), tbl))
        elif s.startswith("> ") or s.startswith("- "):
            story.append(Paragraph(_trans(s), body))
        else:
            story.append(Paragraph(_trans(s), body))
    doc.build(story)
    print(pdf_name, (DLV / pdf_name).stat().st_size, "bytes")

def xlsx_pdf(xlsx_name, pdf_name, title):
    doc = SimpleDocTemplate(str(DLV / pdf_name), pagesize=A4,
                            leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    h1 = ParagraphStyle("h1", fontName=FONT_B, fontSize=13, leading=16, spaceAfter=8)
    h2 = ParagraphStyle("h2", fontName=FONT_B, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    cell = ParagraphStyle("cell", fontName=FONT, fontSize=6.2, leading=8.2)
    story = [Paragraph(_trans(f"# {title}"), h1)]
    wb = openpyxl.load_workbook(DLV / xlsx_name, read_only=True, data_only=True)
    for sname in wb.sheetnames:
        ws = wb[sname]
        rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
        if not rows:
            continue
        story.append(Paragraph(_trans(f"## Sheet: {sname}"), h2))
        data = [[Paragraph(_trans("" if v is None else str(v)), cell) for v in row] for row in rows]
        col_w = [24, 60, 34, 22, 70][: len(rows[0])]
        t = Table(data, colWidths=[w * mm for w in col_w], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.92, 0.92, 0.95)),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.97, 0.97, 0.97)]),
        ]))
        story.append(t)
        story.append(Spacer(1, 8))
    doc.build(story)
    print(pdf_name, (DLV / pdf_name).stat().st_size, "bytes")

md_pdf("design_spec.md", "design_spec.pdf", "Cell Design Specification - archN_robust (VBF-T5R1SINGLEMODEL-DS-001)")
xlsx_pdf("bom.xlsx", "bom.pdf", "Bill of Materials (VBF-T5R1SINGLEMODEL-BOM-001)")
md_pdf("datasheet.md", "datasheet.pdf", "Technical Datasheet - archN_robust (VBF-T5R1SINGLEMODEL-DSH-001)")
xlsx_pdf("calc.xlsx", "calc.pdf", "Design Calculation Sheet (VBF-T5R1SINGLEMODEL-CALC-001)")
md_pdf("dvpr.md", "dvpr.pdf", "Design Verification Plan & Report - Virtual Test (VBF-T5R1SINGLEMODEL-DVPR-001)")
md_pdf("dfmea.md", "dfmea.pdf", "Design FMEA - Qualitative (VBF-T5R1SINGLEMODEL-DFMEA-001)")
md_pdf("delivery_index.md", "delivery_index.pdf", "Delivery Index - t5_r1_singlemodel (VBF-T5R1SINGLEMODEL-IDX-001)")
print("all PDFs written")
