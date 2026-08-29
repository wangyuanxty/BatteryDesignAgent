# -*- coding: utf-8 -*-
"""t5_r2 close-out (c): PDF releases for all deliverables + delivery index (md + pdf)."""
from pathlib import Path

import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable, PageBreak)

CASE = Path(r"runs\exp\t5_r2")
OUT = CASE / "deliverables"

DEEP = colors.HexColor("#14283C")
MID = colors.HexColor("#1E5A8A")
COPPER = colors.HexColor("#C97B3D")
LIGHT = colors.HexColor("#EEF3F8")

S = {
    "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=16, textColor=DEEP,
                          spaceAfter=6, spaceBefore=10, leading=19),
    "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12.5, textColor=MID,
                          spaceAfter=4, spaceBefore=8, leading=15),
    "h3": ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=10.5, textColor=COPPER,
                          spaceAfter=3, spaceBefore=6, leading=13),
    "body": ParagraphStyle("body", fontName="Helvetica", fontSize=8.6, leading=11.5,
                            spaceAfter=3, textColor=colors.black),
    "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=7.2, leading=9.0,
                            textColor=colors.black),
    "cellb": ParagraphStyle("cellb", fontName="Helvetica-Bold", fontSize=7.2, leading=9.0,
                             textColor=colors.white),
    "bullet": ParagraphStyle("bullet", fontName="Helvetica", fontSize=8.6, leading=11.5,
                              leftIndent=12, spaceAfter=2),
}


def _esc(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    i = 0
    while True:
        i = t.find("**", i)
        if i < 0:
            break
        j = t.find("**", i + 2)
        if j < 0:
            break
        t = t[:i] + "<b>" + t[i+2:j] + "</b>" + t[j+2:]
        i = j
    return t


def _pdf_from_md(md_path, pdf_path, title):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=16*mm, bottomMargin=14*mm,
                            title=title)
    story = []
    tbl = []
    for ln in lines:
        s = ln.strip()
        if not s:
            if tbl:
                story.append(_mk_table(tbl))
                tbl = []
            continue
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            tbl.append(cells)
            continue
        if tbl:
            story.append(_mk_table(tbl))
            tbl = []
        if s.startswith("### "):
            story.append(Paragraph(_esc(s[4:]), S["h3"]))
        elif s.startswith("## "):
            story.append(Paragraph(_esc(s[3:]), S["h2"]))
        elif s.startswith("# "):
            story.append(Paragraph(_esc(s[2:]), S["h1"]))
        elif s == "---":
            story.append(HRFlowable(width="100%", thickness=0.7, color=MID, spaceAfter=6))
        elif s.startswith("- "):
            story.append(Paragraph(_esc(s[2:]), S["bullet"]))
        else:
            story.append(Paragraph(_esc(s), S["body"]))
    if tbl:
        story.append(_mk_table(tbl))
    doc.build(story)


def _mk_table(rows):
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    data = []
    for i, r in enumerate(rows):
        if i == 0:
            data.append([Paragraph(_esc(c), S["cellb"]) for c in r])
        else:
            data.append([Paragraph(_esc(c), S["cell"]) for c in r])
    t = Table(data, repeatRows=1, hAlign="LEFT")
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), MID),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9FB3C8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    t.setStyle(TableStyle(style))
    return t


def _pdf_from_xlsx(xlsx_path, pdf_path, title):
    wb = openpyxl.load_workbook(xlsx_path)
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                            leftMargin=12*mm, rightMargin=12*mm,
                            topMargin=16*mm, bottomMargin=14*mm, title=title)
    story = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            rows.append(["" if v is None else str(v) for v in row])
        if rows:
            story.append(Paragraph(f"Sheet: {name}", S["h2"]))
            story.append(_mk_table(rows))
            story.append(PageBreak())
    doc.build(story)


# ---- 1. markdown PDFs --------------------------------------------------------
_pdf_from_md(OUT / "design_spec.md", OUT / "design_spec.pdf", "VBF-T5R2-DS-01 design specification")
_pdf_from_md(OUT / "datasheet.md", OUT / "datasheet.pdf", "VBF-T5R2-DSH-01 datasheet")
_pdf_from_md(OUT / "dvpr.md", OUT / "dvpr.pdf", "VBF-T5R2-DVPR-01 DVP&R")
_pdf_from_md(OUT / "dfmea.md", OUT / "dfmea.pdf", "VBF-T5R2-DFMEA-01 DFMEA")

# ---- 2. xlsx PDFs -------------------------------------------------------------
_pdf_from_xlsx(OUT / "bom.xlsx", OUT / "bom.pdf", "VBF-T5R2-BOM-01 bill of materials")
_pdf_from_xlsx(OUT / "calc.xlsx", OUT / "calc.pdf", "VBF-T5R2-CALC-01 calculation sheet")

# ---- 3. delivery index (registers only files that now exist) -----------------
rows = [
    ("VBF-T5R2-DS-01",   "design_spec.md",   "md",   "Design specification - stack/parameters/process/verification; generated per deliverable-design-spec spec"),
    ("VBF-T5R2-DS-02",   "design_spec.pdf",  "pdf",  "PDF release of VBF-T5R2-DS-01 (reportlab)"),
    ("VBF-T5R2-DSH-01",  "datasheet.md",     "md",   "Cell datasheet - finalist f2_h70_C verified figures"),
    ("VBF-T5R2-DSH-02",  "datasheet.pdf",    "pdf",  "PDF release of VBF-T5R2-DSH-01 (reportlab)"),
    ("VBF-T5R2-BOM-01",  "bom.xlsx",         "xlsx", "Bill of materials - dual caliber g/cell + kg/kWh, estimates marked"),
    ("VBF-T5R2-BOM-02",  "bom.pdf",          "pdf",  "PDF release of VBF-T5R2-BOM-01 (openpyxl -> reportlab)"),
    ("VBF-T5R2-CALC-01", "calc.xlsx",        "xlsx", "5-sheet calculation chain: inputs -> mass -> energy -> stack/process -> verification"),
    ("VBF-T5R2-CALC-02", "calc.pdf",         "pdf",  "PDF release of VBF-T5R2-CALC-01"),
    ("VBF-T5R2-DVPR-01", "dvpr.md",          "md",   "Design verification plan & report - 6 PASS rows + 5 honest N/A rows"),
    ("VBF-T5R2-DVPR-02", "dvpr.pdf",         "pdf",  "PDF release of VBF-T5R2-DVPR-01"),
    ("VBF-T5R2-DFMEA-01","dfmea.md",         "md",   "Design FMEA - qualitative S/O/RPN with simulation signals"),
    ("VBF-T5R2-DFMEA-02","dfmea.pdf",        "pdf",  "PDF release of VBF-T5R2-DFMEA-01"),
    ("VBF-T5R2-IDX-01",  "delivery_index.md","md",   "This index - cover + controlled file list of the design package"),
    ("VBF-T5R2-IDX-02",  "delivery_index.pdf","pdf", "PDF release of this index (blueprint cover)"),
]
missing = [r for r in rows if not (OUT / r[1]).exists() and r[1] != "delivery_index.md"
           and r[1] != "delivery_index.pdf"]
if missing:
    raise SystemExit(f"refusing to register missing files: {missing}")

idx_md = """# Delivery Package Index - VBF-T5R2-IDX-01

**Case**: t5_r2 **Verdict**: achieved **Date**: 2026-08-26 **Protocol**: Virtual Battery Factory

## Status summary

| Criterion | Threshold | Measured | Verdict |
|---|---|---|---|
| Energy density | >= 500.94 Wh/kg | 666.30 | PASS |
| Max temperature @4C | <= 333.15 K | 330.98 K | PASS |
| Lithium plating @4C | none | anode potential min +0.0506 V | PASS |

closing audit chain: entry 0 contract -> plan -> ceiling -> 7 propose rounds each with
same-round evaluate -> endorse (skipped, real_compute=false, justified) -> final (achieved).

## Controlled file list (only actually generated files are registered)

| VBF number | File | Format | Source |
|---|---|---|---|
"""
for num, fn, fmt, desc in rows:
    if fn in ("delivery_index.md", "delivery_index.pdf"):
        continue
    idx_md += f"| {num} | {fn} | {fmt} | {desc} |\n"
idx_md += """| VBF-T5R2-IDX-01 | delivery_index.md | md | This index |
| VBF-T5R2-IDX-02 | delivery_index.pdf | pdf | This index (PDF release) |

Notes:
- cell_model.stl was **not** requested by the task and is therefore not provided (recorded
  honestly in log.jsonl `final` entry caveats).
- report.html is a post-close audit presentation produced by `bda render`; it is not part
  of this controlled package.
- Simulation outputs and the audit ledger live in the case workspace (cell/*.json, log.jsonl,
  bridge/p_r6_f2.json) as the evidence chain for every number above.
"""
(OUT / "delivery_index.md").write_text(idx_md, encoding="utf-8")

# ---- index PDF with blueprint cover ------------------------------------------
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate
from reportlab.lib.styles import getSampleStyleSheet

doc = SimpleDocTemplate(str(OUT / "delivery_index.pdf"), pagesize=A4,
                        leftMargin=15*mm, rightMargin=15*mm,
                        topMargin=18*mm, bottomMargin=14*mm,
                        title="VBF-T5R2 delivery index")
story = []
cover = Table(
    [[Paragraph("VIRTUAL BATTERY FACTORY", ParagraphStyle(
        "covk", fontName="Helvetica-Bold", fontSize=11, textColor=COPPER,
        leading=14, spaceAfter=2))],
     [Paragraph(f"Delivery Package Index", ParagraphStyle(
        "covt", fontName="Helvetica-Bold", fontSize=22, textColor=colors.white,
        leading=26))],
     [Paragraph(f"Case T5-R2 &nbsp;|&nbsp; verdict: ACHIEVED &nbsp;|&nbsp; 2026-08-26", ParagraphStyle(
        "covs", fontName="Helvetica", fontSize=10.5, textColor=LIGHT, leading=14))],
     [Paragraph("ED 666.30 Wh/kg >= 500.94 &nbsp;&middot;&nbsp; T_max 330.98 K <= 333.15 &nbsp;&middot;&nbsp; "
                "plated = false (ap_min +0.0506 V)", ParagraphStyle(
        "covm", fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#BFD1E0"),
        leading=12))],
     ], colWidths=[180*mm])
cover.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, -1), DEEP),
    ("LEFTPADDING", (0, 0), (-1, -1), 10*mm),
    ("RIGHTPADDING", (0, 0), (-1, -1), 10*mm),
    ("TOPPADDING", (0, 0), (-1, -1), 8*mm),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 8*mm),
    ("LINEBELOW", (0, 0), (-1, 0), 1.2, COPPER),
]))
story.append(cover)
story.append(Spacer(1, 8*mm))
idx_lines = idx_md.splitlines()
tbl = []
for ln in idx_lines:
    s = ln.strip()
    if s.startswith("|"):
        tbl.append([c.strip() for c in s.strip("|").split("|")])
    else:
        if tbl:
            story.append(_mk_table(tbl))
            tbl = []
        if s.startswith("# "):
            story.append(Paragraph(_esc(s[2:]), S["h1"]))
        elif s.startswith("## "):
            story.append(Paragraph(_esc(s[3:]), S["h2"]))
        elif s.startswith("- "):
            story.append(Paragraph(_esc(s[2:]), S["bullet"]))
        elif s and not s.startswith("```"):
            story.append(Paragraph(_esc(s), S["body"]))
if tbl:
    story.append(_mk_table(tbl))
doc.build(story)

print("PDFs done:")
for p in sorted(OUT.glob("*")):
    print(" ", p.name, p.stat().st_size, "bytes")