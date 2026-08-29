# -*- coding: utf-8 -*-
"""t3_r3 closing deliverables, part B: write design_spec.md / dvpr.md / dfmea.md /
delivery_index.md (all numbers from deliv_metrics.json + tool outputs + entry 0),
then render every required PDF with reportlab.

Agent-built input files: allowed. Command-output JSONs: read-only.
"""
import json
from pathlib import Path

DLV = Path("runs/exp/t3_r3/deliverables")
CELL = Path("runs/exp/t3_r3/cell")
M = json.load(open(DLV / "deliv_metrics.json", encoding="utf-8"))
entry0 = json.loads(open("runs/exp/t3_r3/log.jsonl", encoding="utf-8").readline())
CR = entry0["criteria"]
TH = {
    "cap_min": CR["stage2"]["capacity_ah"]["min"],
    "ret_min": CR["stage2"]["retention_5c"]["min"],
    "pd_min": CR["stage2"]["power_density_w_kg"]["min"],
    "tmax": CR["stage3"]["T_max_K"]["max"],
}
c1 = json.load(open(CELL / "r5_v12_1c_dfn.json", encoding="utf-8"))
c5 = json.load(open(CELL / "r5_v12_5c_dfn.json", encoding="utf-8"))
c4 = json.load(open(CELL / "r5_v12_4c_dfn.json", encoding="utf-8"))

# ---------- shared formatted values ----------
cap1 = c1["capacity_ah"]
cap5 = c5["capacity_ah"]
ret = M["retention_5c"]
t5 = c5["T_max_K"]
t4 = c4["T_max_K"]
amin = min(c4["anode_potential_v"])
chg4 = c4["capacity_ah"]
mk = M  # metrics dict alias

def kc(k):
    return f"{k:.2f} K ({k-273.15:.2f} C)"

F = {
    "cap1": f"{cap1:.4f}",
    "cap5": f"{cap5:.4f}",
    "ret_pct": f"{ret*100:.2f}%",
    "ret_raw": f"{ret:.6f}",
    "t5s": kc(t5), "t4s": kc(t4),
    "amin_mv": f"+{amin*1e3:.1f}",
    "chg4": f"{chg4:.3f}",
    "ewh": f"{mk['energy_wh']:.3f}",
    "ed": f"{mk['ed_wh_kg']:.1f}",
    "edl": f"{mk['ed_wh_l']:.1f}",
    "pd": f"{mk['power_density_w_kg']:,.0f}",
    "mass": f"{mk['mass_kg']*1e3:.3f}",
    "gstack": f"{mk['g_stack']:.3f}",
    "gelec": f"{mk['g_elec']:.3f}",
    "gtot": f"{mk['g_total']:.3f}",
    "thick": f"{mk['total_thickness_m']*1e6:.2f}",
    "lpos": f"{mk['L_p']*1e6:.3f}", "lneg": f"{mk['L_n']*1e6:.3f}",
    "lsep": f"{mk['L_sep']*1e6:.1f}",
    "rpos": f"{mk['r_p']*1e6:.2f}", "rneg": f"{mk['r_n']*1e6:.2f}",
    "ccp": f"{mk['th_cc_p']*1e6:.0f}", "ccn": f"{mk['th_cc_n']*1e6:.0f}",
    "hmm": f"{mk['height_m']*1e3:.1f}", "wm": f"{mk['width_m']:.3f}",
    "np": f"{mk['np_ratio']:.4f}", "npb": f"{mk['np_baseline']:.4f}",
    "npx": f"{mk['np_x']:.2f}",
    "arealp": f"{mk['areal_pos']:.2f}", "arealn": f"{mk['areal_neg']:.2f}",
    "compp": f"{mk['comp_pos_gcm3']:.3f}", "compn": f"{mk['comp_neg_gcm3']:.3f}",
    "kappa": f"{mk['kappa_e']:.1f}", "tp": f"{mk['tplus']:.2f}",
    "De": f"{mk['D_e']:.1e}", "hh": f"{mk['h']:.0f}",
    "noms": f"{mk['nominal_ah']:.1f}", "vv": f"{mk['v_min']:.1f}-{mk['v_max']:.1f}",
    "mid": f"{mk['midpoint_v']:.3f}", "dcr": f"{mk['dcr_ohm']*1e3:.3f}",
    "area": f"{mk['area_m2']:.4f}", "vol": f"{mk['volume_m3']:.6e}",
    "capmin": f"{TH['cap_min']:.1f}", "retmin": f"{TH['ret_min']*100:.0f}%",
    "pdmin": f"{TH['pd_min']:,.0f}", "tmaxth": f"{TH['tmax']:.2f}",
}

# ================= MARKDOWN DOCS =================

design_spec = """# Cell Design Specification - Power-Tool Battery (case t3_r3)

> VBF-T3R3-DS-01. All values are mechanically taken from the parameter set, simulation results, and literature, with per-line sources. Missing items are written "Not provided"; no numbers from memory.

## 1. Basic Specification

| Field | Value | Source |
|------|------|------|
| Electrochemical system | NMC811 (LiNi0.8Mn0.1Co0.1O2) / graphite | Chen2020 parameter set (Chen et al., J. Electrochem. Soc. 167 (2020) 080534); deterministic default mapping (task text names no system), recorded in funnel log entry |
| Nominal capacity | {noms} Ah nominal (parameter set); {cap1} Ah verified at 1C (DFN simulation) | r5_v12_1c_dfn.json:capacity_ah; criterion >= {capmin} Ah |
| Voltage window | {vv} V | pybamm Chen2020 parameter set |
| Cell dimensions | height {hmm} mm x width {wm} m unwound strip (wound geometry) x layer stack {thick} um; shell dimensions Not provided (no shell parameter) | pybamm Chen2020 + params |
| Electrolyte | EC/EMC + LiPF6 base (c_e0 = {ce0:.0f} mol/m3); design-lever overrides: kappa = {kappa} S/m, t+ = {tp}, D = {De} m2/s; additive candidates: none exercised (no Stage-2 escalation; architecture space judged sufficient - funnel log) | pybamm Chen2020 (base), r5_v12_hip_final_params.json (overrides) |
| Cation transference number | {tp} | r5_v12_hip_final_params.json |

## 2. Electrode and Separator

| Layer | Thickness (um) | Porosity | Particle radius (um) | Other | Source |
|------|------|------|------|------|------|
| Positive electrode (NMC811) | {lpos} | {epsp:.2f} | {rpos} | density {rhop:.0f} kg/m3 | r5_v12_hip_final_params.json / pybamm Chen2020 |
| Negative electrode (graphite) | {lneg} | {epsn:.2f} | {rneg} | density {rhon:.0f} kg/m3 | r5_v12_hip_final_params.json / pybamm Chen2020 |
| Separator | {lsep} | {epssep:.2f} | - | density {rhos:.0f} kg/m3 | params override / pybamm Chen2020 |
| Positive current collector (Al) | {ccp} | - | - | density 2700 kg/m3 | r5_v12_hip_final_params.json |
| Negative current collector (Cu) | {ccn} | - | - | density 8960 kg/m3 | r5_v12_hip_final_params.json |

N/P ratio: {np} loading ratio (negative/positive areal density) divided by the baseline loading ratio {npb} = **x{npx} in capacity terms** - a +30% negative capacity margin against the Chen2020 baseline stoichiometric parity (mechanical derivation from r5_v12_energy.json:layer_kg_m2 and r1_base_energy.json:layer_kg_m2; per-electrode specific capacities are not separately exposed by the parameter set).

## 3. Process Design Parameters

| Parameter | Formula | Value | Source |
|------|------|------|------|
| Positive areal density | thickness x (1 - porosity) x density | {arealp} g/m2 | mechanical from params |
| Negative areal density | thickness x (1 - porosity) x density | {arealn} g/m2 | mechanical from params |
| Positive compaction density | density x (1 - porosity) | {compp} g/cm3 (kg/m3 divided by 1000) | mechanical from params |
| Negative compaction density | density x (1 - porosity) | {compn} g/cm3 (kg/m3 divided by 1000) | mechanical from params |
| Electrolyte fill amount | pore volume x electrolyte density x fill factor | {gelec} g | pore volume from porosity params x 1.2 g/cm3 (literature value, annotated; fill factor 1) |
| Formation recommendation | e.g. 0.1C CC to 4.2 V, 25 C, 2 cycles | design-recommended value | actual production-line value requires tuning |
| Thermal management | cooling via h coefficient | {hh} W/m2/K | r5_v12_hip_final_params.json (design requirement on pack thermal system) |

## 4. Mass Breakdown

| Component | Mass (g) | Note |
|------|------|------|
| Positive coating (94% NMC811 / 3% CB / 3% PVDF, literature default split) | {gpos:.3f} | layer_kg_m2 x area_m2 |
| Negative coating (94% graphite / 1% CB / 3% CMC+SBR, literature default split) | {gneg:.3f} | layer_kg_m2 x area_m2 |
| Separator | {gsep:.3f} | layer_kg_m2 x area_m2 |
| Positive current collector (Al) | {gccp:.3f} | layer_kg_m2 x area_m2 |
| Negative current collector (Cu) | {gccn:.3f} | layer_kg_m2 x area_m2 |
| Electrolyte (fill) | {gelec} | pore volume x 1.2 g/cm3 (lit.) |
| **Total stack mass** (calc-energy contract caliber, electrolyte excluded) | **{gstack}** | r5_v12_energy.json:mass_kg |
| Total including electrolyte | {gtot} | mechanical sum |

The calc-energy contract formula excludes electrolyte from mass and volume (parameter set lacks electrolyte density); the electrolyte row is provided for completeness.

## 5. Performance Verification (vs entry-0 criteria)

| Metric | Simulated value | Criterion | Determination | Source |
|------|------|------|------|------|
| Nominal capacity | {cap1} Ah at 1C | >= {capmin} Ah | PASS | r5_v12_1c_dfn.json:capacity_ah |
| 5C discharge retention | {ret_pct} (5C capacity {cap5} Ah / 1C capacity) | >= {retmin} | PASS | r5_v12_derived.json:retention_5c |
| 4C fast-charge max temperature | {t4s} | <= 333.15 K (60 C) | PASS | r5_v12_4c_dfn.json:T_max_K |
| 5C discharge max temperature | {t5s} | <= 333.15 K | PASS (informational) | r5_v12_5c_dfn.json:T_max_K |
| Lithium plating at 4C | anode potential min {amin_mv} mV | < 0 V -> plated | PASS (no plating) | r5_v12_4c_dfn.json:anode_potential_v |
| Power density | {pd} W/kg | >= {pdmin} W/kg | PASS (non-binding) | r5_v12_energy.json:power_density_w_kg |
| Energy density | {ed} Wh/kg | not thresholded (information) | - | r5_v12_energy.json:energy_density_wh_kg |

## 6. Design Notes

Round-by-round changes with evaluate-log citations (all values DFN judge grade at round 5):

- **R1 baseline (evaluate_r1_Chen2020-baseline = fail)**: Chen2020 default thick-electrode energy cell is polarization-limited: 5C retention 8.7%, anode min -537 mV (plating), T_max 354 K, 4C CC acceptance 0.176 Ah.
- **R2 V1-V3 (evaluate_r2, all fail)**: porosity +0.05 and loading scale 0.55 -> retention reaches 95.0-98.2% but 4C T_max remains 340+ K; particle refinement 1.5 um alone insufficient.
- **R3 V4-V7 (evaluate_r3, all fail)**: N/P x1.3 (V4) gives clean no-plating margin (+36 mV); h=20 cooling (V5) clears 5C T_max (325.3 K); combos V6/V7 push 4C acceptance to 0.774 Ah but 4C T_max 334.2-334.5 K still over.
- **R4 V8-V11 (evaluate_r4)**: fine positive particles 0.8 um (V8, T_max 332.65 K - 0.5 K margin), fast D 8e-10 alone insufficient (V9, 333.76 K), h=25 alone passes (V10, 332.34 K); combo V11 champion at screening: T_max 329.87 K, retention 98.5%, anode +39.6 mV, acceptance 0.729 Ah.
- **R5 judge grade (evaluate_r5, both pass)**: V11 re-verified at DFN (329.85 K, +37.9 mV); V12-HiP-final combines the V7 high-porosity loading profile with V11 refinements: T_max 329.75 K, anode +41.7 mV, acceptance 0.811 Ah, capacity 4.318 Ah -> **champion**.

Honest residual: at 4C CC charge the 4.2 V ceiling is reached after 0.811 Ah (effective ~4.6C vs the 5.0 Ah nominal); the no-plating criterion is judged on anode-potential evidence, and charging beyond the window was not simulated.
""".format(
    noms=F["noms"], cap1=F["cap1"], capmin=F["capmin"], vv=F["vv"], hmm=F["hmm"], wm=F["wm"],
    thick=F["thick"], ce0=mk["c_e0"], kappa=F["kappa"], tp=F["tp"], De=F["De"],
    epsp=mk["eps_p"], epsn=mk["eps_n"], epssep=mk["eps_sep"],
    rhop=mk["rho_pos"], rhon=mk["rho_neg"], rhos=mk["rho_sep"],
    lpos=F["lpos"], lneg=F["lneg"], lsep=F["lsep"], rpos=F["rpos"], rneg=F["rneg"],
    ccp=F["ccp"], ccn=F["ccn"], np=F["np"], npb=F["npb"], npx=F["npx"],
    arealp=F["arealp"], arealn=F["arealn"], compp=F["compp"], compn=F["compn"],
    gelec=F["gelec"], hh=F["hh"], gpos=mk["g_pos_coat"], gneg=mk["g_neg_coat"],
    gsep=mk["g_sep"], gccp=mk["g_cc_p"], gccn=mk["g_cc_n"],
    gstack=F["gstack"], gtot=F["gtot"], ret_pct=F["ret_pct"], cap5=F["cap5"],
    t4s=F["t4s"], t5s=F["t5s"], amin_mv=F["amin_mv"], pd=F["pd"], pdmin=F["pdmin"],
    ed=F["ed"], retmin=F["retmin"],
)

dvpr = """# Design Verification Plan & Report - Power-Tool Battery (case t3_r3)

> VBF-T3R3-DVPR-01. Virtual test version: all results from simulation protocols. Uncovered conditions honestly written "N/A". No numbers from memory.

| Item | Condition | Result | Determination | Source |
|------|------|------|------|------|
| 1C discharge capacity | run-pyamm --protocol 1C_discharge (DFN) | {cap1} Ah | PASS vs >= {capmin} Ah | r5_v12_1c_dfn.json:capacity_ah |
| 5C discharge capacity retention | run-pyamm --protocol 5C_discharge (DFN) / 1C | {ret_pct} (4.2581 Ah / 4.3176 Ah) | PASS vs >= {retmin} | r5_v12_derived.json:retention_5c |
| 4C fast-charge temperature rise | run-pyamm --protocol 4C_charge_45C --thermal lumped (DFN) | T_max {t4s} | PASS vs <= 333.15 K | r5_v12_4c_dfn.json:T_max_K |
| 4C fast-charge lithium plating | same protocol + --plating; anode potential < 0 V | anode min {amin_mv} mV | PASS (no plating) | r5_v12_4c_dfn.json:anode_potential_v |
| 5C discharge temperature | run-pyamm --protocol 5C_discharge --thermal lumped (DFN) | T_max {t5s} | PASS vs <= 333.15 K (informational) | r5_v12_5c_dfn.json:T_max_K |
| Power density | calc-energy contract formula V_OC2/(4*DCR)/mass | {pd} W/kg | PASS vs >= {pdmin} W/kg | r5_v12_energy.json:power_density_w_kg |
| Voltage window | parameter set upper/lower cut-offs | {vv} V | PASS (design basis) | pybamm Chen2020 parameter set |
| Nail penetration | physical abuse test | - | N/A (beyond pure simulation boundary, requires physical experiment) | - |
| Overcharge to thermal runaway | physical safety test | - | N/A (beyond pure simulation boundary, requires physical experiment) | - |
| Crush / drop | mechanical abuse test | - | N/A (beyond pure simulation boundary, requires physical experiment) | - |
| Cycle life | aging protocol (SEI growth) | - | N/A (not simulated; no durability objective in task text) | - |
| Rate-pulse DC internal resistance | pulse protocol | - | N/A (not simulated; DCR estimate {dcr} mOhm at 1C mid-point for reference) | r5_v12_energy.json:dcr_ohm |

## Conclusion

All five entry-0 criteria verified PASS at DFN judge grade (round 5). Uncovered items: nail penetration, overcharge-to-runaway, crush/drop, cycle life, rate-pulse DCIR - all physical-experiment or aging-model territory, cited directly as paper limitations.
""".format(
    cap1=F["cap1"], capmin=F["capmin"], ret_pct=F["ret_pct"], retmin=F["retmin"],
    t4s=F["t4s"], t5s=F["t5s"], amin_mv=F["amin_mv"], pd=F["pd"],
    pdmin=F["pdmin"], vv=F["vv"], dcr=F["dcr"],
)

dfmea = """# Design FMEA - Power-Tool Battery (case t3_r3)

> VBF-T3R3-DFMEA-01. Qualitative version, based on simulation risk signals (annotated caliber). Severity/occurrence: three-level high/medium/low; basis = magnitude of simulation value vs threshold. No numbers from memory.

| Failure mode | Failure cause | Simulation signal (detection basis) | Severity | Occurrence | Design-side mitigation |
|------|------|------|------|------|------|
| Negative electrode lithium plating (4C fast charge) | anode polarization during high-rate charge | anode min {amin_mv} mV vs 0 V (margin engineered upward 21 -> 41.7 mV across rounds) | high | low | N/P x1.3 negative margin; 1.0 um fine anode particles; D 8e-10 m2/s + t+ 0.35 transport; margin monitoring kept as charge-window gate |
| Thermal excursion (exceeds 60 C) | resistive heating at 5C/4C rates | T_max {t4s} vs 333.15 K (3.4 K margin at 4C; 5C margin 13.6 K) | medium | low | h = {hh} W/m2/K thermal design requirement; temperature sensors + BMS charge/discharge derating |
| Electrolyte oxidative decomposition | high upper cut-off voltage vs electrolyte stability | voltage window {vv} V; real_compute=false: no DFT stability endorsement this case (skip recorded in endorse entry) | medium | low | NMC811 4.2 V practice from literature; additive strategy available as Stage-2 path (not exercised; architecture solved the criteria) |
| Insufficient capacity | loading reduction inherent to power architecture | 1C capacity {cap1} Ah vs >= {capmin} Ah (2.2x margin) | low | low | margin verified at DFN judge grade; no action |
| Reduced 4C charge acceptance (functional, not safety) | 4.2 V ceiling reached before full charge at 4C | CC acceptance {chg4} Ah before cutoff (effective ~4.6C vs 5 Ah nominal) | low | high | CV phase at ceiling; pack-level nominal sizing; documented as honest residual in final entry |

## Conclusion

Highest-risk item: plating (severity high) - mitigated in design (occurrence low) via N/P margin + electrode kinetics + transport; second: thermal (3.4 K margin at 4C). Mitigations are implemented in the champion design. Complete FMEA including process/supplier failure modes: N/A (beyond pure simulation boundary).
""".format(
    amin_mv=F["amin_mv"], t4s=F["t4s"], hh=F["hh"], vv=F["vv"],
    cap1=F["cap1"], capmin=F["capmin"], chg4=F["chg4"],
)

# ---------- delivery index ----------
code_rows = [
    ("DS", "Specification", "design_spec.md"),
    ("BOM", "Bill of Materials", "bom.xlsx"),
    ("DSH", "Datasheet technical parameter sheet", "datasheet.docx"),
    ("CALC", "Calculation sheet", "calc.xlsx"),
    ("DVPR", "Design verification report", "dvpr.md"),
    ("DFMEA", "Failure analysis", "dfmea.md"),
    ("CAD", "Structure model", "cell_model.stl (not generated: no 3D structure model requested this case)"),
]
file_rows = [
    ("design_spec.md", "VBF-T3R3-DS-01", "md", "Cell design specification generated per deliverable-design-spec format (build script, values from tool outputs)"),
    ("design_spec.pdf", "VBF-T3R3-DS-01", "pdf", "PDF release of design_spec.md, rendered with reportlab"),
    ("report.html", "VBF-T3R3-DS-02", "html", "Case report deterministically rendered from log.jsonl (bda render --case-dir runs/exp/t3_r3); belongs to DS"),
    ("bom.xlsx", "VBF-T3R3-BOM-01", "xlsx", "Bill of Materials, dual caliber g/cell and kg/kWh (openpyxl; formulas and sources per row)"),
    ("bom.pdf", "VBF-T3R3-BOM-01", "pdf", "PDF release of bom.xlsx, sheets rendered with reportlab"),
    ("datasheet.docx", "VBF-T3R3-DSH-01", "docx", "Customer-facing technical datasheet (python-docx; per-line sources)"),
    ("datasheet.pdf", "VBF-T3R3-DSH-01", "pdf", "PDF release of datasheet.docx, rendered with reportlab"),
    ("calc.xlsx", "VBF-T3R3-CALC-01", "xlsx", "Design calculation sheet: inputs -> capacity and energy -> energy density -> N/P and mass -> process parameters (openpyxl)"),
    ("calc.pdf", "VBF-T3R3-CALC-01", "pdf", "PDF release of calc.xlsx, sheets rendered with reportlab"),
    ("dvpr.md", "VBF-T3R3-DVPR-01", "md", "Design Verification Plan and Report, virtual test version (simulation protocols)"),
    ("dvpr.pdf", "VBF-T3R3-DVPR-01", "pdf", "PDF release of dvpr.md, rendered with reportlab"),
    ("dfmea.md", "VBF-T3R3-DFMEA-01", "md", "Design FMEA, qualitative version based on simulation risk signals"),
    ("dfmea.pdf", "VBF-T3R3-DFMEA-01", "pdf", "PDF release of dfmea.md, rendered with reportlab"),
    ("delivery_index.md", "VBF-T3R3-IDX-01", "md", "Delivery package index: cover + controlled file list (this file)"),
    ("delivery_index.pdf", "VBF-T3R3-IDX-01", "pdf", "PDF release of delivery_index.md, reportlab with blueprint colors"),
]

idx_lines = [
    "# Delivery Package Index - Power-Tool Battery (case t3_r3)",
    "",
    "> VBF-T3R3-IDX-01. Cover + controlled list of all deliverables generated for this case. The index registers only actually generated files.",
    "",
    "## Cover Information",
    "",
    "- **Case name**: t3_r3 - power-tool battery (nominal capacity >= 2 Ah, 5C retention >= 95%, 4C fast charge without plating, T_max <= 60 C, power density >= 4000 W/kg)",
    "- **Numbering scheme**: VBF-<CASE-ID-UPPERCASE>-<DOC-CODE>-<SEQ-NO>; case ID t3_r3 -> T3R3 (non-alphanumeric removed)",
    "- **Generation date**: 2026-08-26",
    "- **Signature block**: Prepared ______ / Reviewed ______ / Approved ______ (left blank for manual signing)",
    "",
    "## Document Code Reference (fixed by protocol)",
    "",
    "| Document code | Meaning | Corresponding file |",
    "|------|------|------|",
]
for code, meaning, cfile in code_rows:
    idx_lines.append(f"| {code} | {meaning} | {cfile} |")
idx_lines += [
    "",
    "## File List",
    "",
    "| File name | Number | Format | Source description |",
    "|------|------|------|------|",
]
for name, num, fmt, src in file_rows:
    idx_lines.append(f"| {name} | {num} | {fmt} | {src} |")
idx_lines.append("")
delivery_index = "\n".join(idx_lines)

for name, text in [
    ("design_spec.md", design_spec),
    ("dvpr.md", dvpr),
    ("dfmea.md", dfmea),
    ("delivery_index.md", delivery_index),
]:
    (DLV / name).write_text(text, encoding="utf-8")
print("markdown docs written")

# ================= PDF RENDER =================
import re
import xml.sax.saxutils as sx

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, PageBreak, HRFlowable)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

NAVY = "#14283C"
BLUE = "#1E5A8A"
COPPER = "#C97B3D"
ALT = "#F2F6FA"

font_dir = Path("C:/Windows/Fonts")
try:
    pdfmetrics.registerFont(TTFont("Body", str(font_dir / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("Body-Bold", str(font_dir / "DejaVuSans-Bold.ttf")))
    BODY, BODY_B = "Body", "Body-Bold"
except Exception:
    pdfmetrics.registerFont(TTFont("Body", str(font_dir / "arial.ttf")))
    pdfmetrics.registerFont(TTFont("Body-Bold", str(font_dir / "arialbd.ttf")))
    BODY, BODY_B = "Body", "Body-Bold"

ss = getSampleStyleSheet()
st_title = ParagraphStyle("T", parent=ss["Title"], fontName=BODY_B, fontSize=15,
                          textColor=colors.HexColor(NAVY), spaceAfter=6)
st_h2 = ParagraphStyle("H2", parent=ss["Heading2"], fontName=BODY_B, fontSize=11.5,
                       textColor=colors.HexColor(BLUE), spaceBefore=10, spaceAfter=4)
st_body = ParagraphStyle("B", parent=ss["BodyText"], fontName=BODY, fontSize=9,
                         leading=12.5, textColor=colors.HexColor("#1a1a1a"))
st_cell = ParagraphStyle("C", parent=ss["BodyText"], fontName=BODY, fontSize=8,
                         leading=10.5)
st_cell_h = ParagraphStyle("CH", parent=ss["BodyText"], fontName=BODY_B, fontSize=8.5,
                           leading=10.5, textColor=colors.white)
st_note = ParagraphStyle("N", parent=ss["BodyText"], fontName=BODY, fontSize=8.5,
                         leading=12, textColor=colors.HexColor("#444444"), spaceAfter=6)


def esc(s):
    return sx.escape(str(s))


def inline(s):
    s = esc(s).replace("**", "\x00")
    parts = s.split("\x00")
    out = []
    for i, p in enumerate(parts):
        if i % 2 == 1:
            out.append(f"<b>{p}</b>")
        else:
            p = re.sub(r"`([^`]*)`", r'<font face="Courier">\1</font>', p)
            out.append(p)
    return "".join(out)


def md_tables_to_flowables(md_text):
    flow = []
    table_buf = []
    prev = None
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if line.strip() == "":
            if table_buf:
                flow.append(md_table(table_buf))
                table_buf = []
            prev = "blank"
            continue
        if line.startswith(("#", "|", ">", "- ")):
            if table_buf and not line.startswith("|"):
                flow.append(md_table(table_buf))
                table_buf = []
        if line.startswith("# "):
            flow.append(Paragraph(inline(line[2:]), st_title))
        elif line.startswith("## "):
            flow.append(Paragraph(inline(line[3:]), st_h2))
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue  # separator row
            table_buf.append(cells)
        elif line.startswith("> "):
            flow.append(Paragraph(inline("<i>" + line[2:] + "</i>"), st_note))
        elif line.startswith("- "):
            flow.append(Paragraph(inline(line[2:]), st_body))
        else:
            flow.append(Paragraph(inline(line), st_body))
        prev = "text"
    if table_buf:
        flow.append(md_table(table_buf))
    return flow


def md_table(rows):
    data = [[Paragraph(inline(c), st_cell_h) for c in rows[0]]]
    for r in rows[1:]:
        data.append([Paragraph(inline(c), st_cell) for c in r])
    n = len(rows[0])
    widths = _col_widths(n)
    t = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C6D2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(ALT)))
    t.setStyle(TableStyle(style))
    return t


def _col_widths(n):
    avail = A4[0] - 3.2 * cm
    if n <= 2:
        return [avail * 0.32, avail * 0.68]
    if n == 3:
        return [avail * 0.20, avail * 0.40, avail * 0.40]
    if n == 4:
        return [avail * 0.22, avail * 0.18, avail * 0.12, avail * 0.46] if avail > 500 else [avail/4]*4
    if n == 5:
        return [avail * 0.24, avail * 0.22, avail * 0.16, avail * 0.18, avail * 0.20] if avail > 500 else [avail/5]*5
    if n == 6:
        return [avail * 0.16, avail * 0.14, avail * 0.16, avail * 0.14, avail * 0.16, avail * 0.24] if avail > 500 else [avail/6]*5 + [avail - avail/6*5]
    return [avail / n] * n


def build_pdf(out_path, story, title_header=True):
    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm,
                            title=Path(out_path).stem)
    doc.build(story)


# 1) markdown docs -> pdf
for md_name in ["design_spec.md", "dvpr.md", "dfmea.md"]:
    text = (DLV / md_name).read_text(encoding="utf-8")
    build_pdf(DLV / md_name.replace(".md", ".pdf"), md_tables_to_flowables(text))

# 2) xlsx -> pdf (bom, calc)
import openpyxl as px

for base in ["bom", "calc"]:
    wb = px.load_workbook(DLV / f"{base}.xlsx", data_only=True)
    story = []
    for wsn in wb.sheetnames:
        ws = wb[wsn]
        rows = [[c.value if c.value is not None else "" for c in row] for row in ws.iter_rows()]
        rows = [r for r in rows if any(str(v).strip() for v in r)]
        if not rows:
            continue
        story.append(Paragraph(inline(f"{base}.xlsx - sheet '{wsn}'"), st_h2))
        rows = [[str(v) for v in r] for r in rows]
        data = [[Paragraph(inline(c), st_cell_h) for c in rows[0]]]
        for r in rows[1:]:
            data.append([Paragraph(inline(c), st_cell) for c in r])
        inher_w = _col_widths(len(rows[0]))
        t = Table(data, colWidths=inher_w, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE)),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C6D2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for i in range(1, len(rows)):
            if i % 2 == 0:
                style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(ALT)))
        t.setStyle(TableStyle(style))
        story.append(t)
        story.append(Spacer(1, 14))
    build_pdf(DLV / f"{base}.pdf", story)

# 3) datasheet docx -> pdf
import docx

d = docx.Document(str(DLV / "datasheet.docx"))
story = []
for par in d.paragraphs:
    txt = par.text.strip()
    if not txt:
        continue
    if par.style.name.startswith(("Heading", "Title")):
        story.append(Paragraph(inline(txt), st_title))
    else:
        story.append(Paragraph(inline(txt), st_note))
for tbl in d.tables:
    rows = [[cell.text for cell in row.cells] for row in tbl.rows]
    data = [[Paragraph(inline(c), st_cell_h) for c in rows[0]]]
    for r in rows[1:]:
        data.append([Paragraph(inline(c), st_cell) for c in r])
    t = Table(data, colWidths=_col_widths(2), repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C6D2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(Spacer(1, 6))
    story.append(t)
build_pdf(DLV / "datasheet.pdf", story)

# 4) delivery_index.pdf with blueprint cover
st_cov_title = ParagraphStyle("CT", fontName=BODY_B, fontSize=19, leading=24,
                              textColor=colors.white, spaceAfter=2)
st_cov_sub = ParagraphStyle("CS", fontName=BODY, fontSize=11, leading=15,
                            textColor=colors.HexColor("#D8E4F0"))
st_cov_f = ParagraphStyle("CF", fontName=BODY, fontSize=10, leading=15,
                          textColor=colors.white)
st_sig = ParagraphStyle("SG", fontName=BODY, fontSize=9.5, leading=15,
                        textColor=colors.HexColor("#1a1a1a"))

cover_panel = Table(
    [[Paragraph("POWER-TOOL BATTERY CELL - DELIVERY PACKAGE", st_cov_title)],
     [Paragraph("Virtual Battery Factory / case t3_r3 / champion V12-HiP-final<br/>"
                "Nominal capacity >= 2 Ah / 5C retention >= 95% / 4C fast charge (no plating) / "
                "T_max <= 60 C / power density >= 4000 W/kg", st_cov_sub)],
     [Spacer(1, 6)]],
    colWidths=[A4[0] - 3.2 * cm])
cover_panel.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(NAVY)),
    ("TOPPADDING", (0, 0), (-1, -1), 14),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ("LEFTPADDING", (0, 0), (-1, -1), 16),
    ("RIGHTPADDING", (0, 0), (-1, -1), 16),
]))

sig_table = Table(
    [[Paragraph("Prepared: ______________________", st_sig), Paragraph("Reviewed: ______________________", st_sig), Paragraph("Approved: ______________________", st_sig)],
     [Paragraph("<i>(left blank for manual signing)</i>", st_note), Paragraph("", st_sig), Paragraph("", st_sig)]],
    colWidths=[(A4[0] - 3.2 * cm) / 3] * 3)
sig_table.setStyle(TableStyle([
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("TOPPADDING", (0, 0), (-1, -1), 4),
]))

cover_info = Table(
    [[Paragraph("<b>Case name</b>", st_sig), Paragraph("t3_r3 - power-tool battery", st_sig)],
     [Paragraph("<b>Numbering scheme</b>", st_sig), Paragraph("VBF-&lt;CASE-ID-UPPERCASE&gt;-&lt;DOC-CODE&gt;-&lt;SEQ-NO&gt; (case ID t3_r3 -> T3R3)", st_sig)],
     [Paragraph("<b>Generation date</b>", st_sig), Paragraph("2026-08-26", st_sig)]],
    colWidths=[3.2 * cm, A4[0] - 6.4 * cm])
cover_info.setStyle(TableStyle([
    ("TOPPADDING", (0, 0), (-1, -1), 3),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
]))

code_pdf_rows = [[Paragraph(c, st_cell_h) for c in ["Document code", "Meaning", "Corresponding file"]]
                 ] + [[Paragraph(c, st_cell) for c in [r[0], inline(r[1]), inline(r[2])]] for r in code_rows]
code_tbl = Table(code_pdf_rows, colWidths=_col_widths(3), repeatRows=1)
code_tbl.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE)),
    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C6D2")),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("TOPPADDING", (0, 0), (-1, -1), 3),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
] + [("BACKGROUND", (0, i), (-1, i), colors.HexColor(ALT)) for i in range(2, len(code_pdf_rows), 2)]))

file_pdf_rows = [[Paragraph(c, st_cell_h) for c in ["File name", "Number", "Format", "Source description"]]
                 ] + [[Paragraph(inline(r[0]), st_cell), Paragraph(r[1], st_cell),
                       Paragraph(r[2], st_cell), Paragraph(inline(r[3]), st_cell)] for r in file_rows]
avail = A4[0] - 3.2 * cm
file_tbl = Table(file_pdf_rows, colWidths=[avail * 0.20, avail * 0.17, avail * 0.08, avail * 0.55], repeatRows=1)
file_tbl.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE)),
    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C6D2")),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("TOPPADDING", (0, 0), (-1, -1), 3),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
] + [("BACKGROUND", (0, i), (-1, i), colors.HexColor(ALT)) for i in range(2, len(file_pdf_rows), 2)]))

idx_story = [
    Spacer(1, 8), cover_panel, Spacer(1, 10),
    HRFlowable(width="100%", thickness=1.6, color=colors.HexColor(COPPER)),
    Spacer(1, 10), cover_info, Spacer(1, 10), sig_table,
    PageBreak(),
    Paragraph("Document Code Reference (fixed by protocol)", st_h2),
    code_tbl, Spacer(1, 16),
    Paragraph("File List", st_h2),
    file_tbl,
]
build_pdf(DLV / "delivery_index.pdf", idx_story)

print("PDFs written")
for f in sorted(DLV.iterdir()):
    print(f"  {f.name}  {f.stat().st_size} B")