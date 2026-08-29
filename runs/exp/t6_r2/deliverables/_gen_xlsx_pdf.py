# -*- coding: utf-8 -*-
"""Deliverables generator: bom.xlsx, calc.xlsx, and 7 PDF releases (reportlab)."""
import io
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

D = Path(r"D:\research\degradation_prognostics\Battery_Design_Agent\runs\exp\t6_r2\deliverables")
D.mkdir(parents=True, exist_ok=True)

FONT = "Helvetica"
try:
    pdfmetrics.registerFont(TTFont("MSYaHei", r"C:\Windows\Fonts\msyh.ttc", subfontIndex=0))
    FONT = "MSYaHei"
except Exception:
    pass

# ---------------------------------------------------------------- BOM xlsx
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BOM"
hdr = ["#", "Component", "Specification / parameter", "Value", "Unit", "Source / note"]
ws.append(hdr)
rows = [
    [1, "Cathode (LNMO spinel)", "Positive electrode thickness", 7.475e-05, "m", "LNMO system file (Chen2020 base)"],
    [2, "Cathode", "Positive particle radius", 3.0e-06, "m", "r12_a_h140.json (DFN-optimized)"],
    [3, "Cathode", "Density", 4400, "kg/m3", "LNMO system file"],
    [4, "Cathode", "Porosity", 0.335, "-", "Chen2020 base"],
    [5, "Cathode", "Areal mass", 0.2212, "kg/m2", "cell/r12_a_energy_dfn.json:layer_kg_m2"],
    [6, "Anode (graphite)", "Negative electrode thickness", 1.2e-04, "m", "r12_a_h140.json (N/P margin)"],
    [7, "Anode", "Negative particle radius", 1.5e-06, "m", "r12_a_h140.json (4C kinetics)"],
    [8, "Anode", "Porosity", 0.25, "-", "Chen2020 base"],
    [9, "Anode", "Areal mass", 0.1491, "kg/m2", "cell/r12_a_energy_dfn.json:layer_kg_m2"],
    [10, "Separator", "Thickness", 2.5e-05, "m", "Chen2020 base (25 um)"],
    [11, "Separator", "Areal mass", 0.0025, "kg/m2", "cell/r12_a_energy_dfn.json:layer_kg_m2"],
    [12, "Electrolyte", "EC/EMC + LiPF6", "-", "-", "Chen2020 baseline formulation"],
    [13, "Al current collector", "Areal mass", 0.0432, "kg/m2", "cell/r12_a_energy_dfn.json:layer_kg_m2"],
    [14, "Cu current collector", "Areal mass", 0.1075, "kg/m2", "cell/r12_a_energy_dfn.json:layer_kg_m2"],
    [15, "SEI coating (dense FEC/VC-class)", "EC diffusivity", 5.0e-19, "m2/s", "r12_a_h140.json (SEI diffusion-limited lever)"],
    [16, "SEI coating", "SEI kinetic rate constant", 3.0e-13, "m/s", "r12_a_h140.json"],
    [17, "Thermal interface", "Cell cooling surface area", 0.0106, "m2", "r12_a_h140.json (double-sided pouch)"],
    [18, "Thermal interface", "Total heat transfer coefficient", 140.0, "W/m2.K", "r12_a_h140.json (vapor-chamber grade, HARD req.)"],
]
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
widths = [5, 28, 34, 14, 12, 52]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws2 = wb.create_sheet("Notes")
ws2.append(["Note"])
ws2.append(["All values tool-output sourced; see design_spec.md for caveats (4C current basis, LNMO file artifacts, true-compute skip)."])
bom_path = D / "bom.xlsx"
wb.save(bom_path)

# ---------------------------------------------------------------- CALC xlsx
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Criteria"
ws.append(["Criterion", "Threshold", "Result (DFN)", "Unit", "Margin", "Verdict", "Source file:key"])
ws.append(["Volumetric energy density", 950, 1096.0, "Wh/L", "+146.0", "PASS", "cell/r12_a_energy_dfn.json:energy_density_wh_l"])
ws.append(["Voltage plateau (midpoint)", 4.1, 4.1301, "V", "+0.030", "PASS", "cell/r12_a_energy_dfn.json:midpoint_voltage_v"])
ws.append(["Max temperature (4C, 45C amb.)", 323.15, 322.70, "K", "-0.45", "PASS", "cell/r12_a_4c45_dfn.json:T_max_K"])
ws.append(["Lithium plating", "False (min ap > 0)", "False (min +0.0382)", "V", "+0.0382", "PASS", "cell/r12_a_4c45_dfn.json:anode_potential_v"])
ws.append(["SEI after 100 cycles", 500, 415.6, "nm", "-84.4", "PASS", "cell/r12_a_aging.json:sei_thickness_nm_end"])
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
for i, w in enumerate([30, 22, 20, 10, 12, 10, 46], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws2 = wb.create_sheet("Derived")
ws2.append(["Quantity", "Formula", "Value", "Unit", "Source"])
ws2.append(["Energy", "from 1C DFN discharge integral", 26.518, "Wh", "cell/r12_a_energy_dfn.json:energy_wh"])
ws2.append(["Stack thickness", "sum of layer thicknesses", 235.6, "um", "cell/r12_a_energy_dfn.json:thickness_m"])
ws2.append(["Geometric area", "electrode height x width", 0.1027, "m2", "cell/r12_a_energy_dfn.json:area_m2"])
ws2.append(["Volume", "thickness x area", 2.4196e-05, "m3", "cell/r12_a_energy_dfn.json:volume_m3"])
ws2.append(["ED volumetric", "energy / volume (electrolyte excluded per contract)", 1096.0, "Wh/L", "cell/r12_a_energy_dfn.json:energy_density_wh_l"])
ws2.append(["Mass", "sum(layer areal mass x area)", 0.053772, "kg", "cell/r12_a_energy_dfn.json:mass_kg"])
ws2.append(["ED gravimetric", "energy / mass", 493.16, "Wh/kg", "cell/r12_a_energy_dfn.json:energy_density_wh_kg"])
ws2.append(["DCR", "(V(0) - V(10% t)) / I_1C", 0.006866, "ohm", "cell/r12_a_energy_dfn.json:dcr_ohm"])
ws2.append(["Peak power density", "V_OC^2 / (4 R_DC) / mass", 12254.5, "W/kg", "cell/r12_a_energy_dfn.json:power_density_w_kg"])
ws2.append(["SPMe cross-check ED", "same formulas, SPMe 1C", 1106.3, "Wh/L", "cell/r12_a_energy_spme.json:energy_density_wh_l"])
ws2.append(["SPMe cross-check midV", "same formulas, SPMe 1C", 4.2202, "V", "cell/r12_a_energy_spme.json:midpoint_voltage_v"])
for c in ws2[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
for i, w in enumerate([26, 40, 14, 10, 46], 1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
calc_path = D / "calc.xlsx"
wb.save(calc_path)

# ---------------------------------------------------------------- PDF generation
def _text_pdf(out_path: Path, title: str, md_path: Path):
    c = canvas.Canvas(str(out_path), pagesize=A4)
    w, h = A4
    y = h - 20 * mm
    c.setFont(FONT, 16)
    c.drawString(18 * mm, y, title)
    y -= 8 * mm
    c.setFont(FONT, 8.5)
    lines = md_path.read_text(encoding="utf-8-sig").splitlines()
    for line in lines:
        if y < 14 * mm:
            c.showPage()
            y = h - 20 * mm
            c.setFont(FONT, 8.5)
        if line.startswith("# "):
            c.setFont(FONT, 12)
            c.drawString(18 * mm, y, line[2:].strip())
            c.setFont(FONT, 8.5)
        elif line.startswith("## ") or line.startswith("### "):
            c.setFont(FONT, 10)
            c.drawString(18 * mm, y, line.strip("# ").strip())
            c.setFont(FONT, 8.5)
        elif line.startswith("|") and line.strip("| -") == "":
            y += 2 * mm
            continue
        elif line.startswith("|"):
            cells = [x.strip() for x in line.strip("|").split("|")]
            x = 12 * mm
            for cell in cells:
                c.drawString(x, y, cell[:44])
                x += 78 * mm
        elif line.strip():
            c.drawString(18 * mm, y, line[:170])
        y -= 3.6 * mm
    c.showPage()
    c.save()

def _xlsx_pdf(out_path: Path, title: str, xlsx_path: Path):
    wb = openpyxl.load_workbook(xlsx_path)
    c = canvas.Canvas(str(out_path), pagesize=A4)
    w, h = A4
    c.setFont(FONT, 16)
    c.drawString(18 * mm, h - 20 * mm, title)
    c.setFont(FONT, 7)
    y = h - 28 * mm
    for sheet in wb.worksheets:
        c.setFont(FONT, 11)
        c.drawString(18 * mm, y, f"Sheet: {sheet.title}")
        c.setFont(FONT, 7)
        y -= 5 * mm
        for row in sheet.iter_rows(values_only=True):
            if y < 14 * mm:
                c.showPage()
                y = h - 20 * mm
                c.setFont(FONT, 7)
            vals = [str(v)[:38] if v is not None else "" for v in row]
            if row == next(sheet.iter_rows(values_only=True)):
                c.setFont(FONT, 7)
            x = 12 * mm
            for v in vals:
                c.drawString(x, y, v)
                x += 52 * mm
            y -= 3.4 * mm
        y -= 6 * mm
    c.showPage()
    c.save()

docs = [
    ("design_spec.md", "Design Specification — VBF-T6R2-DS-001"),
    ("datasheet.md", "Datasheet — VBF-T6R2-DSH-001"),
    ("dvpr.md", "Design Verification Plan & Report — VBF-T6R2-DVPR-001"),
    ("dfmea.md", "Design FMEA — VBF-T6R2-DFMEA-001"),
    ("delivery_index.md", "Delivery Index — VBF-T6R2-IDX-001"),
]
for name, title in docs:
    _text_pdf(D / (Path(name).stem + ".pdf"), title, D / name)
_xlsx_pdf(D / "bom.pdf", "Bill of Materials — VBF-T6R2-BOM-001", bom_path)
_xlsx_pdf(D / "calc.pdf", "Compliance Calculation Book — VBF-T6R2-CALC-001", calc_path)

for p in sorted(D.iterdir()):
    if p.suffix in (".pdf", ".xlsx"):
        print(p.name, p.stat().st_size, "bytes")
print("OK")
