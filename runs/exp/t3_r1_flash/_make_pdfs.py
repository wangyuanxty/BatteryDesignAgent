# -*- coding: utf-8 -*-
"""Generate PDF releases for all deliverable categories (md -> pdf, xlsx -> pdf)."""
import os, re
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, Preformatted)
from openpyxl import load_workbook

WS = r"D:\research\degradation_prognostics\Battery_Design_Agent\runs\exp\t3_r1_flash"
OUT = os.path.join(WS, "deliverables")

styles = getSampleStyleSheet()
H1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=15, spaceBefore=10, spaceAfter=6)
H2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12.5, spaceBefore=8, spaceAfter=4)
BODY = ParagraphStyle("BODY", parent=styles["BodyText"], fontSize=9, leading=12.5, spaceAfter=4)
CELL = ParagraphStyle("CELL", parent=BODY, fontSize=8, leading=10, spaceAfter=0)
CELLB = ParagraphStyle("CELLB", parent=CELL, fontName="Helvetica-Bold")
TITLE = ParagraphStyle("TITLE", parent=styles["Title"], fontSize=16, spaceAfter=8)

def esc(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", t)
    return t

def md_to_pdf(md_path, pdf_path):
    doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                            leftMargin=16*mm, rightMargin=16*mm,
                            topMargin=14*mm, bottomMargin=14*mm,
                            title=os.path.basename(pdf_path))
    story = []
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    i, n = 0, len(lines)
    while i < n:
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# ") and not ln.startswith("## "):
            story.append(Paragraph(esc(ln[2:]), TITLE)); i += 1; continue
        if ln.startswith("## "):
            story.append(Paragraph(esc(ln[3:]), H2)); i += 1; continue
        if ln.startswith("### "):
            story.append(Paragraph(esc(ln[4:]), H1)); i += 1; continue
        if ln.startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i+1]):
            header = [c.strip() for c in ln.strip().strip("|").split("|")]
            data = []
            j = i + 2
            while j < n and lines[j].startswith("|"):
                data.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            tbl = [[Paragraph(esc(c), CELLB) for c in header]] + \
                  [[Paragraph(esc(c), CELL) for c in row] for row in data]
            t = Table(tbl, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF4FA")]),
            ]))
            story.append(t)
            story.append(Spacer(1, 5))
            i = j
            continue
        if ln.startswith("- "):
            story.append(Paragraph("• " + esc(ln[2:]), BODY)); i += 1; continue
        story.append(Paragraph(esc(ln), BODY)); i += 1
    doc.build(story)

def xlsx_to_pdf(xlsx_path, pdf_path, title):
    wb = load_workbook(xlsx_path)
    doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                            leftMargin=14*mm, rightMargin=14*mm,
                            topMargin=14*mm, bottomMargin=14*mm,
                            title=title)
    story = [Paragraph(esc(title), TITLE)]
    for ws in wb.worksheets:
        story.append(Paragraph(esc(ws.title), H1))
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            story.append(Paragraph("(empty sheet)", BODY)); continue
        tbl = [[Paragraph(esc(str(c)) if c is not None else "", CELLB if r == 0 else CELL)
                for c in row] for r, row in enumerate(rows)]
        t = Table(tbl, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(t)
        story.append(Spacer(1, 8))
    doc.build(story)

md_to_pdf(os.path.join(OUT, "design_spec.md"), os.path.join(OUT, "design_spec.pdf"))
md_to_pdf(os.path.join(OUT, "datasheet.md"), os.path.join(OUT, "datasheet.pdf"))
md_to_pdf(os.path.join(OUT, "dvpr.md"), os.path.join(OUT, "dvpr.pdf"))
md_to_pdf(os.path.join(OUT, "dfmea.md"), os.path.join(OUT, "dfmea.pdf"))
md_to_pdf(os.path.join(OUT, "delivery_index.md"), os.path.join(OUT, "delivery_index.pdf"))
xlsx_to_pdf(os.path.join(OUT, "bom.xlsx"), os.path.join(OUT, "bom.pdf"),
            "Bill of Materials — VBF Power-Tool Cell (V4)")
xlsx_to_pdf(os.path.join(OUT, "calc.xlsx"), os.path.join(OUT, "calc.pdf"),
            "Calculation Workbook — VBF Power-Tool Cell (V4)")
print("PDFs written")
