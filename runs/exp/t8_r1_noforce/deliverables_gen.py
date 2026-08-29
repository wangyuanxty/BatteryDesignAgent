# -*- coding: utf-8 -*-
"""Generate BOM/calc workbooks + PDF releases + delivery index for t8_r1_noforce.
All values are mechanical copies of bda tool outputs (r6_v9_*.json); see per-row source columns."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable, PageBreak)
from reportlab.pdfgen import canvas as rl_canvas

BASE = 'runs/exp/t8_r1_noforce'
DLV = os.path.join(BASE, 'deliverables')
os.makedirs(DLV, exist_ok=True)

# ---------------- shared mechanical values (tool outputs) ----------------
AREA = 0.122915          # r6_v9_energy.json area_m2
E_WH = 19.4147           # r6_v9_energy.json energy_wh
E_KWH = E_WH / 1000.0
LAYERS = [
    ('Positive electrode active material (NMC811)', 0.147964),
    ('Negative electrode active material (graphite)', 0.091765),
    ('Positive current collector (Al, 8 um)', 0.0216),
    ('Negative current collector (Cu, 6 um)', 0.05376),
    ('Separator (10 um)', 0.002104),
]
TOTAL_KG = 0.0389878      # r6_v9_energy.json mass_kg
ELEC_G = 9.552            # pore vol 7.960e-6 m3 x 1200 kg/m3 (literature density)

def g_per_cell(kg_m2):
    return kg_m2 * AREA * 1000.0

def kg_kwh(kg):
    return kg / E_KWH

# BOM row data: (component, g/cell, contract?, note)
POS_G = g_per_cell(LAYERS[0][1])
NEG_G = g_per_cell(LAYERS[1][1])
AL_G  = g_per_cell(LAYERS[2][1])
CU_G  = g_per_cell(LAYERS[3][1])
SEP_G = g_per_cell(LAYERS[4][1])
ADD_POS_B = POS_G * 0.02   # PVDF 2 wt% literature default
ADD_POS_C = POS_G * 0.02   # conductive carbon 2 wt% literature default
ADD_NEG_B = NEG_G * 0.02   # CMC+SBR 2 wt% literature default
ADD_NEG_C = NEG_G * 0.01   # conductive carbon 1 wt% literature default

BOM_MAIN = [
    ('Positive active material (NMC811)', POS_G, True,
     'r6_v9_energy.json layer_kg_m2 x area'),
    ('Negative active material (graphite)', NEG_G, True,
     'r6_v9_energy.json layer_kg_m2 x area'),
    ('Positive current collector (Al, 8 um)', AL_G, True,
     'r6_v9_energy.json layer_kg_m2 x area'),
    ('Negative current collector (Cu, 6 um)', CU_G, True,
     'r6_v9_energy.json layer_kg_m2 x area'),
    ('Separator (10 um)', SEP_G, True,
     'r6_v9_energy.json layer_kg_m2 x area'),
]
BOM_SUPP = [
    ('Electrolyte (1 M LiPF6 EC:DMC)', ELEC_G,
     'pore volume 7.960e-6 m3 x 1200 kg/m3; density = literature value; fill factor 1.0'),
    ('Positive binder (PVDF)', ADD_POS_B, '2 wt% of active mass - literature default (excluded from contract mass)'),
    ('Positive conductive carbon', ADD_POS_C, '2 wt% of active mass - literature default (excluded from contract mass)'),
    ('Negative binder (CMC+SBR)', ADD_NEG_B, '2 wt% of active mass - literature default (excluded from contract mass)'),
    ('Negative conductive carbon', ADD_NEG_C, '1 wt% of active mass - literature default (excluded from contract mass)'),
    ('Enclosure, tabs, termination', 0.0, 'not modeled in this virtual case'),
]

# ---------------- xlsx: bom ----------------
def style_header(cell):
    cell.font = Font(bold=True, color='FFFFFF', size=10)
    cell.fill = PatternFill('solid', fgColor='1E5A8A')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = 'BOM'
ws['A1'] = 'Bill of Materials - VBF-T8R1NOFORCE-BOM-01 (cell R6-V9, t8_r1_noforce)'
ws['A1'].font = Font(bold=True, size=12, color='14283C')
ws['A2'] = 'Dual caliber: g/cell and kg/kWh. Energy basis: 19.4147 Wh (1C contract, r6_v9_energy.json). Generated 2026-08-25.'
ws['A2'].font = Font(italic=True, size=9)
hdrs = ['Component', 'Mass (g/cell)', 'Mass (kg/kWh)', 'In contract mass', 'Source / basis']
for j, h in enumerate(hdrs, 1):
    c = ws.cell(row=4, column=j, value=h)
    style_header(c)
rows = []
for name, g, in_contract, src in BOM_MAIN:
    rows.append((name, round(g, 4), round(kg_kwh(g / 1000.0), 4),
                 'YES' if in_contract else 'no', src))
rows.append(('CONTRACT TOTAL (electrolyte excluded, calc-energy caliber)', round(TOTAL_KG * 1000, 4),
             round(kg_kwh(TOTAL_KG), 4), 'YES', 'r6_v9_energy.json mass_kg = 0.0389878 kg'))
for name, g, src in BOM_SUPP:
    rows.append((name, round(g, 4), round(kg_kwh(g / 1000.0), 4), 'no', src))
grand = TOTAL_KG * 1000 + ELEC_G + ADD_POS_B + ADD_POS_C + ADD_NEG_B + ADD_NEG_C
rows.append(('GRAND TOTAL (incl. electrolyte + binder/additive)', round(grand, 4),
             round(kg_kwh(grand / 1000.0), 4), 'no', 'annotated total'))
r = 5
for row in rows:
    for j, v in enumerate(row, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.alignment = Alignment(vertical='center', wrap_text=(j in (1, 5)))
        c.font = Font(size=9)
        if row[0].startswith(('CONTRACT', 'GRAND')):
            c.font = Font(size=9, bold=True)
    r += 1
for j, w in enumerate([48, 15, 15, 16, 60], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'A5'
bom_path = os.path.join(DLV, 'bom.xlsx')
wb.save(bom_path)

# ---------------- xlsx: calc ----------------
wb2 = Workbook()
NAVY, BLUE = '14283C', '1E5A8A'

def sheet_table(wsx, title, subtitle, hdrs, rows_, widths):
    wsx['A1'] = title
    wsx['A1'].font = Font(bold=True, size=12, color=NAVY)
    wsx['A2'] = subtitle
    wsx['A2'].font = Font(italic=True, size=9)
    for j, h in enumerate(hdrs, 1):
        c = wsx.cell(row=4, column=j, value=h)
        style_header(c)
    for i, row in enumerate(rows_):
        for j, v in enumerate(row, 1):
            c = wsx.cell(row=5 + i, column=j, value=v)
            c.font = Font(size=9)
            c.alignment = Alignment(vertical='center', wrap_text=True)
    for j, w in enumerate(widths, 1):
        wsx.column_dimensions[get_column_letter(j)].width = w
    wsx.freeze_panes = 'A5'

# Sheet 1 Energy
ws = wb2.active
ws.title = 'Energy'
sheet_table(ws,
    'Energy calculation - contract caliber (VBF-T8R1NOFORCE-CALC-01)',
    'Source: bda calc-energy on R6-V9 params (r6_v9_energy.json). Electrolyte excluded.',
    ['Metric', 'Value', 'Unit', 'Source'],
    [
        ('Discharge capacity (1C)', 5.2725, 'Ah', 'r6_v9_1c_dfn.json capacity_ah'),
        ('Energy (1C contract)', 19.4147, 'Wh', 'r6_v9_energy.json energy_wh'),
        ('Gravimetric energy density', 497.968, 'Wh/kg', 'r6_v9_energy.json energy_density_wh_kg'),
        ('Volumetric energy density', 854.72, 'Wh/L', 'r6_v9_energy.json energy_density_vol_wh_l'),
        ('Cell mass (contract)', 0.0389878, 'kg', 'r6_v9_energy.json mass_kg'),
        ('Electrode area', 0.122915, 'm2', 'r6_v9_energy.json area_m2'),
        ('Stack thickness', 1.848e-4, 'm', 'r6_v9_energy.json thickness_m'),
        ('Discharge midpoint voltage', 3.9335, 'V', 'r6_v9_energy.json midpoint_voltage_v'),
        ('DC resistance (formulation)', 0.0007275, 'ohm', 'r6_v9_energy.json dcr_ohm'),
        ('Power density (formulation)', 150480.7, 'W/kg', 'r6_v9_energy.json power_density_w_kg'),
    ],
    [36, 16, 12, 44])
# Sheet 2 Mass
ws = wb2.create_sheet('Mass')
mrows = [(n, kgm2, round(g_per_cell(kgm2), 4)) for n, kgm2 in LAYERS]
mrows.append(('TOTAL (contract)', round(sum(x[1] for x in LAYERS), 6), round(TOTAL_KG * 1000, 4)))
mrows.append(('Electrolyte (supplementary, excluded)', '-', ELEC_G))
sheet_table(ws,
    'Mass budget (VBF-T8R1NOFORCE-CALC-01)',
    'g/cell = layer_kg_m2 x area (0.122915 m2). Contract = electrolyte excluded.',
    ['Layer', 'kg/m2', 'g/cell'],
    mrows, [44, 16, 14])
# Sheet 3 Retention
ws = wb2.create_sheet('Retention')
sheet_table(ws,
    '5C capacity retention (VBF-T8R1NOFORCE-CALC-01)',
    'Retention = Q(5C discharge DFN) / Q(1C discharge DFN), same parameters (mechanical, r6_v9_retention.json).',
    ['Quantity', 'Value', 'Unit', 'Source'],
    [
        ('1C discharge capacity (25 C)', 5.2725, 'Ah', 'r6_v9_1c_dfn.json'),
        ('5C discharge capacity (25 C)', 5.2178, 'Ah', 'r6_v9_5c_dfn.json'),
        ('Capacity retention @5C', 0.989615, '-', 'r6_v9_retention.json capacity_retention_5c'),
        ('Criterion (entry 0)', 0.90, '-', 'log.jsonl entry 0 stage2.capacity_retention_5c min'),
        ('Verdict', 'PASS', '-', '0.9896 >= 0.90'),
        ('5C discharge T_max', 313.64, 'K', 'r6_v9_5c_dfn.json (<= 333.15 K)'),
    ],
    [38, 16, 12, 46])
# Sheet 4 Safety
ws = wb2.create_sheet('Safety')
sheet_table(ws,
    'Safety exam - fast charge (VBF-T8R1NOFORCE-CALC-01)',
    '4C charge protocol at 45 C ambient, lumped thermal + plating module (r6_v9_4c_charge45_plating.json).',
    ['Quantity', 'Value', 'Unit', 'Determination'],
    [
        ('Anode potential minimum (4C charge)', 0.0063475, 'V', 'plated = false -> PASS (margin +6.3 mV)'),
        ('Lithium plating flag', 'false', '-', 'PASS (criterion: plated = false)'),
        ('T_max during 4C charge', 326.81, 'K', 'PASS (<= 333.15 K red line)'),
        ('T_max during 5C discharge', 313.64, 'K', 'PASS (<= 333.15 K red line)'),
        ('Maximum temperature red line (entry 0)', 333.15, 'K', 'protocol default (task silent)'),
    ],
    [40, 16, 12, 48])
# Sheet 5 Notes
ws = wb2.create_sheet('Notes')
notes = [
    ('Assumption', 'Note'),
    ('Electrolyte excluded from contract mass',
     'calc-energy contract caliber excludes electrolyte; parameter set carries no electrolyte density (electrolyte_included: false).'),
    ('Electrolyte transport parameters',
     'kappa = 1.4 S/m, D = 3.7e-10 m2/s are literature values (Nyman 2008 transport parameterization) set as formulation design levers, recorded in R6 propose.'),
    ('Electrolyte density 1.2 g/cm3',
     'Literature value used for BOM electrolyte row and datasheet mass incl. electrolyte (48.54 g).'),
    ('Binder/additive fractions',
     '2 wt% PVDF / 2 wt% carbon (positive), 2 wt% CMC+SBR / 1 wt% carbon (negative) - literature defaults, excluded from contract mass.'),
    ('N/P ratio 0.641',
     'Capacity-density basis, mechanical from parameter set (neg 49.18 Ah/m2 vs pos 76.72 Ah/m2). Cell is negative-limited - deliberate transport trade.'),
    ('Formation recommendation', '0.1C CC charge to 4.2 V, 25 C, 2 cycles (design recommended; production tuning required).'),
    ('Not simulated', 'Cycle life, low-temperature retention, abuse (nail/crush/overcharge), self-discharge - recorded as N/A in DVPR.'),
    ('V8 solver failure', 'R5 V8 5C discharge IDA_CONV_FAIL recorded verbatim in log.jsonl; superseded by V9 design.'),
]
for i, row in enumerate(notes):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=1 + i, column=j, value=v)
        c.alignment = Alignment(vertical='top', wrap_text=True)
        c.font = Font(size=9, bold=(i == 0), color='FFFFFF' if i == 0 else '000000')
        if i == 0:
            c.fill = PatternFill('solid', fgColor=BLUE)
ws.column_dimensions['A'].width = 40
ws.column_dimensions['B'].width = 110
calc_path = os.path.join(DLV, 'calc.xlsx')
wb2.save(calc_path)
print('xlsx done:', bom_path, calc_path)

# ---------------- PDF releases ----------------
NAVY_C, BLUE_C, ORANGE_C = colors.HexColor('#14283C'), colors.HexColor('#1E5A8A'), colors.HexColor('#C97B3D')
H1 = ParagraphStyle('H1', fontName='Helvetica-Bold', fontSize=15, textColor=NAVY_C, spaceAfter=4, leading=18)
H2 = ParagraphStyle('H2', fontName='Helvetica-Bold', fontSize=11, textColor=BLUE_C, spaceBefore=10, spaceAfter=4, leading=14)
META = ParagraphStyle('META', fontName='Helvetica', fontSize=8.5, textColor=colors.HexColor('#555555'), spaceAfter=8)
CELL = ParagraphStyle('CELL', fontName='Helvetica', fontSize=7.5, leading=9.5)
CELLB = ParagraphStyle('CELLB', fontName='Helvetica-Bold', fontSize=7.5, leading=9.5)
NOTE = ParagraphStyle('NOTE', fontName='Helvetica-Oblique', fontSize=8, textColor=colors.HexColor('#333333'), spaceBefore=6, spaceAfter=4, leading=11)
BODY = ParagraphStyle('BODY', fontName='Helvetica', fontSize=9, leading=12.5, spaceAfter=4)

class CountingCanvas(rl_canvas.Canvas):
    def __init__(self, *a, **k):
        rl_canvas.Canvas.__init__(self, *a, **k)
        self._pages = 0
    def showPage(self):
        self._pages += 1
        rl_canvas.Canvas.showPage(self)

def esc(s):
    return str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def build_pdf(fname, story, title):
    path = os.path.join(DLV, fname)
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title=title, author='VBF Design Agent',
                            canvasmaker=CountingCanvas)
    def footer(cnv, doc_):
        cnv.saveState()
        cnv.setFont('Helvetica', 7)
        cnv.setFillColor(colors.HexColor('#888888'))
        cnv.drawString(2 * cm, 1.1 * cm, 'Virtual Battery Factory - case t8_r1_noforce - simulation-based virtual design')
        cnv.drawRightString(A4[0] - 2 * cm, 1.1 * cm, 'Page %d' % cnv._pageNumber)
        cnv.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path, doc.page

def head(number, title, subtitle):
    return [Paragraph(esc(title), H1),
            Paragraph('%s  |  %s' % (esc(number), esc(subtitle)), META),
            HRFlowable(width='100%', thickness=1.2, color=ORANGE_C), Spacer(1, 6)]

def md_table(hdr, rows, widths=None, font=CELL, header_bg=BLUE_C):
    usable = A4[0] - 4 * cm
    if widths is None:
        widths = [usable / len(hdr)] * len(hdr)
    data = [[Paragraph('<b>%s</b>' % esc(h), font) for h in hdr]]
    for row in rows:
        data.append([Paragraph(esc(v), font) for v in row])
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), header_bg),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#BBBBBB')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F6FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    return t

U = A4[0] - 4 * cm
W = [U * f for f in [0.34, 0.30, 0.36]]

def ds_story():
    s = head('VBF-T8R1NOFORCE-DS-01', 'Cell Design Specification - Long-Endurance Drone Battery',
             'case t8_r1_noforce | 2026-08-25 | virtual design (simulation-based)')
    s.append(Paragraph('1. Basic specification', H2))
    s.append(md_table(['Item', 'Value', 'Source'], [
        ['Electrochemical system', 'NMC811 / graphite (Chen2020 parameterization)', 'parameter set (base selection per anchor table)'],
        ['Nominal capacity', '5.0 Ah (set) / 5.2725 Ah (simulated 1C)', 'r6_v9_1c_dfn.json capacity_ah'],
        ['Voltage window', '2.5 - 4.2 V', 'parameter set'],
        ['Electrode dimensions', '65 mm height x 1891 mm width x 0.1848 mm stack', 'parameter set + r6_v9_energy.json'],
        ['Electrolyte', '1 M LiPF6 EC:DMC 1:1; kappa 1.4 S/m, D 3.7e-10 m2/s (literature, Nyman 2008); t+ 0.2594 (set)', 'r6_v9_kin_el.json + set'],
        ['Cooling requirement', 'h = 60 W/m2/K forced-air (ducted prop-wash)', 'r6_v9_kin_el.json'],
        ['Shell/casing thickness', 'not provided (no parameter)', '-'],
    ], W))
    s.append(Paragraph('2. Electrode and separator', H2))
    s.append(md_table(['Layer', 'Thickness (um)', 'Porosity', 'AM vol. frac.', 'Particle radius (um)'], [
        ['Positive (NMC811)', '75.6', '0.40', '0.60', '0.7 (nano-NMC)'],
        ['Negative (graphite)', '85.2', '0.35', '0.65', '1.2 (fine graphite)'],
        ['Separator', '10', '0.47', '-', '-'],
        ['Positive current collector (Al)', '8', '-', '-', '-'],
        ['Negative current collector (Cu)', '6', '-', '-', '-'],
    ]))
    s.append(Paragraph('N/P = (33133 x 0.65 x 85.2e-6) / (63104 x 0.60 x 75.6e-6) = 0.641 (capacity-density basis, mechanical). '
                       'The cell is negative-limited - a deliberate transport trade (see design notes).', NOTE))
    s.append(Paragraph('3. Process design parameters', H2))
    s.append(md_table(['Parameter', 'Formula', 'Value'], [
        ['Positive areal density', 't x (1-e) x rho', '75.6e-6 x 0.60 x 3262 = 147.96 g/m2'],
        ['Negative areal density', 't x (1-e) x rho', '85.2e-6 x 0.65 x 1657 = 91.76 g/m2'],
        ['Positive compaction density', 'rho x (1-e)', '3262 x 0.60 = 1.96 g/cm3'],
        ['Negative compaction density', 'rho x (1-e)', '1657 x 0.65 = 1.08 g/cm3'],
        ['Electrolyte fill amount', 'pore vol x density x fill', '7.960e-6 m3 x 1200 kg/m3 x 1.0 = 9.55 g (density literature)'],
        ['Formation (recommended)', '-', '0.1C CC to 4.2 V, 25 C, 2 cycles (design rec.; production tuning required)'],
    ]))
    s.append(Paragraph('4. Mass breakdown (contract caliber; electrolyte excluded per calc-energy)', H2))
    s.append(md_table(['Layer', 'kg/m2', 'g/cell (area 0.122915 m2)'], [
        ['Positive electrode (active)', '0.14796', '18.19'],
        ['Negative electrode (active)', '0.09176', '11.28'],
        ['Positive current collector (Al 8 um)', '0.02160', '2.65'],
        ['Negative current collector (Cu 6 um)', '0.05376', '6.61'],
        ['Separator (10 um)', '0.00210', '0.26'],
        ['TOTAL (contract)', '0.31718', '38.99 g  (criterion: <= 40 g PASS)'],
        ['Electrolyte (excluded)', '-', '9.55 g (BOM row, literature density)'],
    ]))
    s.append(Paragraph('5. Performance verification (entry-0 criteria)', H2))
    s.append(md_table(['Metric', 'Result', 'Criterion', 'Verdict'], [
        ['Energy density', '497.97 Wh/kg', '>= 446.18 Wh/kg', 'PASS'],
        ['5C capacity retention', '0.9896', '>= 0.90', 'PASS'],
        ['Cell mass', '38.99 g', '<= 40 g', 'PASS'],
        ['Lithium plating (4C charge, 45 C)', 'anode min +0.00635 V', 'plated = false', 'PASS (margin +6.3 mV)'],
        ['T_max 4C charge', '326.81 K', '<= 333.15 K', 'PASS'],
        ['T_max 5C discharge', '313.64 K', '<= 333.15 K', 'PASS'],
        ['1C capacity', '5.2725 Ah', 'no threshold (feeds ED)', 'informational'],
    ]))
    s.append(Paragraph('6. Design notes (change log, cited from log.jsonl)', H2))
    for line in [
        'R1 baseline (Chen2020 defaults): ED 400.75, 5C retention 0.087, mass 43.45 g - all three fail; diagnosis: solid-diffusion-limited at 5C (tau ~ r^2) + inert mass 0.153 kg/m2.',
        'R2 attribution: inert lean-out solves ED/mass but not rate; particles 0.2325; porosity 0.1998.',
        'R3: rate levers combined (r_p 0.7 um, r_n 2 um, porosity 0.40/0.35, high-kappa electrolyte) -> retention 0.9524, ED 493.6, mass 33.0 g.',
        'R4: sizing (width 1.891 m -> 39.0 g) + thermal (h=40): stage2 pass, stage3 fail - plated (anode -0.030 V), 4C T_max 334.14 K.',
        'R5 (V8 thicker anode, N/P 1.13): backfired - anode min -0.194 V (plating dip is charge-end electrolyte transport polarization, not capacity) + 5C IDA_CONV_FAIL recorded verbatim.',
        'R6 (final V9): anode kinetics r_n 1.2 um + EC:DMC electrolyte (kappa 1.4, D 3.7e-10) + h=60 -> anode min +0.00635 V, 4C T_max 326.81 K, retention 0.9896, ED 497.97, mass 38.99 g. All five criteria PASS.',
    ]:
        s.append(Paragraph('* %s' % esc(line), BODY))
    return s

def dsh_story():
    s = head('VBF-T8R1NOFORCE-DSH-01', 'Cell Datasheet - Long-Endurance Drone Battery',
             'case t8_r1_noforce | 2026-08-25 | all values from bda tool output; not a physical cell')
    s.append(Paragraph('General', H2))
    s.append(md_table(['Field', 'Value', 'Source'], [
        ['Cell designation', 'VBF-T8R1NOFORCE-CELL-R6V9', 'delivery index'],
        ['Electrochemical system', 'NMC811 / graphite (pouch geometry)', 'Chen2020 parameter set'],
        ['Rated capacity', '5.2725 Ah (1C discharge, 25 C, DFN)', 'r6_v9_1c_dfn.json'],
        ['Nominal voltage', '3.933 V (1C midpoint)', 'r6_v9_energy.json'],
        ['Voltage window', '2.5 - 4.2 V', 'parameter set'],
        ['Rated energy', '19.4147 Wh', 'r6_v9_energy.json'],
        ['Energy density (gravimetric)', '497.968 Wh/kg (contract; electrolyte excluded)', 'r6_v9_energy.json'],
        ['Energy density (volumetric)', '854.72 Wh/L', 'r6_v9_energy.json'],
        ['DC resistance (formulation)', '0.7275 mohm', 'r6_v9_energy.json dcr_ohm'],
        ['Power density (formulation)', '150.5 kW/kg', 'r6_v9_energy.json'],
    ], W))
    s.append(Paragraph('Physical', H2))
    s.append(md_table(['Field', 'Value', 'Source'], [
        ['Electrode dimensions', '65 mm height x 1891 mm width; area 0.122915 m2', 'parameter set / r6_v9_energy.json'],
        ['Stack thickness', '184.8 um (pos 75.6 + sep 10 + neg 85.2 + Al 8 + Cu 6)', 'r6_v9_energy.json'],
        ['Cell mass (contract)', '38.99 g (electrolyte excluded per calc-energy)', 'r6_v9_energy.json mass_kg'],
        ['Cell mass incl. electrolyte', '48.54 g (electrolyte 9.55 g, literature density)', 'BOM sheet'],
    ], W))
    s.append(Paragraph('Electrical performance (simulated, DFN)', H2))
    s.append(md_table(['Condition', 'Result', 'Source'], [
        ['1C discharge, 25 C', '5.2725 Ah; T_max 299.18 K', 'r6_v9_1c_dfn.json'],
        ['5C discharge, 25 C', '5.2178 Ah; T_max 313.64 K; retention vs 1C = 0.9896', 'r6_v9_5c_dfn.json / r6_v9_retention.json'],
        ['Continuous discharge rating', '5C (requirement retention >= 0.90 - met at 0.9896)', 'r6_v9_retention.json'],
        ['Fast charge capability', '4C at 45 C: anode min +0.00635 V -> no plating; T_max 326.81 K', 'r6_v9_4c_charge45_plating.json'],
    ]))
    s.append(Paragraph('Environmental and safety', H2))
    s.append(md_table(['Field', 'Value', 'Source'], [
        ['Operating temperature (simulated)', '25 - 45 C (1C/5C at 25 C; 4C charge at 45 C)', 'protocol runs, honest scope'],
        ['Cooling requirement', 'h = 60 W/m2/K forced-air (ducted prop-wash)', 'r6_v9_kin_el.json'],
        ['Lithium plating (4C charge)', 'none - anode min +0.00635 V (plated = false); thin margin +6.3 mV', 'r6_v9_4c_charge45_plating.json'],
        ['Thermal runaway', 'not simulated (no abuse runs in this case)', 'honest N/A'],
        ['Cycle life', 'not simulated', 'honest N/A (aging model present; not a task criterion)'],
    ], W))
    s.append(Paragraph('Application note', H2))
    s.append(Paragraph(
        'Target: long-endurance drone. 5C discharge support (retention 0.9896) covers launch/climb bursts; energy density '
        '497.97 Wh/kg (contract) with 38.99 g cell mass supports endurance-optimized packs. The design is negative-limited '
        '(N/P 0.641 capacity-density basis) - a deliberate trade for transport: the 85.2 um graphite electrode keeps '
        'electrolyte-path polarization low at charge end, protecting the plating margin. Thin 6-8 um current collectors and '
        '10 um separator carry the energy-density budget - aggressive engineering choices requiring pouch-fabrication capability.', BODY))
    s.append(Paragraph(
        'Simulation honesty note: all values are DFN simulation outputs on the Chen2020 parameterization with design-parameter '
        'overrides; electrolyte transport parameters (kappa 1.4 S/m, D 3.7e-10 m2/s) are literature values (Nyman 2008) '
        'substituted as formulation levers; electrolyte mass uses literature density 1.2 g/cm3. Physical-cell qualification '
        '(formation, mechanical integrity, cycle life) is out of scope.', NOTE))
    return s

def dvpr_story():
    s = head('VBF-T8R1NOFORCE-DVPR-01', 'Design Verification Plan & Report',
             'case t8_r1_noforce | 2026-08-25 | virtual design verification (simulation-based)')
    s.append(Paragraph('Executed verification items', H2))
    s.append(md_table(['#', 'Item', 'Method / condition', 'Result', 'Criterion', 'Verdict'], [
        ['1', 'Gravimetric energy density', 'calc-energy contract caliber', '497.968 Wh/kg', '>= 446.18 Wh/kg', 'PASS'],
        ['2', '5C capacity retention', 'DFN 5C vs 1C, same params', '0.9896', '>= 0.90', 'PASS'],
        ['3', 'Cell mass', 'calc-energy layer mass x area', '38.99 g', '<= 40 g', 'PASS'],
        ['4', 'Plating under fast charge', '4C at 45 C, plating module', 'anode min +0.00635 V -> plated=false', 'plated = false', 'PASS (margin +6.3 mV - flagged)'],
        ['5', 'T_max fast charge', 'same run, lumped thermal', '326.81 K', '<= 333.15 K', 'PASS'],
        ['6', 'T_max 5C discharge', 'DFN 5C discharge, 25 C', '313.64 K', '<= 333.15 K', 'PASS'],
        ['7', '1C discharge capacity', 'DFN 1C discharge, 25 C', '5.2725 Ah', 'no entry-0 threshold (feeds 1-2)', 'PASS (informational)'],
    ], [U * f for f in [0.05, 0.17, 0.22, 0.22, 0.16, 0.18]]))
    s.append(Paragraph('Result: 7 / 7 executed items PASS. All five entry-0 stage2/stage3 criteria verified green on final design '
                       '(R6-V9); verification trail: log.jsonl round 6 propose -> evaluate (checked = 5, mechanical).', BODY))
    s.append(Paragraph('Uncovered conditions (honest N/A list)', H2))
    s.append(md_table(['Condition', 'Reason not executed', 'Materiality'], [
        ['Cycle life / long-term aging', 'not a task criterion; not simulated (no fabricated values)', 'medium - endurance drones accumulate cycles'],
        ['Nail penetration', 'requires physical experiment', 'low for pouch drone pack'],
        ['Crush / drop / vibration', 'requires physical experiment', 'medium - airframe vibration real; recommend qualification'],
        ['Thermal runaway abuse', 'protocol exists but out of task scope; not run', 'low (4C exam covers task thermal red line)'],
        ['Low-temperature retention (-20 C)', 'not a task criterion', 'medium - altitude operation is cold'],
        ['Self-discharge / storage', 'not simulated', 'low (single-flight-cycle pattern)'],
        ['DCR pulse measurement', 'calc-energy DCR is formulation estimate, not pulse test', 'low (power density far exceeds demand)'],
    ], W))
    s.append(Paragraph('Verification conclusion', H2))
    s.append(Paragraph(
        'The R6-V9 design satisfies all entry-0 acceptance criteria. Two margins deserve engineering attention before physical '
        'prototyping: (a) plating margin on 4C charge is +6.3 mV - thin; recommended mitigation is a 45 C charge-temperature '
        'floor and/or 3C charge in cold conditions (see DFMEA); (b) the 10 um separator and 6-8 um current collectors are '
        'aggressive and must be validated for pouch winding/lasering capability. These are recorded as risks, not criterion '
        'failures - the criteria themselves are all met.', BODY))
    return s

def dfmea_story():
    s = head('VBF-T8R1NOFORCE-DFMEA-01', 'DFMEA - Long-Endurance Drone Cell (design R6-V9)',
             'case t8_r1_noforce | 2026-08-25 | qualitative analysis - ratings are engineering judgement, not measurements')
    s.append(Paragraph(
        'Scope: cell-level failure modes of the final virtual design (NMC811/graphite pouch, 65 x 1891 mm electrodes, '
        '10 um separator, 6/8 um collectors, 4C fast charge at 45 C, 5C discharge). S/O/D 1-10 (10 worst); RPN = S x O x D.', NOTE))
    s.append(md_table(['#', 'Failure mode', 'Effect', 'S', 'Cause', 'O', 'Current controls', 'D', 'RPN', 'Recommended actions'], [
        ['1', 'Plating at charge end under cold conditions', 'dendrites, short risk, fade', '7',
         'anode-side transport polarization grows at low T (margin +6.3 mV at 45 C)', '4',
         'anode kinetics sized (r_n 1.2 um) + high-D electrolyte + h=60', '3', '84',
         'BMS 45 C charge-temperature floor; 3C max below 25 C; re-simulate plating at 25 C'],
        ['2', 'Separator (10 um) puncture / melt-shrink', 'internal short, thermal event', '8',
         'mechanical defect, dendrite (mode 1), winding tension', '3', 'thin separator is ED trade; no coating in simulation', '5', '120',
         'ceramic-coated separator qualification; supplier defect screening; abuse runs in next case'],
        ['3', 'Cu foil (6 um) tear during winding', 'local overcurrent, hot spot', '5',
         'ultra-thin foil slitting/winding stress', '4', 'design assumes intact collector', '3', '60',
         'hardened Cu alloy foil; tension control on 1.89 m winding'],
        ['4', 'Incomplete wetting of 1.891 m electrode', 'local starvation, early plating', '5',
         'long electrode path; 10 um separator restricts wicking', '4', 'high D reduces transport consequence, not wetting', '4', '80',
         'vacuum-fill with wetting hold; injection point design; formation capacity check'],
        ['5', 'Pack cooling short of h=60 assumption', 'over-temperature at 5C / 4C', '6',
         'pack airflow design shortfall', '2', 'h is design freedom; BMS cut-off not modeled', '3', '36',
         'ducted prop-wash CFD at pack level; derate 5C above 30 C ambient'],
        ['6', 'Current crowding on 1.89 m electrode', 'local heating, uneven aging', '4',
         'electrode length, single-tab assumption', '4', 'uniform-current assumption (1D model limit)', '3', '48',
         'multi-tab / full-width tab design study before build'],
        ['7', 'Nano-NMC exothermic release at elevated T', 'thermal runaway in abuse', '6',
         'high specific surface of nano-particles', '2', 'thermal red line 333.15 K enforced (T_max 326.8 K simulated)', '5', '60',
         'run overcharge/nail protocols in safety follow-up case'],
    ], [U * f for f in [0.04, 0.15, 0.12, 0.04, 0.15, 0.04, 0.15, 0.04, 0.06, 0.21]]))
    s.append(Paragraph('Top risks: #2 separator (RPN 120), #1 plating margin (RPN 84), #4 wetting (RPN 80). None are criterion '
                       'failures in the entry-0 sense - they are manufacturing/operational risks of an aggressive virtual design, '
                       'flagged for the next iteration (physical prototype case).', BODY))
    return s

def bom_pdf_story():
    s = head('VBF-T8R1NOFORCE-BOM-01', 'Bill of Materials (PDF release)',
             'case t8_r1_noforce | 2026-08-25 | dual caliber g/cell + kg/kWh on 19.4147 Wh (1C contract)')
    rows = []
    for name, g, in_contract, src in BOM_MAIN:
        rows.append([name, '%.4f' % g, '%.4f' % kg_kwh(g / 1000.0), 'YES' if in_contract else 'no', src])
    rows.append(['CONTRACT TOTAL (electrolyte excluded)', '%.4f' % (TOTAL_KG * 1000), '%.4f' % kg_kwh(TOTAL_KG),
                 'YES', 'r6_v9_energy.json mass_kg = 0.0389878 kg'])
    for name, g, src in BOM_SUPP:
        rows.append([name, '%.4f' % g, '%.4f' % kg_kwh(g / 1000.0), 'no', src])
    rows.append(['GRAND TOTAL (incl. electrolyte + binder/additive)', '%.4f' % grand, '%.4f' % kg_kwh(grand / 1000.0),
                 'no', 'annotated total'])
    s.append(md_table(['Component', 'Mass (g/cell)', 'Mass (kg/kWh)', 'In contract mass', 'Source / basis'],
                      rows, [U * f for f in [0.32, 0.12, 0.12, 0.10, 0.34]]))
    s.append(Paragraph('Contract mass = 38.99 g (criterion <= 40 g, PASS). Electrolyte, binder and additives are listed '
                       'supplementary: electrolyte density 1.2 g/cm3 and binder/additive fractions (2/2 wt% positive, 2/1 wt% '
                       'negative) are literature defaults, recorded as annotations, excluded from the calc-energy contract '
                       'caliber. Enclosure/tabs not modeled.', NOTE))
    return s

def calc_pdf_story():
    s = head('VBF-T8R1NOFORCE-CALC-01', 'Calculation Sheet (PDF release)',
             'case t8_r1_noforce | 2026-08-25 | contract-caliber energy + mass + retention + safety (all from bda tool outputs)')
    s.append(Paragraph('1. Energy calculation (calc-energy, contract caliber)', H2))
    s.append(md_table(['Metric', 'Value', 'Unit', 'Source'], [
        ['Discharge capacity (1C)', '5.2725', 'Ah', 'r6_v9_1c_dfn.json'],
        ['Energy (1C contract)', '19.4147', 'Wh', 'r6_v9_energy.json'],
        ['Gravimetric energy density', '497.968', 'Wh/kg', 'r6_v9_energy.json'],
        ['Volumetric energy density', '854.72', 'Wh/L', 'r6_v9_energy.json'],
        ['Cell mass (contract)', '0.0389878', 'kg', 'r6_v9_energy.json'],
        ['Electrode area', '0.122915', 'm2', 'r6_v9_energy.json'],
        ['Stack thickness', '1.848e-4', 'm', 'r6_v9_energy.json'],
        ['Midpoint voltage', '3.9335', 'V', 'r6_v9_energy.json'],
        ['DC resistance (formulation)', '0.0007275', 'ohm', 'r6_v9_energy.json'],
        ['Power density (formulation)', '150480.7', 'W/kg', 'r6_v9_energy.json'],
    ], [U * f for f in [0.30, 0.16, 0.12, 0.42]]))
    s.append(Paragraph('2. Mass budget (g/cell = layer_kg_m2 x area)', H2))
    mrows = [[n, '%.6f' % kgm2, '%.4f' % g_per_cell(kgm2)] for n, kgm2 in LAYERS]
    mrows.append(['TOTAL (contract)', '%.6f' % sum(x[1] for x in LAYERS), '%.4f' % (TOTAL_KG * 1000)])
    mrows.append(['Electrolyte (supplementary)', '-', '%.3f' % ELEC_G])
    s.append(md_table(['Layer', 'kg/m2', 'g/cell'], mrows, [U * f for f in [0.55, 0.20, 0.25]]))
    s.append(Paragraph('3. 5C capacity retention', H2))
    s.append(md_table(['Quantity', 'Value', 'Unit', 'Source'], [
        ['1C discharge capacity (25 C)', '5.2725', 'Ah', 'r6_v9_1c_dfn.json'],
        ['5C discharge capacity (25 C)', '5.2178', 'Ah', 'r6_v9_5c_dfn.json'],
        ['Capacity retention @5C', '0.989615', '-', 'r6_v9_retention.json'],
        ['Criterion (entry 0)', '0.90', '-', 'log.jsonl entry 0'],
        ['Verdict', 'PASS', '-', '0.9896 >= 0.90'],
        ['5C discharge T_max', '313.64', 'K', 'r6_v9_5c_dfn.json (<= 333.15 K)'],
    ], [U * f for f in [0.30, 0.16, 0.12, 0.42]]))
    s.append(Paragraph('4. Safety exam (4C charge at 45 C, lumped thermal + plating module)', H2))
    s.append(md_table(['Quantity', 'Value', 'Unit', 'Determination'], [
        ['Anode potential minimum', '0.0063475', 'V', 'plated = false -> PASS (margin +6.3 mV)'],
        ['Lithium plating flag', 'false', '-', 'PASS (criterion: plated = false)'],
        ['T_max 4C charge', '326.81', 'K', 'PASS (<= 333.15 K red line)'],
        ['T_max 5C discharge', '313.64', 'K', 'PASS (<= 333.15 K red line)'],
        ['T_max red line (entry 0)', '333.15', 'K', 'protocol default (task silent)'],
    ], [U * f for f in [0.30, 0.16, 0.12, 0.42]]))
    s.append(Paragraph('5. Notes and assumptions', H2))
    for line in [
        'Electrolyte excluded from contract mass (calc-energy contract caliber; parameter set carries no electrolyte density).',
        'Electrolyte transport kappa = 1.4 S/m, D = 3.7e-10 m2/s: literature values (Nyman 2008) set as formulation levers, recorded in R6 propose.',
        'Electrolyte density 1.2 g/cm3: literature value for BOM row / datasheet mass incl. electrolyte (48.54 g).',
        'Binder/additive fractions: literature defaults, excluded from contract mass.',
        'N/P = 0.641 capacity-density basis (neg 49.18 Ah/m2 vs pos 76.72 Ah/m2) - negative-limited by design.',
        'Formation recommendation: 0.1C CC to 4.2 V, 25 C, 2 cycles (design rec.; production tuning required).',
        'Not simulated: cycle life, low-T retention, abuse (nail/crush/overcharge), self-discharge - recorded N/A in DVPR.',
        'R5 V8 5C discharge IDA_CONV_FAIL recorded verbatim in log.jsonl; superseded by V9 design.',
    ]:
        s.append(Paragraph('* %s' % esc(line), BODY))
    return s

pdfs = [
    ('design_spec.pdf', ds_story(), 'Design Specification - t8_r1_noforce'),
    ('datasheet.pdf', dsh_story(), 'Cell Datasheet - t8_r1_noforce'),
    ('dvpr.pdf', dvpr_story(), 'DVPR - t8_r1_noforce'),
    ('dfmea.pdf', dfmea_story(), 'DFMEA - t8_r1_noforce'),
    ('bom.pdf', bom_pdf_story(), 'BOM - t8_r1_noforce'),
    ('calc.pdf', calc_pdf_story(), 'Calculation Sheet - t8_r1_noforce'),
]
pages = {}
for fname, story, title in pdfs:
    path, p = build_pdf(fname, story, title)
    pages[fname] = p
    print('pdf:', path, 'pages', p)

# ---------------- delivery index (md + pdf) ----------------
files = [
    ('Design Specification', 'design_spec.md', 'VBF-T8R1NOFORCE-DS-01', 'Core', 'md'),
    ('Design Specification (PDF)', 'design_spec.pdf', 'VBF-T8R1NOFORCE-DS-01', 'Core', 'pdf'),
    ('Bill of Materials (workbook)', 'bom.xlsx', 'VBF-T8R1NOFORCE-BOM-01', 'Core', 'xlsx'),
    ('Bill of Materials (PDF)', 'bom.pdf', 'VBF-T8R1NOFORCE-BOM-01', 'Core', 'pdf'),
    ('Cell Datasheet', 'datasheet.md', 'VBF-T8R1NOFORCE-DSH-01', 'Core', 'md'),
    ('Cell Datasheet (PDF)', 'datasheet.pdf', 'VBF-T8R1NOFORCE-DSH-01', 'Core', 'pdf'),
    ('Calculation Sheet (workbook)', 'calc.xlsx', 'VBF-T8R1NOFORCE-CALC-01', 'Core', 'xlsx'),
    ('Calculation Sheet (PDF)', 'calc.pdf', 'VBF-T8R1NOFORCE-CALC-01', 'Core', 'pdf'),
    ('DVPR', 'dvpr.md', 'VBF-T8R1NOFORCE-DVPR-01', 'Core', 'md'),
    ('DVPR (PDF)', 'dvpr.pdf', 'VBF-T8R1NOFORCE-DVPR-01', 'Core', 'pdf'),
    ('DFMEA', 'dfmea.md', 'VBF-T8R1NOFORCE-DFMEA-01', 'Core', 'md'),
    ('DFMEA (PDF)', 'dfmea.pdf', 'VBF-T8R1NOFORCE-DFMEA-01', 'Core', 'pdf'),
    ('Delivery Index', 'delivery_index.md', 'VBF-T8R1NOFORCE-IDX-01', 'Index', 'md'),
    ('Delivery Index (PDF)', 'delivery_index.pdf', 'VBF-T8R1NOFORCE-IDX-01', 'Index', 'pdf'),
    ('Rendered Design Report (ancillary)', 'report.html', 'VBF-T8R1NOFORCE-DS-02', 'Ancillary', 'html'),
]

def index_story():
    s = [Spacer(1, 2 * cm),
         Paragraph('Battery Design Deliverables', ParagraphStyle('T', fontName='Helvetica-Bold', fontSize=24, textColor=NAVY_C, alignment=1, leading=28)),
         Paragraph('Long-Endurance Drone Cell - case t8_r1_noforce', ParagraphStyle('ST', fontName='Helvetica', fontSize=13, textColor=BLUE_C, alignment=1, leading=18)),
         Spacer(1, 0.8 * cm),
         HRFlowable(width='60%', thickness=1.5, color=ORANGE_C, hAlign='CENTER'),
         Spacer(1, 1.2 * cm)]
    s.append(md_table(['Field', 'Value'], [
        ['Case', 't8_r1_noforce'],
        ['Document scheme', 'VBF-T8R1NOFORCE-&lt;DOC&gt;-01'],
        ['Document number (this file)', 'VBF-T8R1NOFORCE-IDX-01'],
        ['Delivery date', '2026-08-25'],
        ['Status', 'ALL PASS (entry-0 criteria; verify-deliverables)'],
        ['Prepared by', 'VBF Design Agent (virtual, simulation-based)'],
        ['Reviewed by', ''],
        ['Approved by', ''],
        ['Ancillary', 'report.html - VBF-T8R1NOFORCE-DS-02'],
        ['Audit chain', 'log.jsonl (workspace root; entries 0..final)'],
    ], [U * f for f in [0.30, 0.70]]))
    s.append(PageBreak())
    s.append(Paragraph('File table', H2))
    rows = []
    for i, (name, fname, num, role, kind) in enumerate(files, 1):
        pg = str(pages[fname]) if fname in pages else 'n/a'
        rows.append([str(i), name, fname, num, role, pg])
    s.append(md_table(['#', 'Document', 'File', 'Number', 'Role', 'Pages'],
                      rows, [U * f for f in [0.05, 0.27, 0.25, 0.26, 0.10, 0.07]]))
    s.append(Spacer(1, 1.2 * cm))
    s.append(Paragraph('Verification note', H2))
    s.append(Paragraph(
        'All criterion-grade values in these deliverables are mechanical copies of bda tool outputs '
        '(r6_v9_energy.json, r6_v9_1c_dfn.json, r6_v9_5c_dfn.json, r6_v9_4c_charge45_plating.json, r6_v9_retention.json) '
        'with per-row source annotations; the audit chain is log.jsonl entry 0 through final.', BODY))
    s.append(Spacer(1, 1.2 * cm))
    s.append(Paragraph('Signatures', H2))
    s.append(md_table(['Role', 'Name', 'Signature', 'Date'],
                      [['Prepared', '', '', ''], ['Reviewed', '', '', ''], ['Approved', '', '', '']],
                      [U * f for f in [0.25, 0.25, 0.30, 0.20]]))
    return s

path, p = build_pdf('delivery_index.pdf', index_story(), 'Delivery Index - t8_r1_noforce')
pages['delivery_index.pdf'] = p
print('pdf:', path, 'pages', p)

lines = ['# Delivery Index - Battery Design Case t8_r1_noforce', '',
         'Document number: VBF-T8R1NOFORCE-IDX-01 | Delivery date: 2026-08-25', '',
         '## Cover', '',
         '| Field | Value |', '|---|---|',
         '| Case | t8_r1_noforce |',
         '| Title | Battery design deliverables - long-endurance drone cell (5C discharge capable, >= 446.18 Wh/kg, <= 40 g) |',
         '| Document scheme | VBF-T8R1NOFORCE-&lt;DOC&gt;-01 |',
         '| Status | ALL PASS (entry-0 criteria; verify-deliverables) |',
         '| Prepared by | VBF Design Agent (virtual, simulation-based) |',
         '| Reviewed by | (blank) |', '| Approved by | (blank) |',
         '| Release date | 2026-08-25 |',
         '| Ancillary | report.html (rendered 7-section report, VBF-T8R1NOFORCE-DS-02) |',
         '| Audit chain | log.jsonl (workspace root; entry 0 criteria, rounds R1-R6, evaluate, endorse, final) |',
         '', '## File table', '',
         '| # | Document | File | Number | Role | Pages |', '|---|---|---|---|---|---|']
for i, (name, fname, num, role, kind) in enumerate(files, 1):
    if fname in pages:
        pg = str(pages[fname])
    elif fname == 'delivery_index.md':
        pg = 'see PDF'
    else:
        pg = 'n/a'
    lines.append('| %d | %s | %s | %s | %s | %s |' % (i, name, fname, num, role, pg))
lines += ['', '## Signatures', '',
          '| Role | Name | Signature | Date |', '|---|---|---|---|',
          '| Prepared | | | |', '| Reviewed | | | |', '| Approved | | | |',
          '', '## Verification note', '',
          'All criterion-grade values in these deliverables are mechanical copies of bda tool outputs '
          '(r6_v9_energy.json, r6_v9_1c_dfn.json, r6_v9_5c_dfn.json, r6_v9_4c_charge45_plating.json, '
          'r6_v9_retention.json) with per-row source annotations; the audit chain is log.jsonl entry 0 through final.',
          '']
with open(os.path.join(DLV, 'delivery_index.md'), 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))
print('md:', os.path.join(DLV, 'delivery_index.md'))
print('ALL DELIVERABLES GENERATED')
