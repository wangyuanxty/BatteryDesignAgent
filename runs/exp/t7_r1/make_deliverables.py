# -*- coding: utf-8 -*-
"""One-off closing script for case t7_r1 deliverables:
bom.xlsx, calc.xlsx + PDF release versions of all 7 deliverable classes
(md -> PDF via reportlab; xlsx -> PDF via openpyxl read + reportlab tables).
All numeric values copied from tool output files (see sources in comments);
this script writes deliverables only, it reads: cell/r5_v11_*.json + params.
"""
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable)

CASE = Path("D:/research/degradation_prognostics/Battery_Design_Agent/runs/exp/t7_r1")
DLV = CASE / "deliverables"
CELL = CASE / "cell"
ENERGY = json.load(open(CELL / "r5_v11_energy.json", encoding="utf-8"))

DEEP = colors.HexColor("#14283C")
MED = colors.HexColor("#1E5A8A")
COPPER = colors.HexColor("#C97B3D")

# ---------------------------------------------------------------- xlsx part
def style_sheet(ws, widths, header_fill="14283C"):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=header_fill)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

# ---- BOM
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BOM"
ws.append(["Component", "Mass g/cell", "kg/kWh", "Formula / source"])
KWH = ENERGY["energy_wh"] / 1000.0
bom_rows = [
    ("Positive active material (NMC811)", 16.168, None, "pos coating 16.842 g x 0.96 AM mass fraction (literature estimate); coating = cell/r5_v11_energy.json:layer_kg_m2 positive 0.163994 x area 0.1027 m2"),
    ("Positive conductive additive (carbon black)", 0.337, None, "pos coating x 0.02 (literature estimate)"),
    ("Positive binder (PVDF)", 0.337, None, "pos coating x 0.02 (literature estimate)"),
    ("Positive coating total (set caliber)", 16.842, None, "parameter set: 75.6um x (1-0.335) x 3262 kg/m3 x 0.1027 m2 = layer_kg_m2 x area"),
    ("Negative active material (graphite)", 11.961, None, "neg coating 12.525 g x 0.955 (literature estimate); coating = layer_kg_m2 negative 0.121955 x 0.1027 m2"),
    ("Negative binder (CMC)", 0.188, None, "neg coating x 0.015 (literature estimate)"),
    ("Negative binder (SBR)", 0.376, None, "neg coating x 0.030 (literature estimate)"),
    ("Negative coating total (set caliber)", 12.525, None, "parameter set: 115um x (1-0.36) x 1657 kg/m3 x 0.1027 m2"),
    ("Separator (polyolefin)", 0.259, None, "layer_kg_m2 separator 0.00252492 x 0.1027 m2"),
    ("Electrolyte (EC/EMC + LiPF6)", 8.918, None, "pore volume 7.432 cm3 x 1.2 g/cm3 (literature density) x fill factor 1.0; pores = (75.6e-6*0.335 + 12e-6*0.47 + 115e-6*0.36) x 0.1027 m3"),
    ("Positive current collector (Al foil)", 4.437, None, "layer_kg_m2 pos cc 0.0432 x 0.1027 (16um x 2700 kg/m3)"),
    ("Negative current collector (Cu foil)", 11.042, None, "layer_kg_m2 neg cc 0.10752 x 0.1027 (12um x 8960 kg/m3)"),
    ("Enclosure, tabs, terminals", "Not modeled", "Not modeled", "no parameter in set"),
]
for comp, g, _, src in bom_rows:
    if isinstance(g, str):
        ws.append([comp, g, g, src])
    else:
        ws.append([comp, round(g, 3), round(g / 1000.0 / KWH, 4), src])
ws.append(["Total (contract caliber, electrolyte excluded)", round(ENERGY["mass_kg"] * 1000, 3), round(ENERGY["mass_kg"] / KWH, 4), "cell/r5_v11_energy.json:mass_kg = 0.0451052 kg; energy = 21.7970 Wh"])
ws.append(["Total (with electrolyte estimate)", round(ENERGY["mass_kg"] * 1000 + 8.918, 3), round((ENERGY["mass_kg"] + 8.918e-3) / KWH, 4), "contract mass + electrolyte row above (estimate, not adjudicated)"])
style_sheet(ws, [46, 12, 12, 90])
wb.save(DLV / "bom.xlsx")

# ---- CALC (5 sheets)
wb = openpyxl.Workbook()
def sheet(name, header, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r)
    style_sheet(ws, widths)
    return ws

sheet("1 Inputs", ["Parameter", "Value", "Unit", "Source"], [
    ["Base parameter set", "Chen2020 (NMC811/graphite)", "-", "entry 0 meta.base (anchor table; task names no electrode system)"],
    ["Negative particle radius", 1.5e-6, "m", "params_v11.json (design)"],
    ["Positive particle radius", 3.5e-6, "m", "params_v11.json (design)"],
    ["Electrolyte conductivity", 2.5, "S/m", "params_v11.json (design)"],
    ["Cation transference number", 0.55, "-", "params_v11.json (design)"],
    ["Electrolyte diffusivity", 6.0e-10, "m2/s", "params_v11.json (design)"],
    ["Negative electrode thickness", 115e-6, "m", "params_v11.json (design)"],
    ["Negative electrode porosity", 0.36, "-", "params_v11.json (design)"],
    ["Total heat transfer coefficient h", 50.0, "W/(m2 K)", "params_v11.json (design)"],
    ["Nominal cell capacity", 6.0313, "Ah", "params_v11.json (= measured 1C, mechanical rule propose_r2b)"],
    ["Positive electrode thickness", 75.6e-6, "m", "Chen2020 parameter set"],
    ["Positive electrode porosity", 0.335, "-", "Chen2020 parameter set"],
    ["Separator thickness / porosity", "12e-6 / 0.47", "m / -", "Chen2020 parameter set"],
    ["Electrode height x width", "0.065 x 1.58", "m", "Chen2020 parameter set"],
    ["Electrode area", 0.1027, "m2", "height x width = 0.1027 (matches calc-energy area_m2)"],
    ["Voltage window", "2.5 - 4.2", "V", "Chen2020 parameter set (cut-offs)"],
    ["Max conc positive / negative", "63104 / 33133", "mol/m3", "Chen2020 parameter set"],
    ["Positive/negative electrode density", "3262 / 1657", "kg/m3", "Chen2020 parameter set"],
    ["Pos/neg CC thickness + density", "16um Al 2700 / 12um Cu 8960", "-", "Chen2020 parameter set"],
    ["SEI kinetic rate constant", 1.0e-12, "m/s", "Chen2020 parameter set"],
], [38, 22, 12, 80])

sheet("2 Capacity-Energy", ["Quantity", "Value", "Unit", "Formula / source"], [
    ["Measured 1C capacity (DFN)", 6.032209, "Ah", "cell/r5_v11_1c_dfn.json:capacity_ah"],
    ["Discharge energy", 21.796964, "Wh", "cell/r5_v11_energy.json:energy_wh (integral V x I over 1C discharge)"],
    ["Midpoint voltage", 3.723722, "V", "cell/r5_v11_energy.json:midpoint_voltage_v"],
    ["Average voltage", 3.6135, "V", "energy / capacity = 21.7970 / 6.0322 (mechanical)"],
    ["DC resistance", 0.00288198, "ohm", "cell/r5_v11_energy.json:dcr_ohm"],
    ["Power density", 32559.66, "W/kg", "cell/r5_v11_energy.json:power_density_w_kg"],
    ["Areal capacity (cell)", 58.7356, "Ah/m2", "6.032209 Ah / 0.1027 m2"],
], [38, 16, 12, 90])

sheet("3 Energy Density", ["Quantity", "Value", "Unit", "Formula / source"], [
    ["Positive coating mass", 16.842, "g", "layer_kg_m2 0.163994 x 0.1027 (r5_v11_energy.json)"],
    ["Negative coating mass", 12.525, "g", "layer_kg_m2 0.121955 x 0.1027"],
    ["Al current collector", 4.437, "g", "layer_kg_m2 0.0432 x 0.1027"],
    ["Cu current collector", 11.042, "g", "layer_kg_m2 0.10752 x 0.1027"],
    ["Separator", 0.259, "g", "layer_kg_m2 0.00252492 x 0.1027"],
    ["Total mass (contract, electrolyte excluded)", 45.105, "g", "cell/r5_v11_energy.json:mass_kg = 0.0451052 kg"],
    ["Gravimetric energy density", 483.247, "Wh/kg", "21.796964 Wh / 0.0451052 kg -> threshold 327.18 PASS"],
    ["Cell volume", 2.368262e-5, "m3", "cell/r5_v11_energy.json:volume_m3"],
    ["Volumetric energy density", 920.378, "Wh/L", "21.796964 Wh / 23.68262 cm3"],
    ["Stack thickness", 230.6e-6, "m", "75.6 + 12 + 115 + 16 + 12 um (cell/r5_v11_energy.json:thickness_m)"],
    ["Electrolyte mass (estimate, excluded)", 8.918, "g", "pore volume 7.432 cm3 x 1.2 g/cm3 (literature)"],
    ["Reference ED incl. electrolyte (not adjudicated)", 403.47, "Wh/kg", "21.796964 Wh / 0.0540234 kg (reference only)"],
], [44, 14, 12, 90])

sheet("4 N-P and Mass", ["Quantity", "Value", "Unit", "Formula / source"], [
    ["Negative theoretical areal capacity", 65.357, "Ah/m2", "115e-6 x (1-0.36) x 33133 mol/m3 x 96485/3600"],
    ["Positive theoretical areal capacity (full range)", 85.027, "Ah/m2", "75.6e-6 x (1-0.335) x 63104 mol/m3 x 96485/3600"],
    ["Cell measured areal capacity", 58.7356, "Ah/m2", "6.032209 Ah / 0.1027 m2 (= positive usable capacity)"],
    ["Practical N/P ratio", 1.1127, "-", "65.357 / 58.7356 (negative theoretical / positive usable)"],
    ["Theoretical full-range N/P (reference)", 0.7687, "-", "65.357 / 85.027 (positive full Li range never used below 4.2 V)"],
    ["Positive utilization in 2.5-4.2 V window", 0.6908, "-", "58.7356 / 85.027"],
], [46, 14, 12, 90])

sheet("5 Process", ["Parameter", "Value", "Unit", "Formula / source"], [
    ["Positive areal density", 164.0, "g/m2", "75.6e-6 x (1-0.335) x 3262 (design-spec process formula)"],
    ["Negative areal density", 122.0, "g/m2", "115e-6 x (1-0.36) x 1657"],
    ["Positive compaction density", 2.17, "g/cm3", "3262 x 0.665 / 1000"],
    ["Negative compaction density", 1.06, "g/cm3", "1657 x 0.64 / 1000"],
    ["Electrolyte fill amount", 8.92, "g", "7.432 cm3 x 1.2 g/cm3 x 1.0 fill factor (literature density; fill factor assumption)"],
    ["Formation recommendation", "0.1C CC to 4.2 V, 25 C, 2 cycles", "-", "design recommended value; production value requires tuning"],
], [40, 24, 12, 90])

wb.save(DLV / "calc.xlsx")
print("xlsx written")

# ---------------------------------------------------------------- PDF part
def md_to_pdf(md_path, pdf_path, doc_number, subtitle):
    text = md_path.read_text(encoding="utf-8")
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    styles = {
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, textColor=DEEP, spaceAfter=4),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, textColor=MED, spaceBefore=8, spaceAfter=3),
        "h3": ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=9.5, textColor=COPPER, spaceBefore=6, spaceAfter=2),
        "p": ParagraphStyle("p", fontName="Helvetica", fontSize=8, leading=11, spaceAfter=3),
        "li": ParagraphStyle("li", fontName="Helvetica", fontSize=8, leading=11, leftIndent=10, spaceAfter=2),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=6.5, leading=8.5),
        "cellb": ParagraphStyle("cellb", fontName="Helvetica-Bold", fontSize=6.8, leading=8.5, textColor=colors.white),
    }
    story = []
    # blueprint cover block
    cover = Table([[Paragraph(doc_number, styles["cellb"])],
                   [Paragraph(md_path.stem.replace("_", " ").title(), ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=17, textColor=colors.white, leading=20))],
                   [Paragraph(subtitle, ParagraphStyle("s", fontName="Helvetica", fontSize=8.5, textColor=colors.HexColor("#B8C8D8"), leading=11))],
                   [Paragraph("Virtual Battery Factory · case t7_r1 · 2026-08-25", ParagraphStyle("d", fontName="Helvetica", fontSize=7, textColor=colors.HexColor("#C97B3D")))],
                  ], colWidths=[170 * mm])
    cover.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), DEEP),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, COPPER),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(cover)
    story.append(Spacer(1, 5 * mm))

    def cell_par(s):
        s = s.replace("**", "")
        return Paragraph(s if s.strip() else " ", styles["cell"])

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("|"):
            tbl_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip().strip("|").split("|")
                tbl_rows.append([c.strip() for c in row])
                i += 1
            data = [[Paragraph(c.replace("**", ""), styles["cellb"]) for c in tbl_rows[0]]]
            for r in tbl_rows[1:]:
                if all(re.fullmatch(r"-{2,}", c.strip()) for c in r):
                    continue
                data.append([cell_par(c) for c in r])
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), MED),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9FB2C4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 2.5 * mm))
            continue
        if ln.startswith("### "):
            story.append(Paragraph(ln[4:], styles["h3"])); i += 1; continue
        if ln.startswith("## "):
            story.append(Paragraph(ln[3:], styles["h2"])); i += 1; continue
        if ln.startswith("# "):
            story.append(Paragraph(ln[2:], styles["h1"])); i += 1; continue
        if re.fullmatch(r"-{3,}", ln.strip()):
            story.append(HRFlowable(width="100%", color=COPPER, thickness=0.7)); i += 1; continue
        if ln.startswith("- "):
            story.append(Paragraph("• " + ln[2:].replace("**", ""), styles["li"])); i += 1; continue
        story.append(Paragraph(ln.replace("**", ""), styles["p"])); i += 1
    doc.build(story)

md_pairs = [
    ("design_spec.md", "VBF-T7R1-DS-01", "Cell Design Specification"),
    ("datasheet.md", "VBF-T7R1-DSH-01", "Technical Datasheet"),
    ("dvpr.md", "VBF-T7R1-DVPR-01", "Design Verification Plan and Report (Virtual)"),
    ("dfmea.md", "VBF-T7R1-DFMEA-01", "Design FMEA (Qualitative)"),
    ("delivery_index.md", "VBF-T7R1-IDX-01", "Delivery Index"),
]
for name, num, sub in md_pairs:
    md_to_pdf(DLV / name, DLV / name.replace(".md", ".pdf"), num, sub)
    print("pdf:", name.replace(".md", ".pdf"))

def xlsx_to_pdf(xlsx_path, pdf_path, doc_number, subtitle):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    s_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=6.2, leading=8)
    s_h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=6.8, leading=8.5, textColor=colors.white)
    s_sec = ParagraphStyle("sec", fontName="Helvetica-Bold", fontSize=11, textColor=MED, spaceBefore=6, spaceAfter=3)
    story = [Paragraph(doc_number, s_h), Paragraph(subtitle, ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, textColor=colors.white)), Spacer(1, 4 * mm)]
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        story.append(Paragraph(ws.title, s_sec))
        data = [[Paragraph(str(c) if c is not None else " ", s_h) for c in rows[0]]]
        for r in rows[1:]:
            data.append([Paragraph(str(c) if c is not None else " ", s_cell) for c in r])
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), MED),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9FB2C4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(t)
        story.append(Spacer(1, 3 * mm))
    doc.build(story)

xlsx_to_pdf(DLV / "bom.xlsx", DLV / "bom.pdf", "VBF-T7R1-BOM-01", "Bill of Materials")
xlsx_to_pdf(DLV / "calc.xlsx", DLV / "calc.pdf", "VBF-T7R1-CALC-01", "Design Calculation Sheet")
print("bom.pdf + calc.pdf written")
print("done")
