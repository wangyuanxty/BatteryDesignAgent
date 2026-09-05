import json, os
from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import ParagraphStyle
from openpyxl import load_workbook

ROOT = Path(r'D:\research\degradation_prognostics\Battery_Design_Agent')
CASE = ROOT / 'runs' / 't5_r1_mimo'
LOG = CASE / 'log.jsonl'

# append final entry
lines = [json.loads(line) for line in LOG.read_text(encoding='utf-8').splitlines() if line.strip()]
has_final = any(entry.get('action') == 'final' for entry in lines)
if not has_final:
    final = {
        'action': 'final',
        'recommendation': 'Candidate E2 satisfies all three task targets: energy density >= 500.94 Wh/kg, no lithium plating under 4C charging, and maximum lumped temperature <= 60C.',
        'verdict': 'achieved'
    }
    with LOG.open('a', encoding='utf-8') as f:
        f.write(json.dumps(final, ensure_ascii=False) + '\n')
    lines.append(final)

# numbering helper
VBF_PREFIX = 'VBF-T5R1MIMO'

INDEX = {
    'design_spec.md': ('DS', 'Design Specification'),
    'bom.xlsx': ('BOM', 'Bill of Materials'),
    'datasheet.md': ('DSH', 'Datasheet'),
    'calc.xlsx': ('CALC', 'Calculation Sheet'),
    'dvpr.md': ('DVPR', 'Verification Report'),
    'dfmea.md': ('DFMEA', 'Failure Analysis'),
    'delivery_index.md': ('INDEX', 'Delivery Index'),
}

# update delivery_index with explicit VBF numbering
idx_path = CASE / 'delivery_index.md'
idx_text = idx_path.read_text(encoding='utf-8')
if 'VBF-' not in idx_text:
    rows = '\n'.join([f"| {name} | {VBF_PREFIX}-{code}-001 | {name.split('.')[-1].upper()} | {title} |" for name, (code, title) in INDEX.items()])
    addition = f"""
## Official VBF numbering

| Deliverable | Number | Format | Title |
|---|---|---|---|
{rows}
"""
    idx_text += addition
    idx_path.write_text(idx_text, encoding='utf-8')

# render markdown to pdf
styles = getSampleStyleSheet()
body = ParagraphStyle('Body', parent=styles['BodyText'], fontName='Helvetica', fontSize=9, leading=12, spaceAfter=6)
code = ParagraphStyle('Code', parent=styles['Code'], fontName='Courier', fontSize=8, leading=10, backColor=colors.whitesmoke)

def render_md_to_pdf(md_path: Path, pdf_path: Path):
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=14*mm, rightMargin=14*mm, topMargin=14*mm, bottomMargin=14*mm)
    story = []
    for raw_line in md_path.read_text(encoding='utf-8').splitlines():
        line = raw_line.rstrip()
        if not line:
            story.append(Spacer(1, 4))
            continue
        if line.startswith('#'):
            text = line.lstrip('#').strip()
            story.append(Paragraph(f"<b>{text}</b>", styles['Heading2']))
            continue
        if line.startswith('|'):
            # simple table
            cells = [c.strip() for c in line.split('|')[1:-1]]
            story.append(Table([cells], hAlign='LEFT', colWidths=[doc.width/len(cells)]*len(cells), style=TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#D9E1F2')),
                ('GRID', (0,0), (-1,-1), 0.3, colors.grey),
                ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('VALIGN', (0,0), (-1,-1), 'TOP')
            ])))
            continue
        story.append(Paragraph(line.replace('<', '&lt;').replace('>', '&gt;').replace('`', '"'), body))
    if story:
        doc.build(story)


# render xlsx to pdf (sheet-by-sheet tables)
def render_xlsx_to_pdf(xlsx_path: Path, pdf_path: Path):
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=14*mm, rightMargin=14*mm, topMargin=14*mm, bottomMargin=14*mm)
    story = []
    wb = load_workbook(str(xlsx_path), data_only=True)
    for ws in wb.worksheets:
        story.append(Paragraph(f"<b>{ws.title}</b>", styles['Heading2']))
        data = []
        for row in ws.iter_rows(values_only=True):
            data.append([str(v) if v is not None else '' for v in row])
        if data:
            table = Table(data, hAlign='LEFT', repeatRows=1, colWidths=[doc.width/len(data[0])]*len(data[0]))
            table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#D9E1F2')),
                ('GRID', (0,0), (-1,-1), 0.25, colors.grey),
                ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('VALIGN', (0,0), (-1,-1), 'TOP')
            ]))
            story.append(table)
        story.append(Spacer(1, 8))
    if story:
        doc.build(story)

pdf_dir = CASE
pdf_dir.mkdir(exist_ok=True)
for name in ['design_spec.md', 'datasheet.md', 'dvpr.md', 'dfmea.md', 'delivery_index.md']:
    render_md_to_pdf(CASE / name, pdf_dir / f"{Path(name).stem}.pdf")
for name in ['bom.xlsx', 'calc.xlsx']:
    render_xlsx_to_pdf(CASE / name, pdf_dir / f"{Path(name).stem}.pdf")

print('Final entry appended and PDF release versions generated.')
